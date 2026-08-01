import json
import urllib.parse
from collections.abc import Sequence
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from sqlalchemy import create_engine, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from medical_audit_kb.api.agent_store import (
    AGENT_ID_PREFIX,
    InMemoryAgentStore,
    SqlAlchemyAgentStore,
)
from medical_audit_kb.api.analytics_upload_store import (
    ANALYTICS_UPLOAD_ID_PREFIX,
    InMemoryAnalyticsUploadStore,
    SqlAlchemyAnalyticsUploadStore,
)
from medical_audit_kb.api.app import ApiState, create_app
from medical_audit_kb.api.audit_finding_store import SqlAlchemyAuditFindingStore
from medical_audit_kb.api.auth_user_store import (
    AUTH_ROLE_ASSIGNMENT_ID_PREFIX,
    AUTH_USER_ID_PREFIX,
    SqlAlchemyAuthUserStore,
)
from medical_audit_kb.api.chat_models import ChatModelAlias, chat_model_config_from_env
from medical_audit_kb.api.document_upload_ingestion import SqlAlchemyDocumentUploadIndexer
from medical_audit_kb.api.document_upload_store import (
    DOCUMENT_UPLOAD_ID_PREFIX,
    InMemoryDocumentUploadStore,
    SqlAlchemyDocumentUploadStore,
)
from medical_audit_kb.api.project_member_store import (
    DEFAULT_PROJECT_MEMBERS_BY_PROJECT,
    DEFAULT_PROJECT_PAYLOADS,
    PROJECT_MEMBER_ID_PREFIX,
    InMemoryProjectMemberStore,
    SqlAlchemyProjectMemberStore,
    visible_project_keys,
)
from medical_audit_kb.api.query_history_store import (
    InMemoryQueryHistoryStore,
    SqlAlchemyQueryHistoryStore,
)
from medical_audit_kb.api.review_task_store import SqlAlchemyReviewTaskStore
from medical_audit_kb.api.routes_documents import DocumentUploadItem
from medical_audit_kb.core.config import (
    DocumentStorageSettings,
    DocumentUploadGovernanceSettings,
    DocumentUploadIndexingSettings,
    KnowledgeQuerySettings,
    ModelProviderSettings,
)
from medical_audit_kb.db.models import (
    AnalyticsUploadRecord,
    AuditDataSnapshot,
    AuditFinding,
    AuditProject,
    AuditProjectMember,
    AuditRule,
    AuditRun,
    AuditTask,
    ReviewTask,
    RuleVersion,
)
from medical_audit_kb.domain.constants import SourceCollection
from medical_audit_kb.domain.source_collection_registry import (
    KNOWLEDGE_QUERY_CONTRACT_VERSION,
    SYSTEM_SOURCE_COLLECTION_DEFINITIONS,
)
from medical_audit_kb.generation.answer_providers import AnswerProviderError
from medical_audit_kb.generation.citations import Citation
from medical_audit_kb.indexing.bm25_index import BM25Document, InMemoryBM25Index
from medical_audit_kb.indexing.embeddings import DeterministicFakeEmbeddingProvider
from medical_audit_kb.indexing.index_activation import (
    IndexActivationError,
    IndexActivationResult,
    IndexRollbackResult,
)
from medical_audit_kb.indexing.vector_index import (
    ChunkEmbeddingInput,
    InMemoryVectorIndex,
    build_chunk_embedding_records,
)
from medical_audit_kb.ocr.unlimited_ocr import UnlimitedOcrResult
from medical_audit_kb.retrieval.hybrid_search import HybridSearchEngine
from medical_audit_kb.retrieval.rerank import FakeRerankProvider

PROJECT_NAME_HEADER = (
    "%E5%8C%BB%E4%BF%9D%E5%9F%BA%E9%87%91"
    "%E4%BD%BF%E7%94%A8%E5%90%88%E8%A7%84"
    "%E4%B8%93%E9%A1%B9%E8%87%AA%E6%9F%A5"
)


def test_health_endpoint_returns_api_status(tmp_path: Path) -> None:
    client = TestClient(create_app(_api_state(tmp_path)))

    response = client.get("/health")

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok"
    assert body["data_root"] == str(tmp_path / "data")


def test_deployment_metadata_reports_sha_without_runtime_side_effects(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)
    deploy_sha = "5e603f85aa11bb22cc33dd44ee55ff6677889900"
    monkeypatch.setenv("MEDICAL_AUDIT_DEPLOY_SHA", deploy_sha.upper())
    client = TestClient(create_app(state))

    response = client.get("/deployment/metadata")
    versioned_response = client.get("/api/v1/deployment/metadata")
    backend_response = client.get("/api/backend/deployment/metadata")

    assert response.status_code == 200
    assert versioned_response.status_code == 200
    assert backend_response.status_code == 200
    body = response.json()
    serialized = json.dumps(body, ensure_ascii=False)
    assert body["status"] == "deployment_metadata_available"
    assert body["evidence_grade"] == "L1-public-or-runtime"
    assert body["deploy_sha_status"] == "set"
    assert body["deploy_sha"] == deploy_sha
    assert body["deploy_sha_source"] == "env"
    assert body["required_report_fields"] == {
        "expected_deploy_sha": deploy_sha,
        "current_deploy_sha": deploy_sha,
        "deploy_sha_status": "set",
    }
    assert body["boundaries"] == {
        "production_write": False,
        "production_env_write": False,
        "provider_call": False,
        "object_storage_write": False,
        "secret_values_reported": False,
        "allowed_http_methods": ["GET"],
        "non_get_http_methods_allowed": False,
    }
    assert "MEDICAL_AUDIT_DEPLOY_SHA" not in serialized
    assert state.operation_logs == []


def test_deployment_metadata_reads_default_deploy_sha_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)
    deploy_sha = "abcdef0123456789abcdef0123456789abcdef01"
    (tmp_path / ".deploy-sha").write_text(f"{deploy_sha}\n", encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("MEDICAL_AUDIT_DEPLOY_SHA", raising=False)
    monkeypatch.delenv("MEDICAL_AUDIT_DEPLOY_SHA_FILE", raising=False)
    client = TestClient(create_app(state))

    response = client.get("/deployment/metadata")

    assert response.status_code == 200
    body = response.json()
    assert body["deploy_sha_status"] == "set"
    assert body["deploy_sha"] == deploy_sha
    assert body["deploy_sha_source"] == "default_file"


def test_deployment_metadata_is_protected_by_controlled_api_auth(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)
    monkeypatch.setenv(
        "MEDICAL_AUDIT_DEPLOY_SHA",
        "5e603f85aa11bb22cc33dd44ee55ff6677889900",
    )
    client = TestClient(create_app(state, enforce_controlled_api_auth=True))

    anonymous_response = client.get("/api/v1/deployment/metadata")
    authed_response = client.get(
        "/api/v1/deployment/metadata",
        headers={
            "X-User-Id": "admin-1",
            "X-Role": "admin",
            "X-Tenant-Id": "hospital-demo",
        },
    )

    assert anonymous_response.status_code == 401
    assert authed_response.status_code == 200
    assert state.operation_logs[-1]["payload"]["path"] == "/api/v1/deployment/metadata"


def test_versioned_api_prefix_serves_backend_routes_and_auth_middleware(
    tmp_path: Path,
) -> None:
    client = TestClient(create_app(_api_state(tmp_path), enforce_controlled_api_auth=True))

    health_response = client.get("/api/v1/health")
    backend_proxy_health_response = client.get("/api/backend/health")
    anonymous_projects_response = client.get("/api/v1/projects")
    anonymous_backend_proxy_search_response = client.get("/api/backend/index/search-backend")
    search_backend_response = client.get(
        "/api/v1/index/search-backend",
        headers={
            "X-User-Id": "admin-1",
            "X-Role": "admin",
            "X-Tenant-Id": "hospital-demo",
        },
    )
    backend_proxy_search_response = client.get(
        "/api/backend/index/search-backend",
        headers={
            "X-User-Id": "admin-1",
            "X-Role": "admin",
            "X-Tenant-Id": "hospital-demo",
        },
    )

    assert health_response.status_code == 200
    assert health_response.json()["status"] == "ok"
    assert backend_proxy_health_response.status_code == 200
    assert backend_proxy_health_response.json()["status"] == "ok"
    assert anonymous_projects_response.status_code == 401
    assert anonymous_backend_proxy_search_response.status_code == 401
    assert search_backend_response.status_code == 200
    assert search_backend_response.json()["backend"] == "none"
    assert backend_proxy_search_response.status_code == 200
    assert backend_proxy_search_response.json()["backend"] == "none"


def test_static_export_serves_portal_without_swallowing_api_404(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    static_root = tmp_path / "web-out"
    _write_text(static_root / "index.html", "<html><body>AuditScope Portal</body></html>")
    _write_text(static_root / "agents.html", "<html>Agents Portal</html>")
    _write_text(static_root / "agents.txt", "Agents RSC")
    _write_text(static_root / "analytics.html", "<html>Analytics Portal</html>")
    _write_text(static_root / "analytics.txt", "Analytics RSC")
    _write_text(static_root / "documents" / "index.html", "<html>Documents App</html>")
    _write_text(static_root / "graph.html", "<html>Graph Portal</html>")
    _write_text(static_root / "graph.txt", "Graph RSC")
    _write_text(static_root / "workspace.html", "<html>Workspace App</html>")
    _write_text(static_root / "_next" / "static" / "chunk.js", "console.log('ok');")
    monkeypatch.setenv("MEDICAL_AUDIT_WEB_STATIC_ROOT", str(static_root))
    client = TestClient(create_app(_api_state(tmp_path)))

    root_response = client.get("/")
    agents_response = client.get("/agents")
    agents_rsc_response = client.get("/agents.txt")
    analytics_response = client.get("/analytics")
    analytics_rsc_response = client.get("/analytics.txt")
    documents_response = client.get("/documents")
    graph_response = client.get("/graph")
    graph_rsc_response = client.get("/graph.txt")
    workspace_response = client.get("/workspace")
    asset_response = client.get("/_next/static/chunk.js")
    missing_asset_response = client.get("/_next/static/missing.js")
    missing_api_response = client.get("/api/v1/not-found")

    assert root_response.status_code == 200
    assert "AuditScope Portal" in root_response.text
    assert agents_response.status_code == 200
    assert "Agents Portal" in agents_response.text
    assert agents_rsc_response.status_code == 200
    assert "Agents RSC" in agents_rsc_response.text
    assert analytics_response.status_code == 200
    assert "Analytics Portal" in analytics_response.text
    assert analytics_rsc_response.status_code == 200
    assert "Analytics RSC" in analytics_rsc_response.text
    assert documents_response.status_code == 200
    assert "Documents App" in documents_response.text
    assert graph_response.status_code == 200
    assert "Graph Portal" in graph_response.text
    assert graph_rsc_response.status_code == 200
    assert "Graph RSC" in graph_rsc_response.text
    assert workspace_response.status_code == 200
    assert "Workspace App" in workspace_response.text
    assert asset_response.status_code == 200
    assert "console.log" in asset_response.text
    assert missing_asset_response.status_code == 404
    assert missing_api_response.status_code == 404


def test_auth_api_lists_roles_and_manages_users(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'auth-users.db'}"
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(database_url, create_schema=True)
    client = TestClient(create_app(state))

    roles_response = client.get("/auth/roles")

    assert roles_response.status_code == 200
    roles_body = roles_response.json()
    assert roles_body["mode"] == "header_transition_layer"
    assert [item["role"] for item in roles_body["items"]] == [
        "admin",
        "technician",
        "director",
        "member",
    ]
    assert roles_body["compatibility"]["it-admin"] == "admin"
    assert roles_body["compatibility"]["system-admin"] == "admin"
    permissions_by_role = {item["role"]: set(item["permissions"]) for item in roles_body["items"]}
    assert "create_report_draft" in permissions_by_role["admin"]
    assert "create_report_draft" in permissions_by_role["director"]
    assert "create_report_draft" in permissions_by_role["member"]
    assert "create_report_draft" not in permissions_by_role["technician"]
    assert "create_review_task" in permissions_by_role["admin"]
    assert "create_review_task" in permissions_by_role["director"]
    assert "create_review_task" in permissions_by_role["member"]
    assert "create_review_task" not in permissions_by_role["technician"]
    assert "create_project" in permissions_by_role["admin"]
    assert "create_project" not in permissions_by_role["director"]
    assert "create_project" not in permissions_by_role["member"]
    assert "create_project" not in permissions_by_role["technician"]

    session_response = client.get(
        "/auth/session",
        headers={
            "X-User-Id": "next-admin",
            "X-Role": "admin",
            "X-Tenant-Id": "hospital-demo",
        },
    )
    assert session_response.status_code == 200
    session_body = session_response.json()
    assert session_body["role"] == "admin"
    assert session_body["tenant_id"] == "hospital-demo"
    assert session_body["profile"]["user_key"] == "next-admin"

    system_admin_session_response = client.get(
        "/auth/session",
        headers={
            "X-User-Id": "system-admin-user",
            "X-Role": "system-admin",
            "X-Tenant-Id": "hospital-demo",
        },
    )
    assert system_admin_session_response.status_code == 200
    system_admin_session_body = system_admin_session_response.json()
    assert system_admin_session_body["role"] == "admin"
    assert system_admin_session_body["legacy_api_role"] == "it-admin"

    create_response = client.post(
        "/auth/users",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
        json={
            "display_name": "医保办主任",
            "department_key": "medical-insurance-office",
        },
    )
    assert create_response.status_code == 200
    created = create_response.json()["item"]
    assert created["user_key"].startswith(AUTH_USER_ID_PREFIX)
    assert created["display_name"] == "医保办主任"
    assert created["department_key"] == "medical-insurance-office"
    assert state.operation_logs[-1]["action"] == "auth-user-create"

    assign_response = client.post(
        f"/auth/users/{created['user_key']}/role-assignments",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
        json={"role": "director", "scope_type": "project", "scope_key": "CATALOG-LIMIT-202606"},
    )
    assert assign_response.status_code == 200
    assignment = assign_response.json()["item"]
    assert assignment["assignment_key"].startswith(AUTH_ROLE_ASSIGNMENT_ID_PREFIX)
    assert assignment["role"] == "director"
    assert assignment["scope_type"] == "project"
    assert state.operation_logs[-1]["action"] == "auth-user-role-assign"

    users_response = client.get(
        "/auth/users",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
    )
    assert users_response.status_code == 200
    users_body = users_response.json()
    assert users_body["store"]["backend"] == "SqlAlchemyAuthUserStore"
    assert users_body["departments"][0]["department_key"]
    assert any(item["user_key"] == created["user_key"] for item in users_body["items"])

    second_state = _api_state(tmp_path / "second")
    second_state.auth_user_store = SqlAlchemyAuthUserStore(database_url)
    second_client = TestClient(create_app(second_state))
    persisted_users = second_client.get(
        "/auth/users",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
    ).json()["items"]
    persisted = next(item for item in persisted_users if item["user_key"] == created["user_key"])
    assert persisted["role_assignments"][0]["role"] == "director"


def test_auth_api_rejects_member_user_management(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))

    response = client.get(
        "/auth/users",
        headers={"X-User-Id": "member-1", "X-Role": "member"},
    )
    missing_role_response = client.post(
        "/auth/users",
        json={"display_name": "未授权成员"},
    )

    assert response.status_code == 403
    assert missing_role_response.status_code == 401
    assert state.operation_logs[-1]["action"] == "authorization-denied"


def test_controlled_api_auth_middleware_enforces_user_headers_and_status(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state, enforce_controlled_api_auth=True))

    health_response = client.get("/health")
    roles_response = client.get("/auth/roles")
    anonymous_projects_response = client.get("/projects")
    assert health_response.status_code == 200
    assert roles_response.status_code == 200
    assert anonymous_projects_response.status_code == 401
    denied_payload = state.operation_logs[-1]["payload"]
    assert denied_payload["attempted_action"] == "controlled-api-auth"
    assert denied_payload["path"] == "/projects"
    assert denied_payload["permission"] == "access_controlled_api"
    assert denied_payload["tenant_id"] is None
    assert denied_payload["reason"] == "X-Tenant-Id header is required"

    missing_tenant_response = client.get(
        "/projects",
        headers={"X-User-Id": "member-1", "X-Role": "member"},
    )
    missing_tenant_payload = state.operation_logs[-1]["payload"]
    assert missing_tenant_response.status_code == 401
    assert missing_tenant_payload["path"] == "/projects"
    assert missing_tenant_payload["tenant_id"] is None
    assert missing_tenant_payload["reason"] == "X-Tenant-Id header is required"

    member_projects_response = client.get(
        "/projects",
        headers={
            "X-User-Id": "member-1",
            "X-Role": "member",
            "X-Tenant-Id": "hospital-demo",
        },
    )
    create_disabled_response = client.post(
        "/auth/users",
        headers={
            "X-User-Id": "admin-1",
            "X-Role": "admin",
            "X-Tenant-Id": "hospital-demo",
        },
        json={
            "user_key": "disabled-api-user",
            "display_name": "中间件停用用户",
            "department_key": "audit-office",
            "status": "disabled",
        },
    )
    disabled_graph_response = client.get(
        "/graph/workbench",
        headers={
            "X-User-Id": "disabled-api-user",
            "X-Role": "admin",
            "X-Tenant-Id": "hospital-demo",
        },
    )

    assert member_projects_response.status_code == 200
    assert create_disabled_response.status_code == 200
    assert disabled_graph_response.status_code == 403
    disabled_denied_payload = state.operation_logs[-1]["payload"]
    assert disabled_denied_payload["attempted_action"] == "controlled-api-auth"
    assert disabled_denied_payload["path"] == "/graph/workbench"
    assert disabled_denied_payload["tenant_id"] == "hospital-demo"
    assert disabled_denied_payload["reason"] == "auth user status is disabled"


def test_permission_resolver_prefers_persisted_global_role(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    state.agent_store = SqlAlchemyAgentStore(
        f"sqlite:///{tmp_path / 'agents.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))

    create_user_response = client.post(
        "/auth/users",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
        json={
            "user_key": "persistent-technician",
            "display_name": "持久化技术人员",
            "department_key": "it-department",
        },
    )
    assert create_user_response.status_code == 200

    assign_response = client.post(
        "/auth/users/persistent-technician/role-assignments",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
        json={"role": "technician", "scope_type": "global"},
    )
    assert assign_response.status_code == 200

    session_response = client.get(
        "/auth/session",
        headers={"X-User-Id": "persistent-technician", "X-Role": "member"},
    )
    assert session_response.status_code == 200
    session_body = session_response.json()
    assert session_body["role"] == "technician"
    assert session_body["auth_source"] == "persistent_role"
    assert session_body["profile_status"] == "active"

    agent_response = client.post(
        "/agents",
        headers={
            "X-User-Id": "persistent-technician",
            "X-Role": "member",
            "X-Project-Name": PROJECT_NAME_HEADER,
        },
        json={
            "name": "持久化权限助手",
            "category": "业务类",
            "topic": "医保基金使用合规",
            "prompt": "仅基于授权角色进行配置。",
        },
    )
    assert agent_response.status_code == 200
    assert state.operation_logs[-1]["action"] == "agent-create"
    assert state.operation_logs[-1]["payload"]["role"] == "technician"

    project_member_response = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers={"X-User-Id": "persistent-technician", "X-Role": "admin"},
        json={
            "user_identifier": "persistent-denied-member",
            "name": "持久化越权用户",
            "role": "审计员",
            "department": "医保办",
        },
    )
    assert project_member_response.status_code == 403
    denied_payload = state.operation_logs[-1]["payload"]
    assert state.operation_logs[-1]["action"] == "authorization-denied"
    assert denied_payload["effective_role"] == "technician"
    assert denied_payload["auth_source"] == "persistent_role"
    assert denied_payload["permission"] == "manage_project_members"


def test_permission_resolver_uses_project_scoped_role_for_matching_project(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    state.project_member_store = SqlAlchemyProjectMemberStore(
        f"sqlite:///{tmp_path / 'project-members.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    admin_headers = {"X-User-Id": "admin-1", "X-Role": "admin"}

    create_user_response = client.post(
        "/auth/users",
        headers=admin_headers,
        json={
            "user_key": "project-director",
            "display_name": "项目主任",
            "department_key": "audit-office",
        },
    )
    assert create_user_response.status_code == 200
    assign_response = client.post(
        "/auth/users/project-director/role-assignments",
        headers=admin_headers,
        json={
            "role": "admin",
            "scope_type": "project",
            "scope_key": "SELF-CHECK-FUND-20260607",
        },
    )
    assert assign_response.status_code == 200

    scoped_session_response = client.get(
        "/auth/session",
        headers={
            "X-User-Id": "project-director",
            "X-Role": "member",
            "X-Project-Key": "SELF-CHECK-FUND-20260607",
        },
    )
    other_project_session_response = client.get(
        "/auth/session",
        headers={
            "X-User-Id": "project-director",
            "X-Role": "admin",
            "X-Project-Key": "CATALOG-LIMIT-202606",
        },
    )
    allowed_member_response = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers={"X-User-Id": "project-director", "X-Role": "member"},
        json={
            "user_identifier": "project-authorized-member",
            "name": "项目授权成员",
            "role": "审计员",
            "department": "医保办",
        },
    )
    blocked_member_response = client.post(
        "/projects/CATALOG-LIMIT-202606/members",
        headers={"X-User-Id": "project-director", "X-Role": "admin"},
        json={
            "user_identifier": "cross-project-denied-member",
            "name": "跨项目越权成员",
            "role": "审计员",
            "department": "医保办",
        },
    )

    assert scoped_session_response.status_code == 200
    scoped_session = scoped_session_response.json()
    assert scoped_session["role"] == "admin"
    assert scoped_session["auth_source"] == "persistent_project_role"
    assert scoped_session["auth_scope_type"] == "project"
    assert scoped_session["auth_scope_key"] == "SELF-CHECK-FUND-20260607"
    assert other_project_session_response.status_code == 200
    other_project_session = other_project_session_response.json()
    assert other_project_session["role"] == "member"
    assert other_project_session["auth_source"] == "persistent_profile_without_project_role"
    assert allowed_member_response.status_code == 200
    assert allowed_member_response.json()["item"]["created_by"] == "project-director"
    assert blocked_member_response.status_code == 403
    denied_payload = state.operation_logs[-1]["payload"]
    assert denied_payload["attempted_action"] == "project-member-create"
    assert denied_payload["effective_role"] == "member"
    assert denied_payload["auth_source"] == "persistent_profile_without_project_role"
    assert denied_payload["auth_scope_type"] == "project"
    assert denied_payload["auth_scope_key"] == "CATALOG-LIMIT-202606"


def test_permission_resolver_denies_disabled_persisted_user(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))

    create_user_response = client.post(
        "/auth/users",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
        json={
            "user_key": "disabled-admin",
            "display_name": "停用管理员",
            "department_key": "it-department",
            "status": "disabled",
        },
    )
    assert create_user_response.status_code == 200

    agent_response = client.post(
        "/agents",
        headers={"X-User-Id": "disabled-admin", "X-Role": "admin"},
        json={
            "name": "停用用户不应保存",
            "category": "业务类",
            "topic": "医保基金使用合规",
            "prompt": "停用用户不能写入。",
        },
    )

    assert agent_response.status_code == 403
    assert state.operation_logs[-1]["action"] == "authorization-denied"
    assert state.operation_logs[-1]["payload"]["user_identifier"] == "disabled-admin"
    assert state.operation_logs[-1]["payload"]["reason"] == "auth user status is disabled"


def test_auth_api_updates_user_status_and_role_assignment(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    admin_headers = {"X-User-Id": "admin-1", "X-Role": "admin"}

    create_user_response = client.post(
        "/auth/users",
        headers=admin_headers,
        json={
            "user_key": "managed-director",
            "display_name": "待管理主任",
            "department_key": "audit-office",
        },
    )
    assert create_user_response.status_code == 200
    assign_response = client.post(
        "/auth/users/managed-director/role-assignments",
        headers=admin_headers,
        json={"role": "director", "scope_type": "global"},
    )
    assert assign_response.status_code == 200
    assignment_key = assign_response.json()["item"]["assignment_key"]

    disabled_response = client.patch(
        "/auth/users/managed-director",
        headers=admin_headers,
        json={"status": "disabled"},
    )
    assert disabled_response.status_code == 200
    assert disabled_response.json()["item"]["status"] == "disabled"
    assert state.operation_logs[-1]["action"] == "auth-user-update"

    disabled_session_response = client.get(
        "/auth/session",
        headers={"X-User-Id": "managed-director", "X-Role": "director"},
    )
    assert disabled_session_response.status_code == 403
    assert disabled_session_response.json()["detail"] == "auth user status is disabled"

    active_response = client.patch(
        "/auth/users/managed-director",
        headers=admin_headers,
        json={"status": "active", "display_name": "已恢复主任"},
    )
    assert active_response.status_code == 200
    assert active_response.json()["item"]["status"] == "active"
    assert active_response.json()["item"]["display_name"] == "已恢复主任"

    restored_session_response = client.get(
        "/auth/session",
        headers={"X-User-Id": "managed-director", "X-Role": "member"},
    )
    assert restored_session_response.status_code == 200
    assert restored_session_response.json()["role"] == "director"

    revoke_response = client.patch(
        f"/auth/users/managed-director/role-assignments/{assignment_key}",
        headers=admin_headers,
        json={"status": "revoked"},
    )
    assert revoke_response.status_code == 200
    assert revoke_response.json()["item"]["status"] == "revoked"
    assert state.operation_logs[-1]["action"] == "auth-user-role-assignment-update"

    downgraded_session_response = client.get(
        "/auth/session",
        headers={"X-User-Id": "managed-director", "X-Role": "admin"},
    )
    assert downgraded_session_response.status_code == 200
    downgraded_session = downgraded_session_response.json()
    assert downgraded_session["role"] == "member"
    assert downgraded_session["auth_source"] == "persistent_profile_without_global_role"

    restore_assignment_response = client.patch(
        f"/auth/users/managed-director/role-assignments/{assignment_key}",
        headers=admin_headers,
        json={"status": "active"},
    )
    assert restore_assignment_response.status_code == 200
    assert restore_assignment_response.json()["item"]["status"] == "active"
    restored_role_response = client.get(
        "/auth/session",
        headers={"X-User-Id": "managed-director", "X-Role": "member"},
    )
    assert restored_role_response.status_code == 200
    assert restored_role_response.json()["role"] == "director"


def test_agents_api_lists_defaults_and_persists_created_agent(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'agents.db'}"
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(database_url, create_schema=True)
    client = TestClient(create_app(state))

    list_response = client.get("/agents")

    assert list_response.status_code == 200
    list_body = list_response.json()
    assert list_body["store"]["backend"] == "SqlAlchemyAgentStore"
    assert list_body["categories"] == ["业务类", "效率类", "研究类"]
    assert [item["id"] for item in list_body["items"][:4]] == [
        "agent-contract-audit-v2",
        "agent-citation-check",
        "agent-duplicate-charge",
        "agent-report-draft",
    ]

    create_response = client.post(
        "/agents",
        headers={
            "X-User-Id": "director-1",
            "X-Role": "director",
            "X-Project-Name": urllib.parse.quote("医保目录限制条件核验"),
        },
        json={
            "name": "目录限制核验助手",
            "category": "业务类",
            "topic": "医保目录限制条件核验",
            "prompt": "仅基于目录限制字段和引用依据输出待补证问题。",
            "knowledge_base": "医保目录库",
            "project_name": "医保目录限制条件核验",
        },
    )

    assert create_response.status_code == 200
    created = create_response.json()["item"]
    assert created["id"].startswith(AGENT_ID_PREFIX)
    assert created["created_by"] == "director-1"
    assert created["source"] == "custom"
    assert created["status"] == "active"
    assert created["prompt_version"] == 1
    assert created["prompt_version_key"] == f"{created['id']}@v1"
    assert created["visibility_scope"] == "project"
    assert created["allowed_roles"] == ["admin", "technician", "director", "member"]
    assert state.operation_logs[-1]["action"] == "agent-create"
    assert state.operation_logs[-1]["payload"]["role"] == "director"

    second_state = _api_state(tmp_path / "second")
    second_state.agent_store = SqlAlchemyAgentStore(database_url)
    second_client = TestClient(create_app(second_state))
    persisted_items = second_client.get(
        "/agents",
        headers={"X-Project-Name": urllib.parse.quote("医保目录限制条件核验")},
    ).json()["items"]

    assert persisted_items[0]["id"] == created["id"]
    assert persisted_items[0]["name"] == "目录限制核验助手"
    assert any(item["id"] == "agent-citation-check" for item in persisted_items)


def test_agents_api_preserves_catalog_source_row_metadata(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'agents-catalog-row.db'}"
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(database_url, create_schema=True)
    client = TestClient(create_app(state))
    catalog_metadata = {
        "catalog_source": "audit-agent-prompts-0613",
        "catalog_row_id": "audit-agent-prompts-0613-042",
        "source_key": "audit-agent-prompts-0613-042",
        "source_row_index": 42,
        "legacy_source_key": "财务收支审计|会议费用核验",
        "source_category": "财务收支审计",
        "source_title": "会议费用核验",
        "source_scene": "用于会议费报销材料审计",
        "source_file": "提示词分类0613.zip",
        "display_name": "会议费用核验",
        "avatar_seed": "财务收支审计-会议费用核验-042",
    }

    create_response = client.post(
        "/agents",
        headers={
            "X-User-Id": "director-1",
            "X-Role": "director",
            "X-Project-Name": PROJECT_NAME_HEADER,
        },
        json={
            "name": "会议费用核验",
            "category": "业务类",
            "topic": "会议费报销合规",
            "prompt": "仅基于会议通知、签到表、报销凭证和制度依据输出核验意见。",
            "knowledge_base": "系统医保审计知识库",
            "project_name": "医保基金使用合规专项自查",
            "visibility_scope": "project",
            "allowed_roles": ["admin", "technician", "director", "member"],
            "metadata": catalog_metadata,
        },
    )

    assert create_response.status_code == 200
    created = create_response.json()["item"]
    assert created["id"].startswith(AGENT_ID_PREFIX)
    assert created["metadata"] | catalog_metadata == created["metadata"]
    assert created["metadata"]["catalog_row_id"] == "audit-agent-prompts-0613-042"
    assert created["metadata"]["source_row_index"] == 42
    assert created["metadata"]["prompt_version_key"] == f"{created['id']}@v1"

    second_state = _api_state(tmp_path / "second-catalog-row")
    second_state.agent_store = SqlAlchemyAgentStore(database_url)
    second_client = TestClient(create_app(second_state))
    persisted_response = second_client.get(
        "/agents",
        headers={"X-Project-Name": PROJECT_NAME_HEADER},
    )

    assert persisted_response.status_code == 200
    persisted = next(
        item for item in persisted_response.json()["items"] if item["id"] == created["id"]
    )
    assert persisted["metadata"] | catalog_metadata == persisted["metadata"]
    assert persisted["metadata"]["catalog_row_id"] == "audit-agent-prompts-0613-042"
    assert persisted["metadata"]["source_row_index"] == 42


def test_agents_api_fallback_defaults_include_prompt_version_metadata(tmp_path: Path) -> None:
    class UnavailableAgentStore:
        def list_agents(self, *, include_inactive: bool = False) -> list[dict[str, object]]:
            raise SQLAlchemyError("agent store unavailable")

    state = _api_state(tmp_path)
    state.agent_store = UnavailableAgentStore()  # type: ignore[assignment]
    client = TestClient(create_app(state))

    response = client.get("/agents")

    assert response.status_code == 200
    body = response.json()
    assert body["store"] == {"ready": False, "backend": "unavailable"}
    citation_agent = next(item for item in body["items"] if item["id"] == "agent-citation-check")
    assert citation_agent["prompt_versions"][0]["version"] == 1
    assert citation_agent["prompt_versions"][0]["is_active"] is True


def test_agents_api_tracks_prompt_versions_lifecycle_and_history(
    tmp_path: Path,
) -> None:
    database_url = f"sqlite:///{tmp_path / 'agents-governance.db'}"
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(database_url, create_schema=True)
    client = TestClient(create_app(state))
    headers = {"X-User-Id": "director-1", "X-Role": "director"}

    create_response = client.post(
        "/agents",
        headers=headers,
        json={
            "name": "目录限制核验助手",
            "category": "业务类",
            "topic": "医保目录限制条件核验",
            "prompt": "仅基于目录限制字段输出待补证问题。",
            "knowledge_base": "医保目录库",
            "project_name": "医保目录限制条件核验",
            "visibility_scope": "system",
            "allowed_roles": ["director", "member"],
        },
    )
    agent = create_response.json()["item"]
    agent_id = str(agent["id"])

    update_response = client.post(
        f"/agents/{agent_id}/prompt-versions",
        headers=headers,
        json={
            "prompt": "仅基于目录限制字段和原文引用输出待补证问题。",
            "change_summary": "补充原文引用约束。",
            "review_note": "待主任复核原文引用边界。",
        },
    )
    updated = update_response.json()["item"]
    changes_response = client.post(
        f"/agents/{agent_id}/prompt-versions/review",
        headers=headers,
        json={
            "version": 2,
            "review_status": "changes-requested",
            "review_note": "需补充原文引用边界说明。",
        },
    )
    changes_requested = changes_response.json()["item"]
    review_response = client.post(
        f"/agents/{agent_id}/prompt-versions/review",
        headers=headers,
        json={
            "version": 2,
            "review_status": "approved",
            "review_note": "主任已复核提示词引用边界。",
        },
    )
    reviewed = review_response.json()["item"]
    rollback_response = client.post(
        f"/agents/{agent_id}/prompt-versions/rollback",
        headers=headers,
        json={"version": 1},
    )
    rolled_back = rollback_response.json()["item"]
    versions_response = client.get(f"/agents/{agent_id}/prompt-versions")
    invocation_response = client.post(
        f"/agents/{agent_id}/invocations",
        headers={"X-User-Id": "member-1", "X-Role": "member"},
        json={
            "invocation_source": "agent-workspace",
            "question": "目录限制核验试用",
            "conversation_ref": "local-chat-draft",
        },
    )
    invocation = invocation_response.json()["item"]
    feedback_response = client.post(
        f"/agents/{agent_id}/feedback",
        headers={"X-User-Id": "member-1", "X-Role": "member"},
        json={
            "invocation_id": invocation["id"],
            "rating": "effective",
            "comment": "引用约束清晰，可继续使用。",
        },
    )
    invocations_response = client.get(f"/agents/{agent_id}/invocations", headers=headers)
    feedback_list_response = client.get(f"/agents/{agent_id}/feedback", headers=headers)
    inactive_response = client.post(
        f"/agents/{agent_id}/lifecycle",
        headers=headers,
        json={"status": "inactive", "reason": "提示词待复核"},
    )
    archived_response = client.post(
        f"/agents/{agent_id}/lifecycle",
        headers=headers,
        json={"status": "archived", "reason": "软归档，不物理删除"},
    )
    list_response = client.get("/agents")
    detail_response = client.get(f"/agents/{agent_id}")

    assert create_response.status_code == 200
    assert agent["visibility_scope"] == "system"
    assert agent["allowed_roles"] == ["director", "member"]
    assert update_response.status_code == 200
    assert updated["prompt_version"] == 1
    assert updated["prompt"] == "仅基于目录限制字段输出待补证问题。"
    assert updated["prompt_versions"][0]["is_active"] is True
    assert updated["prompt_versions"][1]["review_status"] == "pending-review"
    assert updated["prompt_versions"][1]["is_active"] is False
    assert updated["prompt_versions"][1]["review_note"] == "待主任复核原文引用边界。"
    assert changes_response.status_code == 200
    assert changes_requested["prompt_version"] == 1
    assert changes_requested["prompt"] == "仅基于目录限制字段输出待补证问题。"
    assert changes_requested["prompt_versions"][1]["review_status"] == "changes-requested"
    assert changes_requested["prompt_versions"][1]["is_active"] is False
    assert review_response.status_code == 200
    assert reviewed["prompt_version"] == 2
    assert reviewed["prompt"] == "仅基于目录限制字段和原文引用输出待补证问题。"
    assert reviewed["prompt_versions"][0]["is_active"] is False
    assert reviewed["prompt_versions"][1]["is_active"] is True
    assert reviewed["prompt_versions"][1]["review_status"] == "approved"
    assert reviewed["prompt_versions"][1]["review_note"] == "主任已复核提示词引用边界。"
    assert reviewed["prompt_versions"][1]["reviewed_by"] == "director-1"
    assert "agent-prompt-version-review" in [entry["action"] for entry in state.operation_logs]
    assert rollback_response.status_code == 200
    assert rolled_back["prompt_version"] == 3
    assert rolled_back["prompt"] == "仅基于目录限制字段输出待补证问题。"
    assert rolled_back["prompt_versions"][2]["review_status"] == "approved"
    assert rolled_back["prompt_versions"][2]["is_active"] is True
    assert [item["version"] for item in rolled_back["prompt_versions"]] == [1, 2, 3]
    assert versions_response.status_code == 200
    assert [item["version"] for item in versions_response.json()["items"]] == [1, 2, 3]
    assert versions_response.json()["items"][2]["is_active"] is True
    assert invocation_response.status_code == 200
    assert invocation["agent_key"] == agent_id
    assert invocation["prompt_version"] == 3
    assert invocation["question"] == "目录限制核验试用"
    assert feedback_response.status_code == 200
    assert feedback_response.json()["item"]["invocation_id"] == invocation["id"]
    assert feedback_response.json()["item"]["rating"] == "effective"
    assert feedback_response.json()["summary"]["effective"] == 1
    assert feedback_response.json()["summary"]["latest_rating"] == "effective"
    assert invocations_response.json()["items"][0]["id"] == invocation["id"]
    assert feedback_list_response.json()["items"][0]["comment"] == "引用约束清晰，可继续使用。"
    assert feedback_list_response.json()["summary"] == {
        "total": 1,
        "effective": 1,
        "needs_review": 0,
        "unsafe": 0,
        "latest_rating": "effective",
    }
    assert inactive_response.status_code == 200
    assert inactive_response.json()["item"]["status"] == "inactive"
    assert archived_response.status_code == 200
    assert archived_response.json()["item"]["status"] == "archived"
    assert agent_id not in {item["id"] for item in list_response.json()["items"]}
    assert detail_response.status_code == 200
    assert detail_response.json()["item"]["id"] == agent_id
    assert detail_response.json()["item"]["status"] == "archived"
    assert detail_response.json()["item"]["metadata"]["lifecycle_reason"] == "软归档，不物理删除"
    assert [entry["action"] for entry in state.operation_logs[-9:]] == [
        "agent-prompt-versions-view",
        "agent-invocation-create",
        "agent-feedback-create",
        "agent-invocations-view",
        "agent-feedback-view",
        "agent-lifecycle-update",
        "agent-lifecycle-update",
        "agents-list",
        "agent-detail-view",
    ]


def test_agents_api_restricts_prompt_activation_to_admin_and_director(
    tmp_path: Path,
) -> None:
    database_url = f"sqlite:///{tmp_path / 'agents-prompt-activation-role.db'}"
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(database_url, create_schema=True)
    client = TestClient(create_app(state))
    technician_headers = {
        "X-User-Id": "technician-1",
        "X-Role": "technician",
        "X-Project-Name": PROJECT_NAME_HEADER,
    }
    director_headers = {
        "X-User-Id": "director-1",
        "X-Role": "director",
        "X-Project-Name": PROJECT_NAME_HEADER,
    }

    create_response = client.post(
        "/agents",
        headers=technician_headers,
        json={
            "name": "技术配置候选助手",
            "category": "业务类",
            "topic": "医保基金使用合规",
            "prompt": "输出待补证问题。",
            "knowledge_base": "项目默认知识库",
            "project_name": "医保基金使用合规专项自查",
            "visibility_scope": "project",
        },
    )
    agent_id = str(create_response.json()["item"]["id"])
    update_response = client.post(
        f"/agents/{agent_id}/prompt-versions",
        headers=technician_headers,
        json={
            "prompt": "输出待补证问题，并标注引用依据。",
            "change_summary": "技术人员补充引用依据约束。",
            "review_note": "待主任复核。",
        },
    )
    technician_review_response = client.post(
        f"/agents/{agent_id}/prompt-versions/review",
        headers=technician_headers,
        json={
            "version": 2,
            "review_status": "approved",
            "review_note": "技术人员尝试激活。",
        },
    )
    technician_rollback_response = client.post(
        f"/agents/{agent_id}/prompt-versions/rollback",
        headers=technician_headers,
        json={"version": 1},
    )
    director_review_response = client.post(
        f"/agents/{agent_id}/prompt-versions/review",
        headers=director_headers,
        json={
            "version": 2,
            "review_status": "approved",
            "review_note": "主任复核通过。",
        },
    )

    assert create_response.status_code == 200
    assert update_response.status_code == 200
    assert update_response.json()["item"]["prompt_version"] == 1
    assert technician_review_response.status_code == 403
    assert (
        technician_review_response.json()["detail"]
        == "agent prompt activation requires admin or director role"
    )
    assert technician_rollback_response.status_code == 403
    assert director_review_response.status_code == 200
    reviewed = director_review_response.json()["item"]
    assert reviewed["prompt_version"] == 2
    assert reviewed["prompt_versions"][1]["is_active"] is True
    denied_actions = [
        entry["payload"]["attempted_action"]
        for entry in state.operation_logs
        if entry["action"] == "authorization-denied"
    ]
    assert "agent-prompt-version-review" in denied_actions
    assert "agent-prompt-version-rollback" in denied_actions


def test_agents_api_filters_project_scope_and_blocks_cross_project_invocation(
    tmp_path: Path,
) -> None:
    database_url = f"sqlite:///{tmp_path / 'agents-project-scope.db'}"
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(database_url, create_schema=True)
    client = TestClient(create_app(state))
    project_a_headers = {
        "X-User-Id": "admin-1",
        "X-Role": "admin",
        "X-Project-Name": PROJECT_NAME_HEADER,
    }
    project_b_headers = {
        "X-User-Id": "admin-2",
        "X-Role": "admin",
        "X-Project-Name": "other-project",
    }

    create_response = client.post(
        "/agents",
        headers=project_a_headers,
        json={
            "name": "项目内目录限制核验助手",
            "category": "业务类",
            "topic": "医保目录限制条件核验",
            "prompt": "只在当前项目空间内用于目录限制条件核验。",
            "knowledge_base": "项目默认知识库",
            "project_name": "医保基金使用合规专项自查",
            "visibility_scope": "project",
        },
    )
    created = create_response.json()["item"]
    agent_id = str(created["id"])

    project_a_list = client.get("/agents", headers=project_a_headers)
    project_b_list = client.get("/agents", headers=project_b_headers)
    missing_scope_list = client.get("/agents")
    blocked_detail = client.get(f"/agents/{agent_id}", headers=project_b_headers)
    missing_scope_detail = client.get(f"/agents/{agent_id}")
    blocked_invocation = client.post(
        f"/agents/{agent_id}/invocations",
        headers={**project_b_headers, "X-User-Id": "member-2", "X-Role": "member"},
        json={"invocation_source": "agent-workspace", "question": "跨项目调用"},
    )
    allowed_invocation = client.post(
        f"/agents/{agent_id}/invocations",
        headers={**project_a_headers, "X-User-Id": "member-1", "X-Role": "member"},
        json={"invocation_source": "agent-workspace", "question": "本项目调用"},
    )
    missing_scope_invocation = client.post(
        f"/agents/{agent_id}/invocations",
        headers={"X-User-Id": "member-3", "X-Role": "member"},
        json={"invocation_source": "agent-workspace", "question": "缺少项目范围调用"},
    )
    missing_scope_create = client.post(
        "/agents",
        headers={"X-User-Id": "admin-3", "X-Role": "admin"},
        json={
            "name": "缺少项目范围的助手",
            "category": "业务类",
            "topic": "医保目录限制条件核验",
            "prompt": "不得在缺少当前项目时创建。",
            "knowledge_base": "项目默认知识库",
            "project_name": "医保基金使用合规专项自查",
            "visibility_scope": "project",
        },
    )

    assert create_response.status_code == 200
    assert agent_id in {item["id"] for item in project_a_list.json()["items"]}
    assert agent_id not in {item["id"] for item in project_b_list.json()["items"]}
    assert agent_id not in {item["id"] for item in missing_scope_list.json()["items"]}
    assert "agent-citation-check" in {item["id"] for item in project_b_list.json()["items"]}
    assert blocked_detail.status_code == 403
    assert blocked_detail.json()["detail"] == "agent project scope does not match current project"
    assert missing_scope_detail.status_code == 403
    assert missing_scope_detail.json()["detail"] == "agent project scope requires current project"
    assert blocked_invocation.status_code == 403
    assert missing_scope_invocation.status_code == 403
    assert missing_scope_create.status_code == 403
    assert allowed_invocation.status_code == 200
    assert allowed_invocation.json()["item"]["question"] == "本项目调用"
    assert any(entry["action"] == "agent-project-scope-denied" for entry in state.operation_logs)


def test_agents_api_rejects_unknown_category(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(
        f"sqlite:///{tmp_path / 'agents.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/agents",
        json={
            "name": "未知分类助手",
            "category": "其他",
            "topic": "医保基金使用合规",
            "prompt": "输出审计问题。",
        },
    )

    assert response.status_code == 422


def test_agents_api_enforces_manage_agent_permission(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(
        f"sqlite:///{tmp_path / 'agents.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    payload = {
        "name": "成员自建系统智能体",
        "category": "业务类",
        "topic": "医保基金使用合规",
        "prompt": "输出审计问题。",
        "visibility_scope": "system",
    }

    unauthenticated_response = client.post("/agents", json=payload)
    member_response = client.post(
        "/agents",
        headers={"X-User-Id": "member-1", "X-Role": "member"},
        json=payload,
    )

    assert unauthenticated_response.status_code == 401
    assert member_response.status_code == 403
    assert state.operation_logs[-1]["action"] == "authorization-denied"
    assert state.operation_logs[-1]["payload"]["attempted_action"] == "agent-create"
    assert state.operation_logs[-1]["payload"]["permission"] == "manage_agents"


def test_projects_api_filters_visibility_and_publishes_canonical_statuses(
    tmp_path: Path,
) -> None:
    database_url = f"sqlite:///{tmp_path / 'project-members.db'}"
    state = _api_state(tmp_path)
    state.project_member_store = SqlAlchemyProjectMemberStore(database_url, create_schema=True)
    client = TestClient(create_app(state))

    anonymous_response = client.get("/projects")
    role_only_response = client.get("/projects", headers={"X-Role": "admin"})
    admin_response = client.get(
        "/projects",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
    )
    creator_response = client.get(
        "/projects",
        headers={"X-User-Id": "expert-catalog"},
    )
    self_creator_response = client.get(
        "/projects",
        headers={"X-User-Id": "next-director", "X-Role": "director"},
    )
    active_member_response = client.get(
        "/projects",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )
    pending_member_response = client.get(
        "/projects",
        headers={"X-User-Id": "next-technician", "X-Role": "technician"},
    )
    unrelated_response = client.get(
        "/projects",
        headers={"X-User-Id": "unrelated-member", "X-Role": "member"},
    )

    assert anonymous_response.status_code == 200
    assert anonymous_response.json()["items"] == []
    assert role_only_response.status_code == 200
    assert role_only_response.json()["items"] == []
    assert admin_response.status_code == 200
    projects_body = admin_response.json()
    assert projects_body["store"]["backend"] == "SqlAlchemyProjectMemberStore"
    assert projects_body["roles"] == ["项目负责人", "审计员", "业务专家", "信息科", "只读观察员"]
    assert projects_body["statuses"] == ["在项目中", "待确认"]
    assert projects_body["project_statuses"] == ["待开始", "进行中", "已完成", "已归档"]
    project_statuses = {item["status"] for item in projects_body["items"]}
    assert project_statuses <= set(projects_body["project_statuses"])
    assert "待启动" not in project_statuses
    assert len(projects_body["items"]) == 4
    assert projects_body["items"][0]["id"] == "SELF-CHECK-FUND-20260607"
    assert projects_body["items"][0]["member_count"] == 3
    assert {item["id"]: item["creator_user_identifier"] for item in projects_body["items"]} == {
        "SELF-CHECK-FUND-20260607": "next-director",
        "CATALOG-LIMIT-202606": "expert-catalog",
        "OUTPATIENT-DOSE-202606": "auditor-outpatient-dose",
        "KB-GOVERNANCE-202606": "it-kb-governance",
    }
    assert (
        next(item for item in projects_body["items"] if item["id"] == "OUTPATIENT-DOSE-202606")[
            "status"
        ]
        == "进行中"
    )
    assert [item["id"] for item in creator_response.json()["items"]] == ["CATALOG-LIMIT-202606"]
    assert [item["id"] for item in self_creator_response.json()["items"]] == [
        "SELF-CHECK-FUND-20260607"
    ]
    assert [item["id"] for item in active_member_response.json()["items"]] == [
        "SELF-CHECK-FUND-20260607"
    ]
    assert pending_member_response.json()["items"] == []
    assert unrelated_response.json()["items"] == []
    assert state.operation_logs[-1]["payload"]["actor"] == "unrelated-member"
    assert state.operation_logs[-1]["payload"]["visible_project_count"] == 0


def test_default_project_creators_are_active_members() -> None:
    for project in DEFAULT_PROJECT_PAYLOADS:
        project_key = str(project["id"])
        creator_user_identifier = project["creator_user_identifier"]
        member_identifiers = [
            member["user_identifier"] for member in DEFAULT_PROJECT_MEMBERS_BY_PROJECT[project_key]
        ]
        assert len(member_identifiers) == len(set(member_identifiers))
        assert any(
            member["user_identifier"] == creator_user_identifier and member["status"] == "在项目中"
            for member in DEFAULT_PROJECT_MEMBERS_BY_PROJECT[project_key]
        )


def test_projects_api_protects_project_detail_members_and_dashboard(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.project_member_store = SqlAlchemyProjectMemberStore(
        f"sqlite:///{tmp_path / 'project-members.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    routes = (
        "/projects/SELF-CHECK-FUND-20260607",
        "/projects/SELF-CHECK-FUND-20260607/members",
        "/projects/SELF-CHECK-FUND-20260607/dashboard",
    )

    for route in routes:
        creator_response = client.get(
            route,
            headers={"X-User-Id": "next-director", "X-Role": "director"},
        )
        active_member_response = client.get(
            route,
            headers={"X-User-Id": "next-member", "X-Role": "member"},
        )
        pending_member_response = client.get(
            route,
            headers={"X-User-Id": "next-technician", "X-Role": "technician"},
        )
        unrelated_response = client.get(
            route,
            headers={"X-User-Id": "unrelated-member", "X-Role": "member"},
        )
        anonymous_response = client.get(route)
        role_only_response = client.get(route, headers={"X-Role": "admin"})

        assert creator_response.status_code == 200
        assert active_member_response.status_code == 200
        assert pending_member_response.status_code == 404
        assert unrelated_response.status_code == 404
        assert anonymous_response.status_code == 404
        assert role_only_response.status_code == 404

    members = client.get(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    ).json()["items"]
    assert members[0]["user_identifier"] == "next-member"
    assert state.operation_logs[-1]["payload"]["actor"] == "next-member"

    admin_headers = {"X-User-Id": "admin-1", "X-Role": "admin"}
    for route in (
        "/projects/CATALOG-LIMIT-202606",
        "/projects/CATALOG-LIMIT-202606/members",
        "/projects/CATALOG-LIMIT-202606/dashboard",
    ):
        assert client.get(route, headers=admin_headers).status_code == 200


def test_projects_api_uses_project_scoped_auth_for_detail(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    state.project_member_store = InMemoryProjectMemberStore()
    client = TestClient(create_app(state))
    admin_headers = {"X-User-Id": "admin-1", "X-Role": "admin"}
    assert (
        client.post(
            "/auth/users",
            headers=admin_headers,
            json={
                "user_key": "scoped-project-admin",
                "display_name": "项目级管理员",
                "department_key": "audit-office",
            },
        ).status_code
        == 200
    )
    assert (
        client.post(
            "/auth/users/scoped-project-admin/role-assignments",
            headers=admin_headers,
            json={
                "role": "admin",
                "scope_type": "project",
                "scope_key": "SELF-CHECK-FUND-20260607",
            },
        ).status_code
        == 200
    )

    headers = {"X-User-Id": "scoped-project-admin", "X-Role": "member"}
    list_response = client.get("/projects", headers=headers)
    detail_response = client.get(
        "/projects/SELF-CHECK-FUND-20260607",
        headers=headers,
    )
    other_detail_response = client.get(
        "/projects/CATALOG-LIMIT-202606",
        headers=headers,
    )

    assert list_response.status_code == 200
    assert list_response.json()["items"] == []
    assert detail_response.status_code == 200
    assert other_detail_response.status_code == 404
    assert state.operation_logs[-1]["payload"]["actor"] == "scoped-project-admin"
    assert state.operation_logs[-1]["payload"]["actor_role"] == "admin"


def test_projects_api_lists_defaults_and_persists_created_member(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'project-members.db'}"
    state = _api_state(tmp_path)
    state.project_member_store = SqlAlchemyProjectMemberStore(database_url, create_schema=True)
    client = TestClient(create_app(state))
    admin_headers = {"X-User-Id": "admin-1", "X-Role": "admin"}

    project_response = client.get(
        "/projects/SELF-CHECK-FUND-20260607",
        headers=admin_headers,
    )

    assert project_response.status_code == 200
    project_body = project_response.json()
    assert project_body["item"]["id"] == "SELF-CHECK-FUND-20260607"
    assert project_body["item"]["member_count"] == 3
    assert project_body["store"]["backend"] == "SqlAlchemyProjectMemberStore"
    assert project_body["production_side_effect"] == "none"

    dashboard_response = client.get(
        "/projects/SELF-CHECK-FUND-20260607/dashboard",
        headers=admin_headers,
    )

    assert dashboard_response.status_code == 200
    dashboard_body = dashboard_response.json()
    assert dashboard_body["format"] == "project-dashboard-v1"
    assert dashboard_body["project"]["id"] == "SELF-CHECK-FUND-20260607"
    assert dashboard_body["metrics"][0]["key"] == "open_findings"
    assert dashboard_body["queue"][0]["id"] == "QUEUE-BACKEND-001"
    assert dashboard_body["store"]["backend"]["project_members"] == "SqlAlchemyProjectMemberStore"
    assert dashboard_body["store"]["backend"]["audit_findings"] == "unavailable"
    assert dashboard_body["production_side_effect"] == "none"

    members_response = client.get(
        "/projects/CATALOG-LIMIT-202606/members",
        headers=admin_headers,
    )

    assert members_response.status_code == 200
    members_body = members_response.json()
    assert members_body["project_key"] == "CATALOG-LIMIT-202606"
    assert len(members_body["items"]) == 4
    assert members_body["items"][0]["source"] == "system-default"

    create_response = client.post(
        "/projects/CATALOG-LIMIT-202606/members",
        headers=admin_headers,
        json={
            "user_identifier": "  custom-catalog-auditor  ",
            "name": "赵审计",
            "role": "审计员",
            "department": "医保办",
            "status": "在项目中",
            "metadata": {
                "ticket": "MEMBER-42",
                "user_identifier": "spoofed-catalog-member",
            },
        },
    )

    assert create_response.status_code == 200
    created = create_response.json()["item"]
    assert created["id"].startswith(PROJECT_MEMBER_ID_PREFIX)
    assert created["project_key"] == "CATALOG-LIMIT-202606"
    assert created["status"] == "在项目中"
    assert created["user_identifier"] == "custom-catalog-auditor"
    assert created["metadata"]["ticket"] == "MEMBER-42"
    assert created["metadata"]["user_identifier"] == "custom-catalog-auditor"
    assert created["created_by"] == "admin-1"
    assert state.operation_logs[-1]["action"] == "project-member-create"
    assert state.operation_logs[-1]["payload"]["actor_role"] == "admin"

    second_state = _api_state(tmp_path / "second")
    second_state.project_member_store = SqlAlchemyProjectMemberStore(database_url)
    second_client = TestClient(create_app(second_state))
    persisted_members = second_client.get(
        "/projects/CATALOG-LIMIT-202606/members",
        headers={"X-User-Id": "custom-catalog-auditor", "X-Role": "member"},
    ).json()["items"]
    persisted_projects = second_client.get(
        "/projects",
        headers={"X-User-Id": "custom-catalog-auditor", "X-Role": "member"},
    ).json()["items"]
    catalog_project = next(
        item for item in persisted_projects if item["id"] == "CATALOG-LIMIT-202606"
    )

    assert persisted_members[0]["id"] == created["id"]
    assert persisted_members[0]["name"] == "赵审计"
    assert persisted_members[0]["user_identifier"] == "custom-catalog-auditor"
    assert persisted_members[0]["metadata"]["ticket"] == "MEMBER-42"
    assert persisted_members[0]["metadata"]["user_identifier"] == "custom-catalog-auditor"
    assert any(item["id"] == "member-catalog-owner" for item in persisted_members)
    assert catalog_project["member_count"] == 5


def test_projects_api_keeps_pending_custom_member_invisible_with_persistent_store(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.project_member_store = SqlAlchemyProjectMemberStore(
        f"sqlite:///{tmp_path / 'pending-project-members.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    admin_headers = {"X-User-Id": "admin-1", "X-Role": "admin"}

    active_response = client.post(
        "/projects/OUTPATIENT-DOSE-202606/members",
        headers=admin_headers,
        json={
            "user_identifier": "  custom-dose-active  ",
            "name": "自定义门诊审计员",
            "role": "审计员",
            "department": "内审部",
            "status": "在项目中",
            "metadata": {
                "ticket": "DOSE-1",
                "user_identifier": "spoofed-dose-member",
            },
        },
    )
    pending_response = client.post(
        "/projects/OUTPATIENT-DOSE-202606/members",
        headers=admin_headers,
        json={
            "user_identifier": "custom-dose-pending",
            "name": "待确认门诊审计员",
            "role": "审计员",
            "department": "内审部",
            "metadata": {"ticket": "DOSE-2"},
        },
    )

    assert active_response.status_code == 200
    assert active_response.json()["item"]["user_identifier"] == "custom-dose-active"
    assert active_response.json()["item"]["metadata"]["ticket"] == "DOSE-1"
    assert active_response.json()["item"]["metadata"]["user_identifier"] == "custom-dose-active"
    assert pending_response.status_code == 200
    active_headers = {"X-User-Id": "custom-dose-active", "X-Role": "member"}
    pending_headers = {"X-User-Id": "custom-dose-pending", "X-Role": "member"}
    assert [
        item["id"] for item in client.get("/projects", headers=active_headers).json()["items"]
    ] == ["OUTPATIENT-DOSE-202606"]
    assert (
        client.get(
            "/projects/OUTPATIENT-DOSE-202606",
            headers=active_headers,
        ).status_code
        == 200
    )
    assert client.get("/projects", headers=pending_headers).json()["items"] == []
    assert (
        client.get(
            "/projects/OUTPATIENT-DOSE-202606",
            headers=pending_headers,
        ).status_code
        == 404
    )


def test_projects_api_rejects_duplicate_project_member_identities(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.project_member_store = SqlAlchemyProjectMemberStore(
        f"sqlite:///{tmp_path / 'project-member-conflicts.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    admin_headers = {"X-User-Id": "admin-1", "X-Role": "admin"}

    default_conflict = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers=admin_headers,
        json={
            "user_identifier": "next-member",
            "name": "重复默认成员",
            "role": "审计员",
            "department": "内审部",
            "status": "待确认",
        },
    )
    pending_default_conflict = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers=admin_headers,
        json={
            "user_identifier": "next-technician",
            "name": "重复待确认默认成员",
            "role": "信息科",
            "department": "信息科",
            "status": "在项目中",
        },
    )
    active_create = client.post(
        "/projects/CATALOG-LIMIT-202606/members",
        headers=admin_headers,
        json={
            "user_identifier": "duplicate-active-first",
            "name": "先激活成员",
            "role": "审计员",
            "department": "内审部",
            "status": "在项目中",
            "metadata": {
                "ticket": "IDENTITY-ACTIVE",
                "user_identifier": "spoofed-active-identity",
            },
        },
    )
    active_then_pending = client.post(
        "/projects/CATALOG-LIMIT-202606/members",
        headers=admin_headers,
        json={
            "user_identifier": "  duplicate-active-first  ",
            "name": "后待确认成员",
            "role": "审计员",
            "department": "内审部",
            "status": "待确认",
        },
    )
    pending_create = client.post(
        "/projects/CATALOG-LIMIT-202606/members",
        headers=admin_headers,
        json={
            "user_identifier": "duplicate-pending-first",
            "name": "先待确认成员",
            "role": "审计员",
            "department": "内审部",
            "status": "待确认",
        },
    )
    pending_then_active = client.post(
        "/projects/CATALOG-LIMIT-202606/members",
        headers=admin_headers,
        json={
            "user_identifier": "duplicate-pending-first",
            "name": "后激活成员",
            "role": "审计员",
            "department": "内审部",
            "status": "在项目中",
        },
    )

    assert default_conflict.status_code == 409
    assert default_conflict.json()["detail"] == "project member identity already exists"
    assert pending_default_conflict.status_code == 409
    assert pending_default_conflict.json()["detail"] == "project member identity already exists"
    assert active_create.status_code == 200
    assert active_create.json()["item"]["user_identifier"] == "duplicate-active-first"
    assert active_create.json()["item"]["metadata"] == {
        "ticket": "IDENTITY-ACTIVE",
        "user_identifier": "duplicate-active-first",
    }
    assert active_then_pending.status_code == 409
    assert active_then_pending.json()["detail"] == "project member identity already exists"
    assert pending_create.status_code == 200
    assert pending_then_active.status_code == 409
    assert pending_then_active.json()["detail"] == "project member identity already exists"

    members = client.get(
        "/projects/CATALOG-LIMIT-202606/members",
        headers=admin_headers,
    ).json()["items"]
    assert sum(item["user_identifier"] == "duplicate-active-first" for item in members) == 1
    assert sum(item["user_identifier"] == "duplicate-pending-first" for item in members) == 1


@pytest.mark.parametrize("store_kind", ("memory", "sql"))
def test_project_member_reads_dedupe_legacy_identities_consistently(
    tmp_path: Path,
    store_kind: str,
) -> None:
    project_key = "CATALOG-LIMIT-202606"
    raw_members = [
        _legacy_project_member(
            "custom-default-duplicate",
            "expert-catalog",
            "在项目中",
        ),
        _legacy_project_member(
            "custom-shared-new",
            "legacy-shared",
            "待确认",
        ),
        _legacy_project_member(
            "custom-shared-old",
            "legacy-shared",
            "在项目中",
        ),
        _legacy_project_member(
            "custom-active-new",
            "legacy-active",
            "在项目中",
        ),
        _legacy_project_member(
            "custom-active-old",
            "legacy-active",
            "待确认",
        ),
        _legacy_project_member("legacy-identityless-one", None, "在项目中"),
        _legacy_project_member("legacy-identityless-two", None, "待确认"),
    ]
    state = _api_state(tmp_path)
    if store_kind == "memory":
        store: InMemoryProjectMemberStore | SqlAlchemyProjectMemberStore = (
            InMemoryProjectMemberStore(members={project_key: raw_members})
        )
    else:
        database_url = f"sqlite:///{tmp_path / 'legacy-project-members.db'}"
        store = SqlAlchemyProjectMemberStore(database_url, create_schema=True)
        _seed_legacy_project_members(database_url, project_key, raw_members)
    state.project_member_store = store
    client = TestClient(create_app(state))
    admin_headers = {"X-User-Id": "admin-1", "X-Role": "admin"}

    members_response = client.get(
        f"/projects/{project_key}/members",
        headers=admin_headers,
    )
    projects_response = client.get("/projects", headers=admin_headers)
    shared_response = client.get(
        "/projects",
        headers={"X-User-Id": "legacy-shared", "X-Role": "member"},
    )
    active_response = client.get(
        "/projects",
        headers={"X-User-Id": "legacy-active", "X-Role": "member"},
    )
    default_response = client.get(
        "/projects",
        headers={"X-User-Id": "expert-catalog", "X-Role": "member"},
    )

    assert members_response.status_code == 200
    member_ids = {item["id"] for item in members_response.json()["items"]}
    assert len(member_ids) == 8
    assert "custom-default-duplicate" not in member_ids
    assert "custom-shared-new" in member_ids
    assert "custom-shared-old" not in member_ids
    assert "custom-active-new" in member_ids
    assert "custom-active-old" not in member_ids
    assert {"legacy-identityless-one", "legacy-identityless-two"} <= member_ids
    catalog_project = next(
        item for item in projects_response.json()["items"] if item["id"] == project_key
    )
    assert catalog_project["member_count"] == 8
    assert store.member_counts()[project_key] == 4
    assert shared_response.json()["items"] == []
    assert [item["id"] for item in active_response.json()["items"]] == [project_key]
    assert [item["id"] for item in default_response.json()["items"]] == [project_key]


def test_projects_api_store_failure_preserves_only_default_visibility(tmp_path: Path) -> None:
    class FailingProjectMemberStore:
        def list_members(self, project_key: str) -> list[dict[str, object]]:
            raise SQLAlchemyError("member store unavailable")

        def add_member(
            self,
            project_key: str,
            values: dict[str, object],
        ) -> dict[str, object]:
            raise SQLAlchemyError("member store unavailable")

        def member_counts(self) -> dict[str, int]:
            raise SQLAlchemyError("member store unavailable")

    state = _api_state(tmp_path)
    state.project_member_store = FailingProjectMemberStore()
    client = TestClient(create_app(state))

    creator_response = client.get(
        "/projects",
        headers={"X-User-Id": "expert-catalog", "X-Role": "member"},
    )
    unrelated_response = client.get(
        "/projects",
        headers={"X-User-Id": "unrelated-member", "X-Role": "member"},
    )
    admin_response = client.get(
        "/projects",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
    )
    creator_detail_response = client.get(
        "/projects/CATALOG-LIMIT-202606",
        headers={"X-User-Id": "expert-catalog", "X-Role": "member"},
    )
    creator_dashboard_response = client.get(
        "/projects/CATALOG-LIMIT-202606/dashboard",
        headers={"X-User-Id": "expert-catalog", "X-Role": "member"},
    )

    assert [item["id"] for item in creator_response.json()["items"]] == ["CATALOG-LIMIT-202606"]
    assert creator_response.json()["store"] == {
        "ready": False,
        "backend": "unavailable",
        "persistent_writes_ready": False,
        "history_review_task_writes_ready": False,
    }
    assert unrelated_response.json()["items"] == []
    assert len(admin_response.json()["items"]) == 4
    fallback_statuses = {item["status"] for item in admin_response.json()["items"]}
    assert fallback_statuses <= set(admin_response.json()["project_statuses"])
    assert "待启动" not in fallback_statuses
    assert creator_detail_response.status_code == 200
    assert creator_detail_response.json()["item"]["status"] == "待开始"
    assert creator_dashboard_response.status_code == 200
    creator_dashboard = creator_dashboard_response.json()
    assert creator_dashboard["project"]["status"] == "待开始"
    assert creator_dashboard["evidence_grade"] == "backend-defaults"
    assert creator_dashboard["store"] == {
        "ready": False,
        "project_members_ready": False,
        "audit_findings_ready": False,
        "status": "unavailable",
        "backend": {
            "project_members": "unavailable",
            "audit_findings": "unavailable",
        },
    }


def test_project_dashboard_scopes_findings_and_publishes_full_readiness(
    tmp_path: Path,
) -> None:
    class ProjectScopedAuditFindingStore:
        def __init__(self) -> None:
            self.requested_project_keys: list[str | None] = []
            self.items = {
                "SELF-CHECK-FUND-20260607": [
                    {
                        "finding_key": "finding-self-001",
                        "status": "open",
                        "finding_type": "self-check",
                        "severity": "high",
                        "review_status": "pending-review",
                        "review_task_id": "review-self-001",
                        "metadata": {"owner": "SELF负责人"},
                    }
                ],
                "CATALOG-LIMIT-202606": [
                    {
                        "finding_key": "finding-catalog-001",
                        "status": "open",
                        "finding_type": "catalog",
                        "severity": "medium",
                        "review_status": "needs-evidence",
                        "review_task_id": None,
                        "metadata": {"owner": "CATALOG负责人"},
                    },
                    {
                        "finding_key": "finding-catalog-002",
                        "status": "open",
                        "finding_type": "catalog",
                        "severity": "low",
                        "review_status": "pending-review",
                        "review_task_id": None,
                        "metadata": {"owner": "CATALOG负责人"},
                    },
                ],
            }

        def list_findings(
            self,
            *,
            project_key: str | None = None,
            limit: int = 100,
        ) -> list[dict[str, object]]:
            self.requested_project_keys.append(project_key)
            assert project_key is not None
            return [dict(item) for item in self.items.get(project_key, [])[:limit]]

    state = _api_state(tmp_path)
    state.project_member_store = InMemoryProjectMemberStore()
    finding_store = ProjectScopedAuditFindingStore()
    state.audit_finding_store = finding_store  # type: ignore[assignment]
    client = TestClient(create_app(state))
    headers = {"X-User-Id": "admin-1", "X-Role": "admin"}

    self_response = client.get(
        "/projects/SELF-CHECK-FUND-20260607/dashboard",
        headers=headers,
    )
    catalog_response = client.get(
        "/projects/CATALOG-LIMIT-202606/dashboard",
        headers=headers,
    )

    assert self_response.status_code == 200
    assert catalog_response.status_code == 200
    assert finding_store.requested_project_keys == [
        "SELF-CHECK-FUND-20260607",
        "CATALOG-LIMIT-202606",
    ]
    self_body = self_response.json()
    catalog_body = catalog_response.json()
    assert [item["id"] for item in self_body["queue"]] == ["finding-self-001"]
    assert [item["id"] for item in catalog_body["queue"]] == [
        "finding-catalog-001",
        "finding-catalog-002",
    ]
    assert {item["key"]: item["value"] for item in self_body["metrics"]}["open_findings"] == "1"
    assert {item["key"]: item["value"] for item in catalog_body["metrics"]}["open_findings"] == "2"
    assert "CATALOG负责人" not in {item["name"] for item in self_body["member_workloads"]}
    assert "SELF负责人" not in {item["name"] for item in catalog_body["member_workloads"]}
    assert self_body["evidence_grade"] == "live-db-connected"
    assert self_body["store"] == {
        "ready": True,
        "project_members_ready": True,
        "audit_findings_ready": True,
        "status": "ready",
        "backend": {
            "project_members": "InMemoryProjectMemberStore",
            "audit_findings": "ProjectScopedAuditFindingStore",
        },
    }


def test_project_dashboard_keeps_explicitly_closed_findings_out_of_open_metric(
    tmp_path: Path,
) -> None:
    class ClosedAndLegacyFindingStore:
        def list_findings(
            self,
            *,
            project_key: str | None = None,
            limit: int = 100,
        ) -> list[dict[str, object]]:
            assert project_key is not None
            status = "closed" if project_key == "SELF-CHECK-FUND-20260607" else None
            return [
                {
                    "finding_key": f"finding-{project_key}",
                    "status": status,
                    "severity": "medium",
                    "review_status": "confirmed-violation",
                    "review_task_id": None,
                    "metadata": {},
                }
            ][:limit]

    state = _api_state(tmp_path)
    state.project_member_store = InMemoryProjectMemberStore()
    state.audit_finding_store = ClosedAndLegacyFindingStore()  # type: ignore[assignment]
    client = TestClient(create_app(state))
    headers = {"X-User-Id": "admin-1", "X-Role": "admin"}

    closed_response = client.get(
        "/projects/SELF-CHECK-FUND-20260607/dashboard",
        headers=headers,
    )
    legacy_response = client.get(
        "/projects/CATALOG-LIMIT-202606/dashboard",
        headers=headers,
    )

    assert closed_response.status_code == 200
    assert legacy_response.status_code == 200
    closed_metrics = {item["key"]: item["value"] for item in closed_response.json()["metrics"]}
    legacy_metrics = {item["key"]: item["value"] for item in legacy_response.json()["metrics"]}
    assert closed_metrics["open_findings"] == "0"
    assert legacy_metrics["open_findings"] == "1"


def test_audit_finding_store_filters_findings_by_project_in_sql(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'project-findings.db'}"
    store = SqlAlchemyAuditFindingStore(database_url, create_schema=True)
    _seed_project_findings(database_url)

    global_items = store.list_findings(limit=10)
    self_items = store.list_findings(
        project_key="SELF-CHECK-FUND-20260607",
        limit=10,
    )
    catalog_items = store.list_findings(
        project_key="CATALOG-LIMIT-202606",
        limit=10,
    )
    visible_items = store.list_findings(
        project_keys=frozenset({"SELF-CHECK-FUND-20260607"}),
        limit=10,
    )
    no_visible_items = store.list_findings(project_keys=frozenset(), limit=10)
    missing_items = store.list_findings(project_key="UNKNOWN-PROJECT", limit=10)
    self_count = store.count_findings(project_keys=frozenset({"SELF-CHECK-FUND-20260607"}))
    both_count = store.count_findings(
        project_keys=frozenset({"SELF-CHECK-FUND-20260607", "CATALOG-LIMIT-202606"})
    )
    empty_count = store.count_findings(project_keys=frozenset())

    assert {item["finding_key"] for item in global_items} == {
        "finding-self-sql",
        "finding-catalog-sql",
    }
    assert [item["finding_key"] for item in self_items] == ["finding-self-sql"]
    assert [item["finding_key"] for item in catalog_items] == ["finding-catalog-sql"]
    assert [item["finding_key"] for item in visible_items] == ["finding-self-sql"]
    assert no_visible_items == []
    assert missing_items == []
    assert self_count == 1
    assert both_count == 2
    assert empty_count == 0


def test_audit_finding_store_project_contract_is_stable_across_reads(
    tmp_path: Path,
) -> None:
    database_url = f"sqlite:///{tmp_path / 'project-finding-contract.db'}"
    store = SqlAlchemyAuditFindingStore(database_url, create_schema=True)
    _seed_project_findings(database_url)
    engine = create_engine(database_url, connect_args={"check_same_thread": False})
    with Session(engine) as session:
        self_task = session.scalar(
            select(ReviewTask).where(ReviewTask.external_task_id == "review-self-sql")
        )
        catalog_finding = session.scalar(
            select(AuditFinding).where(AuditFinding.finding_key == "finding-catalog-sql")
        )
        assert self_task is not None
        assert catalog_finding is not None
        catalog_finding.review_task_id = self_task.id
        session.commit()

    listed = store.list_findings(
        project_key="SELF-CHECK-FUND-20260607",
        limit=10,
    )
    loaded = store.get_finding("finding-self-sql")
    synced = store.sync_review_task_status(
        "review-self-sql",
        "confirmed-violation",
        project_keys=frozenset({"SELF-CHECK-FUND-20260607"}),
    )
    catalog_after_sync = store.get_finding("finding-catalog-sql")

    assert [item["project_key"] for item in listed] == ["SELF-CHECK-FUND-20260607"]
    assert loaded["project_key"] == "SELF-CHECK-FUND-20260607"
    assert [item["project_key"] for item in synced] == ["SELF-CHECK-FUND-20260607"]
    assert catalog_after_sync["review_status"] == "pending-review"


def test_graph_workbench_defaults_to_knowledge_catalog_without_project_store_reads(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    member_store = _RecordingProjectMemberStore()
    finding_store = _RecordingGraphFindingStore()
    review_store = _RecordingGraphReviewTaskStore()
    state.project_member_store = member_store
    state.audit_finding_store = finding_store  # type: ignore[assignment]
    state.review_task_store = review_store
    client = TestClient(create_app(state))

    response = client.get(
        "/graph/workbench",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["view"] == "knowledge"
    assert body["project_key"] is None
    assert body["evidence_chain_status"] == "catalog"
    assert body["production_side_effect"] == "none"
    assert "graph-node-project" in {node["id"] for node in body["nodes"]}
    assert any(str(node["href"]).startswith("/knowledge-base") for node in body["nodes"])
    assert member_store.list_reads == 0
    assert finding_store.requested_project_keys == []
    assert review_store.list_calls == 0
    assert state.operation_logs[-1]["payload"]["view"] == "knowledge"
    assert state.operation_logs[-1]["payload"]["project_key"] is None


def test_graph_workbench_rejects_invalid_view_project_pairs_before_data_reads(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    member_store = _RecordingProjectMemberStore()
    finding_store = _RecordingGraphFindingStore()
    review_store = _RecordingGraphReviewTaskStore()
    state.project_member_store = member_store
    state.audit_finding_store = finding_store  # type: ignore[assignment]
    state.review_task_store = review_store
    client = TestClient(create_app(state))

    missing = client.get("/graph/workbench?view=project")
    blank = client.get("/graph/workbench?view=project&project_key=%20%20")
    forbidden = client.get("/graph/workbench?view=knowledge&project_key=SELF-CHECK-FUND-20260607")
    invalid_view = client.get("/graph/workbench?view=workflow")
    too_long = client.get(f"/graph/workbench?view=project&project_key={'P' * 129}")

    assert missing.status_code == 422
    assert missing.json()["detail"] == "project_key is required for project graph view"
    assert blank.status_code == 422
    assert blank.json()["detail"] == "project_key is required for project graph view"
    assert forbidden.status_code == 422
    assert forbidden.json()["detail"] == ("project_key is not allowed for knowledge graph view")
    assert invalid_view.status_code == 422
    assert too_long.status_code == 422
    assert member_store.list_reads == 0
    assert finding_store.requested_project_keys == []
    assert review_store.list_calls == 0
    assert not any(log["action"] == "graph-workbench-view" for log in state.operation_logs)


def test_project_graph_authorizes_exact_project_scope_and_hides_unrelated_projects(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'graph-auth.db'}",
        create_schema=True,
    )
    state.project_member_store = InMemoryProjectMemberStore()
    finding_store = _RecordingGraphFindingStore()
    review_store = _RecordingGraphReviewTaskStore()
    state.audit_finding_store = finding_store  # type: ignore[assignment]
    state.review_task_store = review_store
    state.auth_user_store.add_user(
        {
            "user_key": "graph-project-admin",
            "display_name": "图谱项目管理员",
            "department_key": "audit-office",
        }
    )
    state.auth_user_store.assign_role(
        "graph-project-admin",
        {
            "role": "admin",
            "scope_type": "project",
            "scope_key": "SELF-CHECK-FUND-20260607",
        },
    )
    state.auth_user_store.add_user(
        {
            "user_key": "graph-department-admin",
            "display_name": "图谱科室管理员",
            "department_key": "audit-office",
        }
    )
    state.auth_user_store.assign_role(
        "graph-department-admin",
        {
            "role": "admin",
            "scope_type": "department",
            "scope_key": "audit-office",
        },
    )
    client = TestClient(create_app(state))

    visible_member = client.get(
        "/graph/workbench?view=project&project_key=SELF-CHECK-FUND-20260607",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )
    unrelated = client.get(
        "/graph/workbench?view=project&project_key=SELF-CHECK-FUND-20260607",
        headers={"X-User-Id": "unrelated-member", "X-Role": "member"},
    )
    scoped_allowed = client.get(
        "/graph/workbench?view=project&project_key=SELF-CHECK-FUND-20260607",
        headers={"X-User-Id": "graph-project-admin", "X-Role": "member"},
    )
    scoped_other = client.get(
        "/graph/workbench?view=project&project_key=CATALOG-LIMIT-202606",
        headers={"X-User-Id": "graph-project-admin", "X-Role": "admin"},
    )
    department_scoped = client.get(
        "/graph/workbench?view=project&project_key=SELF-CHECK-FUND-20260607",
        headers={"X-User-Id": "graph-department-admin", "X-Role": "admin"},
    )
    unknown = client.get(
        "/graph/workbench?view=project&project_key=UNKNOWN-PROJECT",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
    )

    assert visible_member.status_code == 200
    assert unrelated.status_code == 404
    assert unrelated.json()["detail"] == "project not found"
    assert scoped_allowed.status_code == 200
    assert scoped_other.status_code == 404
    assert department_scoped.status_code == 404
    assert unknown.status_code == 404
    assert finding_store.requested_project_keys == [
        "SELF-CHECK-FUND-20260607",
        "SELF-CHECK-FUND-20260607",
    ]
    assert review_store.list_calls == 2
    denial_logs = [log for log in state.operation_logs if log["action"] == "authorization-denied"]
    assert len(denial_logs) == 3
    assert all("finding" not in json.dumps(log).lower() for log in denial_logs)
    assert all("task" not in json.dumps(log).lower() for log in denial_logs)
    success_logs = [log for log in state.operation_logs if log["action"] == "graph-workbench-view"]
    assert [log["payload"]["project_key"] for log in success_logs] == [
        "SELF-CHECK-FUND-20260607",
        "SELF-CHECK-FUND-20260607",
    ]


def test_project_graph_fails_closed_when_required_stores_are_unavailable(
    tmp_path: Path,
) -> None:
    class FailingProjectMemberStore(_RecordingProjectMemberStore):
        def list_members(self, project_key: str) -> list[dict[str, object]]:
            self.list_reads += 1
            raise SQLAlchemyError("project member store unavailable")

    class FailingFindingStore(_RecordingGraphFindingStore):
        def list_findings(
            self,
            *,
            review_status: str | None = None,
            project_key: str | None = None,
            limit: int = 100,
        ) -> list[dict[str, object]]:
            self.requested_project_keys.append(project_key)
            raise SQLAlchemyError("finding store unavailable")

    class FailingReviewTaskStore(_RecordingGraphReviewTaskStore):
        def list_tasks(self) -> list[dict[str, object]]:
            self.list_calls += 1
            raise SQLAlchemyError("review task store unavailable")

    route = "/graph/workbench?view=project&project_key=SELF-CHECK-FUND-20260607"
    headers = {"X-User-Id": "next-member", "X-Role": "member"}

    missing_member_state = _api_state(tmp_path / "missing-member")
    missing_member_state.project_member_store = None
    missing_member_finding = _RecordingGraphFindingStore()
    missing_member_review = _RecordingGraphReviewTaskStore()
    missing_member_state.audit_finding_store = missing_member_finding  # type: ignore[assignment]
    missing_member_state.review_task_store = missing_member_review
    missing_member = TestClient(create_app(missing_member_state)).get(route, headers=headers)

    failing_member_state = _api_state(tmp_path / "failing-member")
    failing_member_state.project_member_store = FailingProjectMemberStore()
    failing_member_finding = _RecordingGraphFindingStore()
    failing_member_review = _RecordingGraphReviewTaskStore()
    failing_member_state.audit_finding_store = failing_member_finding  # type: ignore[assignment]
    failing_member_state.review_task_store = failing_member_review
    failing_member = TestClient(create_app(failing_member_state)).get(route, headers=headers)

    missing_finding_state = _api_state(tmp_path / "missing-finding")
    missing_finding_state.project_member_store = InMemoryProjectMemberStore()
    missing_finding_state.audit_finding_store = None
    missing_finding_review = _RecordingGraphReviewTaskStore()
    missing_finding_state.review_task_store = missing_finding_review
    missing_finding = TestClient(create_app(missing_finding_state)).get(route, headers=headers)

    missing_review_state = _api_state(tmp_path / "missing-review")
    missing_review_state.project_member_store = InMemoryProjectMemberStore()
    missing_review_finding = _RecordingGraphFindingStore()
    missing_review_state.audit_finding_store = missing_review_finding  # type: ignore[assignment]
    missing_review_state.review_task_store = None
    missing_review = TestClient(create_app(missing_review_state)).get(route, headers=headers)

    failing_finding_state = _api_state(tmp_path / "failing-finding")
    failing_finding_state.project_member_store = InMemoryProjectMemberStore()
    failing_finding = FailingFindingStore()
    failing_finding_state.audit_finding_store = failing_finding  # type: ignore[assignment]
    failing_finding_state.review_task_store = _RecordingGraphReviewTaskStore()
    finding_error = TestClient(create_app(failing_finding_state)).get(route, headers=headers)

    failing_review_state = _api_state(tmp_path / "failing-review")
    failing_review_state.project_member_store = InMemoryProjectMemberStore()
    failing_review_finding = _RecordingGraphFindingStore()
    failing_review_state.audit_finding_store = failing_review_finding  # type: ignore[assignment]
    failing_review_store = FailingReviewTaskStore()
    failing_review_state.review_task_store = failing_review_store
    review_error = TestClient(create_app(failing_review_state)).get(route, headers=headers)

    assert missing_member.status_code == 503
    assert failing_member.status_code == 503
    assert missing_member_finding.requested_project_keys == []
    assert missing_member_review.list_calls == 0
    assert failing_member_finding.requested_project_keys == []
    assert failing_member_review.list_calls == 0
    assert missing_finding.status_code == 503
    assert missing_finding_review.list_calls == 0
    assert missing_review.status_code == 503
    assert missing_review_finding.requested_project_keys == []
    assert finding_error.status_code == 503
    assert review_error.status_code == 503
    assert failing_review_finding.requested_project_keys == ["SELF-CHECK-FUND-20260607"]
    assert failing_review_store.list_calls == 1


def test_project_graph_admin_can_read_known_project_without_member_store(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.project_member_store = None
    finding_store = _RecordingGraphFindingStore()
    review_store = _RecordingGraphReviewTaskStore()
    state.audit_finding_store = finding_store  # type: ignore[assignment]
    state.review_task_store = review_store
    client = TestClient(create_app(state))

    response = client.get(
        "/graph/workbench?view=project&project_key=CATALOG-LIMIT-202606",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
    )

    assert response.status_code == 200
    assert response.json()["project_key"] == "CATALOG-LIMIT-202606"
    assert finding_store.requested_project_keys == ["CATALOG-LIMIT-202606"]
    assert review_store.list_calls == 1


def test_project_graph_builds_only_exact_persisted_evidence_chain(
    tmp_path: Path,
) -> None:
    database_url = f"sqlite:///{tmp_path / 'project-graph.db'}"
    finding_store = SqlAlchemyAuditFindingStore(database_url, create_schema=True)
    _seed_project_findings(database_url)
    review_store = _CountingSqlAlchemyReviewTaskStore(database_url)
    state = _api_state(tmp_path)
    state.project_member_store = InMemoryProjectMemberStore()
    state.audit_finding_store = finding_store
    state.review_task_store = review_store
    client = TestClient(create_app(state))

    self_response = client.get(
        "/graph/workbench?view=project&project_key=SELF-CHECK-FUND-20260607",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )
    catalog_response = client.get(
        "/graph/workbench?view=project&project_key=CATALOG-LIMIT-202606",
        headers={"X-User-Id": "expert-catalog", "X-Role": "member"},
    )

    assert self_response.status_code == 200
    assert catalog_response.status_code == 200
    assert review_store.list_calls == 2
    self_body = self_response.json()
    catalog_body = catalog_response.json()
    assert self_body["view"] == "project"
    assert self_body["project_key"] == "SELF-CHECK-FUND-20260607"
    assert self_body["evidence_chain_status"] == "ready"
    assert self_body["evidence_grade"] == "live-store-readonly"
    assert self_body["store"] == {
        "ready": True,
        "backend": {
            "audit_findings": "SqlAlchemyAuditFindingStore",
            "review_tasks": "_CountingSqlAlchemyReviewTaskStore",
        },
    }
    self_ids = {node["id"] for node in self_body["nodes"]}
    catalog_ids = {node["id"] for node in catalog_body["nodes"]}
    allowed_statuses = {"已归集", "可引用", "待复核", "门禁中", "跟踪中", "待接入"}
    assert {
        node["status"] for node in [*self_body["nodes"], *catalog_body["nodes"]]
    } <= allowed_statuses
    assert {
        "project:SELF-CHECK-FUND-20260607",
        "finding:finding-self-sql",
        "review:review-self-sql",
        "report:review-self-sql",
        "remediation:rectification-self-sql",
        "review:report-draft-self-sql",
        "report:report-draft-self-sql",
    } <= self_ids
    assert {
        "project:CATALOG-LIMIT-202606",
        "finding:finding-catalog-sql",
        "review:review-catalog-sql",
        "report:review-catalog-sql",
        "remediation:rectification-catalog-sql",
        "review:report-draft-catalog-sql",
        "report:report-draft-catalog-sql",
    } <= catalog_ids
    assert not any("catalog" in node_id.lower() for node_id in self_ids)
    assert not any("self" in node_id.lower() for node_id in catalog_ids)
    assert "review:legacy-project-string-sql" not in self_ids | catalog_ids
    assert "report:legacy-project-string-sql" not in self_ids | catalog_ids
    assert all(
        str(node["href"]).startswith(
            ("/projects?project=", "/findings", "/reports", "/remediation")
        )
        for node in [*self_body["nodes"], *catalog_body["nodes"]]
    )
    serialized = json.dumps([self_body, catalog_body], ensure_ascii=False)
    assert "SENSITIVE-GRAPH-BODY" not in serialized
    assert "field_values" not in serialized
    success_payload = state.operation_logs[-1]["payload"]
    assert success_payload["view"] == "project"
    assert success_payload["project_key"] == "CATALOG-LIMIT-202606"
    assert success_payload["node_count"] == len(catalog_body["nodes"])
    assert "finding" not in success_payload
    assert "task" not in success_payload


def test_project_graph_normalizes_business_statuses_and_requires_canonical_signed_report(
    tmp_path: Path,
) -> None:
    project_key = "SELF-CHECK-FUND-20260607"
    finding_statuses = {
        "pending": "pending-review",
        "needs": "needs-evidence",
        "confirmed": "confirmed-violation",
        "rule": "rule-issue",
        "data": "data-issue",
        "excluded": "not-violation",
        "closed": "closed",
        "unknown": "future-status",
    }
    findings = [
        {
            "project_key": project_key,
            "finding_key": f"finding-status-{suffix}",
            "severity": "medium",
            "review_status": status,
            "review_task_id": f"review-status-{suffix}"
            if suffix in {"pending", "needs", "confirmed", "closed"}
            else None,
            "audit_task_key": f"task-status-{suffix}",
            "rule_version_key": "STATUS-RULE@v1",
            "evidence_items": [],
        }
        for suffix, status in finding_statuses.items()
    ]
    tasks = [
        {
            "task_id": "review-status-pending",
            "status": "pending-review",
            "status_label": "待复核原始标签",
            "citation_count": 0,
            "dossier": {
                "report_draft": {"title": "真实报告草稿"},
                "rectification": {
                    "rectification_id": "rectification-status-pending",
                    "status": "pending-rectification",
                },
            },
        },
        {
            "task_id": "review-status-needs",
            "status": "needs-evidence",
            "status_label": "需补证原始标签",
            "citation_count": 1,
            "dossier": {
                "report_draft": {"title": "缺 hash 时保留的真实草稿"},
                "signed_report": {
                    "status": "signed",
                    "report_id": "signed-report-missing-hash",
                    "content": "SENSITIVE-INCOMPLETE-SIGNED-CONTENT",
                },
                "rectification": {
                    "rectification_id": "rectification-status-in-progress",
                    "status": "in-progress",
                },
            },
        },
        {
            "task_id": "review-status-confirmed",
            "status": "confirmed-violation",
            "status_label": "确认违规原始标签",
            "citation_count": 2,
            "dossier": {
                "signed_report": {
                    "status": "signed",
                    "report_id": "signed-report-canonical",
                    "content_sha256": "sha256-canonical",
                    "content": "SENSITIVE-CANONICAL-SIGNED-CONTENT",
                },
                "rectification": {
                    "rectification_id": "rectification-status-accepted",
                    "status": "accepted",
                },
            },
        },
        {
            "task_id": "review-status-closed",
            "status": "closed",
            "status_label": "已关闭原始标签",
            "citation_count": 0,
            "dossier": {
                "signed_report": {
                    "status": "signed",
                    "report_id": "signed-report-missing-content",
                    "content_sha256": "sha256-without-content",
                }
            },
        },
    ]
    for status in ("submitted", "returned"):
        tasks.append(
            {
                "task_id": f"report-draft-rectification-{status}",
                "status": "closed",
                "status_label": "已关闭原始标签",
                "citation_count": 0,
                "dossier": {
                    "report_template_draft": {
                        "status": "draft",
                        "template_id": f"template-{status}",
                        "project_key": project_key,
                    },
                    "rectification": {
                        "rectification_id": f"rectification-status-{status}",
                        "status": status,
                    },
                },
            }
        )

    state = _api_state(tmp_path)
    state.project_member_store = InMemoryProjectMemberStore()
    state.audit_finding_store = _RecordingGraphFindingStore(  # type: ignore[assignment]
        {project_key: findings}
    )
    state.review_task_store = _RecordingGraphReviewTaskStore(tasks)
    client = TestClient(create_app(state))

    response = client.get(
        f"/graph/workbench?view=project&project_key={project_key}",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )

    assert response.status_code == 200
    body = response.json()
    nodes = {node["id"]: node for node in body["nodes"]}
    allowed_statuses = {"已归集", "可引用", "待复核", "门禁中", "跟踪中", "待接入"}
    assert {node["status"] for node in nodes.values()} <= allowed_statuses
    assert nodes[f"project:{project_key}"]["status"] == "已归集"
    assert "进行中" in nodes[f"project:{project_key}"]["description"]
    assert nodes["finding:finding-status-pending"]["status"] == "待复核"
    assert nodes["finding:finding-status-needs"]["status"] == "门禁中"
    for suffix in ("confirmed", "rule", "data", "excluded", "closed"):
        assert nodes[f"finding:finding-status-{suffix}"]["status"] == "已归集"
    assert nodes["finding:finding-status-unknown"]["status"] == "待复核"
    assert "future-status" in nodes["finding:finding-status-unknown"]["description"]
    assert nodes["review:review-status-pending"]["status"] == "待复核"
    assert nodes["review:review-status-needs"]["status"] == "门禁中"
    assert nodes["review:review-status-confirmed"]["status"] == "已归集"
    assert nodes["review:review-status-closed"]["status"] == "已归集"
    assert "confirmed-violation" in nodes["review:review-status-confirmed"]["description"]
    assert "确认违规原始标签" in nodes["review:review-status-confirmed"]["description"]
    assert nodes["report:review-status-pending"]["status"] == "门禁中"
    assert nodes["report:review-status-needs"]["status"] == "门禁中"
    assert nodes["report:review-status-confirmed"]["status"] == "已归集"
    assert "signed" in nodes["report:review-status-confirmed"]["description"]
    assert "report:review-status-closed" not in nodes
    assert nodes["remediation:rectification-status-pending"]["status"] == "跟踪中"
    assert nodes["remediation:rectification-status-in-progress"]["status"] == "跟踪中"
    assert nodes["remediation:rectification-status-submitted"]["status"] == "跟踪中"
    assert nodes["remediation:rectification-status-returned"]["status"] == "跟踪中"
    assert nodes["remediation:rectification-status-accepted"]["status"] == "已归集"
    assert "accepted" in nodes["remediation:rectification-status-accepted"]["description"]
    serialized = json.dumps(body, ensure_ascii=False)
    assert "SENSITIVE-INCOMPLETE-SIGNED-CONTENT" not in serialized
    assert "SENSITIVE-CANONICAL-SIGNED-CONTENT" not in serialized


@pytest.mark.parametrize(
    "duplicate_project_key",
    ("CATALOG-LIMIT-202606", "SELF-CHECK-FUND-20260607"),
    ids=("cross-project", "same-project"),
)
def test_project_graph_rejects_duplicate_review_task_ids_without_data_leakage(
    tmp_path: Path,
    duplicate_project_key: str,
) -> None:
    project_key = "SELF-CHECK-FUND-20260607"
    finding_store = _RecordingGraphFindingStore(
        {
            project_key: [
                {
                    "project_key": project_key,
                    "finding_key": "finding-duplicate-review-id",
                    "severity": "high",
                    "review_status": "pending-review",
                    "review_task_id": "review-duplicate-id",
                    "audit_task_key": "task-duplicate-id",
                    "rule_version_key": "DUPLICATE-ID-RULE@v1",
                    "evidence_items": [],
                }
            ]
        }
    )
    review_store = _RecordingGraphReviewTaskStore(
        [
            {
                "task_id": "review-duplicate-id",
                "status": "pending-review",
                "status_label": "待复核",
                "citation_count": 0,
                "dossier": {"report_draft": {"title": "A report draft"}},
            },
            {
                "task_id": "review-duplicate-id",
                "status": "confirmed-violation",
                "status_label": "确认违规",
                "citation_count": 1,
                "dossier": {
                    "report_template_draft": {
                        "status": "draft",
                        "template_id": "B-REPORT-MARKER",
                        "project_key": duplicate_project_key,
                    },
                    "signed_report": {
                        "status": "signed",
                        "report_id": "B-REPORT-MARKER",
                        "content_sha256": "B-HASH-MARKER",
                        "content": "B-CONTENT-MARKER",
                    },
                    "rectification": {
                        "rectification_id": "B-RECTIFICATION-MARKER",
                        "status": "accepted",
                    },
                },
            },
        ]
    )
    state = _api_state(tmp_path)
    state.project_member_store = InMemoryProjectMemberStore()
    state.audit_finding_store = finding_store  # type: ignore[assignment]
    state.review_task_store = review_store
    client = TestClient(create_app(state))

    response = client.get(
        f"/graph/workbench?view=project&project_key={project_key}",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )

    assert response.status_code == 503
    assert response.json()["detail"] == "review task store contains duplicate task ids"
    serialized = json.dumps(
        {"response": response.json(), "operation_logs": state.operation_logs},
        ensure_ascii=False,
    )
    assert "B-REPORT-MARKER" not in serialized
    assert "B-RECTIFICATION-MARKER" not in serialized
    assert "B-CONTENT-MARKER" not in serialized
    assert not any(log["action"] == "graph-workbench-view" for log in state.operation_logs)
    assert finding_store.requested_project_keys == [project_key]
    assert review_store.list_calls == 1


def test_project_graph_rejects_linked_review_task_with_conflicting_project_scope(
    tmp_path: Path,
) -> None:
    project_key = "SELF-CHECK-FUND-20260607"
    finding_store = _RecordingGraphFindingStore(
        {
            project_key: [
                {
                    "project_key": project_key,
                    "finding_key": "finding-conflicting-review-project",
                    "severity": "high",
                    "review_status": "confirmed-violation",
                    "review_task_id": "review-conflicting-project",
                    "audit_task_key": "task-conflicting-project",
                    "rule_version_key": "CONFLICTING-PROJECT@v1",
                    "evidence_items": [],
                }
            ]
        }
    )
    review_store = _RecordingGraphReviewTaskStore(
        [
            {
                "task_id": "review-conflicting-project",
                "status": "confirmed-violation",
                "status_label": "确认违规",
                "citation_count": 1,
                "dossier": {
                    "report_template_draft": {
                        "status": "draft",
                        "template_id": "CROSS-PROJECT-REPORT-MARKER",
                        "project_key": "CATALOG-LIMIT-202606",
                    },
                    "signed_report": {
                        "status": "signed",
                        "report_id": "CROSS-PROJECT-REPORT-MARKER",
                        "content_sha256": "CROSS-PROJECT-HASH",
                        "content": "CROSS-PROJECT-CONTENT",
                    },
                    "rectification": {
                        "rectification_id": "CROSS-PROJECT-RECTIFICATION-MARKER",
                        "status": "accepted",
                    },
                },
            }
        ]
    )
    state = _api_state(tmp_path)
    state.project_member_store = InMemoryProjectMemberStore()
    state.audit_finding_store = finding_store  # type: ignore[assignment]
    state.review_task_store = review_store
    client = TestClient(create_app(state))

    response = client.get(
        f"/graph/workbench?view=project&project_key={project_key}",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )

    assert response.status_code == 503
    assert response.json()["detail"] == "review task project scope conflicts with linked finding"
    serialized = json.dumps(
        {"response": response.json(), "operation_logs": state.operation_logs},
        ensure_ascii=False,
    )
    assert "CROSS-PROJECT-REPORT-MARKER" not in serialized
    assert "CROSS-PROJECT-RECTIFICATION-MARKER" not in serialized
    assert "CROSS-PROJECT-CONTENT" not in serialized
    assert not any(log["action"] == "graph-workbench-view" for log in state.operation_logs)


def test_project_graph_ignores_blank_review_task_ids_and_fails_fast_on_malformed_rows(
    tmp_path: Path,
) -> None:
    class RawReviewTaskStore:
        def __init__(self, rows: list[object]) -> None:
            self.rows = rows
            self.list_calls = 0

        def list_tasks(self) -> list[dict[str, object]]:
            self.list_calls += 1
            return self.rows  # type: ignore[return-value]

    project_key = "SELF-CHECK-FUND-20260607"
    state = _api_state(tmp_path / "blank")
    state.project_member_store = InMemoryProjectMemberStore()
    state.audit_finding_store = _RecordingGraphFindingStore()  # type: ignore[assignment]
    blank_store = RawReviewTaskStore(
        [
            {
                "task_id": "   ",
                "dossier": {
                    "report_template_draft": {
                        "project_key": project_key,
                        "template_id": "BLANK-ID-MARKER",
                        "status": "draft",
                    }
                },
            }
        ]
    )
    state.review_task_store = blank_store  # type: ignore[assignment]
    response = TestClient(create_app(state)).get(
        f"/graph/workbench?view=project&project_key={project_key}",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )

    assert response.status_code == 200
    assert response.json()["evidence_chain_status"] == "empty"
    assert "BLANK-ID-MARKER" not in response.text
    assert blank_store.list_calls == 1

    malformed_state = _api_state(tmp_path / "malformed")
    malformed_state.project_member_store = InMemoryProjectMemberStore()
    malformed_state.audit_finding_store = (  # type: ignore[assignment]
        _RecordingGraphFindingStore()
    )
    malformed_store = RawReviewTaskStore(["malformed-row"])
    malformed_state.review_task_store = malformed_store  # type: ignore[assignment]
    malformed_client = TestClient(create_app(malformed_state))

    with pytest.raises(AttributeError):
        malformed_client.get(
            f"/graph/workbench?view=project&project_key={project_key}",
            headers={"X-User-Id": "next-member", "X-Role": "member"},
        )
    assert malformed_store.list_calls == 1
    assert not any(
        log["action"] == "graph-workbench-view" for log in malformed_state.operation_logs
    )


def test_project_graph_empty_chain_is_distinct_from_unavailable_store(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.project_member_store = InMemoryProjectMemberStore()
    state.audit_finding_store = _RecordingGraphFindingStore()  # type: ignore[assignment]
    state.review_task_store = _RecordingGraphReviewTaskStore()
    client = TestClient(create_app(state))

    response = client.get(
        "/graph/workbench?view=project&project_key=SELF-CHECK-FUND-20260607",
        headers={"X-User-Id": "next-member", "X-Role": "member"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["evidence_chain_status"] == "empty"
    assert body["relations"] == []
    assert [node["id"] for node in body["nodes"]] == ["project:SELF-CHECK-FUND-20260607"]
    assert body["nodes"][0]["metric"] == "0 条项目证据"
    assert "0 条项目证据" in body["nodes"][0]["description"]
    assert body["store"]["ready"] is True


def test_projects_api_marks_split_visibility_store_failure_degraded(tmp_path: Path) -> None:
    class SplitFailureProjectMemberStore:
        def list_members(self, project_key: str) -> list[dict[str, object]]:
            raise SQLAlchemyError("member visibility unavailable")

        def add_member(
            self,
            project_key: str,
            values: dict[str, object],
        ) -> dict[str, object]:
            raise SQLAlchemyError("member store unavailable")

        def member_counts(self) -> dict[str, int]:
            return {"SELF-CHECK-FUND-20260607": 99}

    class ReadyAuditFindingStore:
        def __init__(self) -> None:
            self.requested_project_keys: list[str | None] = []

        def list_findings(
            self,
            *,
            project_key: str | None = None,
            limit: int,
        ) -> list[dict[str, object]]:
            self.requested_project_keys.append(project_key)
            return []

    state = _api_state(tmp_path)
    state.project_member_store = SplitFailureProjectMemberStore()
    finding_store = ReadyAuditFindingStore()
    state.audit_finding_store = finding_store  # type: ignore[assignment]
    client = TestClient(create_app(state))
    visible_headers = (
        {"X-User-Id": "next-director", "X-Role": "director"},
        {"X-User-Id": "next-member", "X-Role": "member"},
    )

    for headers in visible_headers:
        list_response = client.get("/projects", headers=headers)
        assert list_response.status_code == 200
        assert [item["id"] for item in list_response.json()["items"]] == [
            "SELF-CHECK-FUND-20260607"
        ]
        assert list_response.json()["items"][0]["member_count"] is None
        assert list_response.json()["store"] == {
            "ready": False,
            "backend": "unavailable",
            "persistent_writes_ready": False,
            "history_review_task_writes_ready": False,
        }
        assert state.operation_logs[-1]["payload"]["visibility_ready"] is False

        detail_response = client.get(
            "/projects/SELF-CHECK-FUND-20260607",
            headers=headers,
        )
        assert detail_response.status_code == 200
        assert detail_response.json()["item"]["member_count"] is None
        assert detail_response.json()["store"] == {
            "ready": False,
            "backend": "unavailable",
        }
        assert state.operation_logs[-1]["payload"]["visibility_ready"] is False

        members_response = client.get(
            "/projects/SELF-CHECK-FUND-20260607/members",
            headers=headers,
        )
        assert members_response.status_code == 200
        assert len(members_response.json()["items"]) == 3
        assert all(item["source"] == "system-default" for item in members_response.json()["items"])
        assert members_response.json()["store"] == {
            "ready": False,
            "backend": "unavailable",
        }
        assert state.operation_logs[-1]["payload"]["visibility_ready"] is False

        dashboard_response = client.get(
            "/projects/SELF-CHECK-FUND-20260607/dashboard",
            headers=headers,
        )
        assert dashboard_response.status_code == 200
        assert dashboard_response.json()["project"]["member_count"] is None
        assert dashboard_response.json()["store"]["backend"] == {
            "project_members": "unavailable",
            "audit_findings": "ReadyAuditFindingStore",
        }
        assert dashboard_response.json()["store"]["ready"] is False
        assert dashboard_response.json()["store"]["project_members_ready"] is False
        assert dashboard_response.json()["store"]["audit_findings_ready"] is True
        assert dashboard_response.json()["store"]["status"] == "partial"
        assert dashboard_response.json()["evidence_grade"] == "partial-live-db-connected"
        assert finding_store.requested_project_keys[-1] == "SELF-CHECK-FUND-20260607"
        assert state.operation_logs[-1]["payload"]["visibility_ready"] is False

    for user_identifier in ("unrelated-member", "custom-member"):
        headers = {"X-User-Id": user_identifier, "X-Role": "member"}
        response = client.get("/projects", headers=headers)
        assert response.status_code == 200
        assert response.json()["items"] == []
        assert response.json()["store"] == {
            "ready": False,
            "backend": "unavailable",
            "persistent_writes_ready": False,
            "history_review_task_writes_ready": False,
        }
        for route in (
            "/projects/SELF-CHECK-FUND-20260607",
            "/projects/SELF-CHECK-FUND-20260607/members",
            "/projects/SELF-CHECK-FUND-20260607/dashboard",
        ):
            assert client.get(route, headers=headers).status_code == 404

    admin_response = client.get(
        "/projects",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
    )
    assert admin_response.status_code == 200
    assert len(admin_response.json()["items"]) == 4
    assert admin_response.json()["store"] == {
        "ready": True,
        "backend": "SplitFailureProjectMemberStore",
        "persistent_writes_ready": False,
        "history_review_task_writes_ready": False,
    }


def test_projects_api_admin_visibility_helper_does_not_read_store() -> None:
    class ExplodingProjectMemberStore:
        def list_members(self, project_key: str) -> list[dict[str, object]]:
            raise AssertionError("admin visibility must not read member store")

        def add_member(
            self,
            project_key: str,
            values: dict[str, object],
        ) -> dict[str, object]:
            raise AssertionError("admin visibility must not write member store")

        def member_counts(self) -> dict[str, int]:
            raise AssertionError("admin visibility must not count member store")

    visible_keys = visible_project_keys(
        user_identifier="admin-1",
        is_admin=True,
        store=ExplodingProjectMemberStore(),
    )
    assert isinstance(visible_keys, frozenset)
    assert visible_keys == {
        "SELF-CHECK-FUND-20260607",
        "CATALOG-LIMIT-202606",
        "OUTPATIENT-DOSE-202606",
        "KB-GOVERNANCE-202606",
    }


def test_projects_api_rejects_unknown_project_and_role(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.project_member_store = SqlAlchemyProjectMemberStore(
        f"sqlite:///{tmp_path / 'project-members.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))

    member_headers = {"X-User-Id": "unrelated-member", "X-Role": "member"}
    missing_detail_response = client.get("/projects/UNKNOWN", headers=member_headers)
    missing_members_response = client.get("/projects/UNKNOWN/members", headers=member_headers)
    invalid_role_response = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        json={
            "user_identifier": "invalid-role-member",
            "name": "未知角色",
            "role": "访客",
            "department": "医保办",
        },
    )
    missing_identifier_response = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
        json={
            "name": "缺少身份",
            "role": "审计员",
            "department": "医保办",
        },
    )
    whitespace_identifier_response = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers={"X-User-Id": "admin-1", "X-Role": "admin"},
        json={
            "user_identifier": "   ",
            "name": "空白身份",
            "role": "审计员",
            "department": "医保办",
        },
    )

    assert missing_detail_response.status_code == 404
    assert missing_members_response.status_code == 404
    assert invalid_role_response.status_code == 422
    assert missing_identifier_response.status_code == 422
    assert whitespace_identifier_response.status_code == 422


def test_projects_api_enforces_member_management_permission(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.project_member_store = SqlAlchemyProjectMemberStore(
        f"sqlite:///{tmp_path / 'project-members.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    payload = {
        "user_identifier": "permission-test-member",
        "name": "赵审计",
        "role": "审计员",
        "department": "医保办",
    }

    unauthenticated_response = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        json=payload,
    )
    director_response = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers={"X-User-Id": "director-1", "X-Role": "director"},
        json=payload,
    )
    member_response = client.post(
        "/projects/SELF-CHECK-FUND-20260607/members",
        headers={"X-User-Id": "member-1", "X-Role": "member"},
        json=payload,
    )

    assert unauthenticated_response.status_code == 401
    assert director_response.status_code == 403
    assert member_response.status_code == 403
    assert state.operation_logs[-1]["action"] == "authorization-denied"
    assert state.operation_logs[-1]["payload"]["attempted_action"] == "project-member-create"
    assert state.operation_logs[-1]["payload"]["permission"] == "manage_project_members"


def test_analytics_table_upload_profiles_csv_file(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'analytics-uploads.db'}"
    state = _api_state(tmp_path)
    state.analytics_upload_store = SqlAlchemyAnalyticsUploadStore(
        database_url=database_url,
        upload_root=tmp_path / "retained-uploads",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    headers = {"X-User-Id": "  analytics-owner  ", "X-Role": "member"}

    response = client.post(
        "/analytics/table-upload",
        headers=headers,
        files={
            "file": (
                "charge-sample.csv",
                "\n".join(
                    [
                        "patient_id,visit_date,item_code,charge_amount,insurance_pay",
                        "P001,2026-01-01,A100,120.00,80.00",
                        "P001,2026-01-01,A100,120.00,80.00",
                        "P002,2026-01-02,B200,,50.00",
                    ]
                ),
                "text/csv",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["name"] == "charge-sample.csv"
    assert body["upload_id"].startswith(ANALYTICS_UPLOAD_ID_PREFIX)
    assert body["sha256"]
    assert body["retention_status"] == "retained"
    assert body["created_at"]
    assert body["extension"] == "csv"
    assert body["status"] == "parsed"
    assert body["row_count"] == 3
    assert body["duplicate_row_count"] == 1
    assert body["empty_cell_count"] == 1
    assert body["audit_signals"] == [
        "金额/费用字段",
        "患者/就诊字段",
        "日期/时间字段",
        "项目/药品/目录字段",
        "医保支付字段",
    ]
    assert body["columns"][0]["type"] == "标识"
    assert "发现 1 条完全重复行。" in body["quality_findings"]
    assert state.operation_logs[-1]["action"] == "analytics-table-upload"
    assert state.operation_logs[-1]["payload"]["retention_status"] == "retained"
    assert state.operation_logs[-1]["payload"]["actor"] == "analytics-owner"

    history_response = client.get("/analytics/table-uploads", headers=headers)
    assert history_response.status_code == 200
    history_body = history_response.json()
    assert history_body["store"]["backend"] == "SqlAlchemyAnalyticsUploadStore"
    assert history_body["items"][0]["id"] == body["upload_id"]
    assert history_body["items"][0]["sha256"] == body["sha256"]
    assert history_body["items"][0]["row_count"] == 3
    assert history_body["items"][0]["column_count"] == 5
    assert "storage_path" not in history_body["items"][0]

    retained_record = state.analytics_upload_store.list_uploads(created_by="analytics-owner")[0]
    retained_path = tmp_path / "retained-uploads" / str(retained_record["storage_path"])
    assert retained_path.exists()
    assert retained_path.read_text(encoding="utf-8").startswith("patient_id,visit_date")

    second_state = _api_state(tmp_path / "second")
    second_state.analytics_upload_store = SqlAlchemyAnalyticsUploadStore(
        database_url=database_url,
        upload_root=tmp_path / "retained-uploads",
    )
    second_client = TestClient(create_app(second_state))
    persisted_items = second_client.get(
        "/analytics/table-uploads",
        headers=headers,
    ).json()["items"]
    assert persisted_items[0]["id"] == body["upload_id"]


def test_analytics_table_upload_profiles_xlsx_file(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "收费明细"
    worksheet.append(["患者编号", "结算日期", "项目编码", "收费金额"])
    worksheet.append(["P001", "2026-01-01", "A100", 120])
    worksheet.append(["P002", "2026-01-02", "B200", 80])
    buffer = BytesIO()
    workbook.save(buffer)

    response = client.post(
        "/analytics/table-upload",
        headers={"X-User-Id": "analytics-owner", "X-Role": "member"},
        files={
            "file": (
                "charge-workbook.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["name"] == "charge-workbook.xlsx"
    assert body["extension"] == "xlsx"
    assert body["sheet_name"] == "收费明细"
    assert body["row_count"] == 2
    assert body["columns"][3]["name"] == "收费金额"
    assert body["columns"][3]["type"] == "数值"
    assert body["message"] == "后端已完成 XLSX 工作簿（sheet: 收费明细）的字段画像。"


def test_analytics_table_upload_rejects_unsupported_extension(tmp_path: Path) -> None:
    client = TestClient(create_app(_api_state(tmp_path)))

    response = client.post(
        "/analytics/table-upload",
        headers={"X-User-Id": "analytics-owner", "X-Role": "member"},
        files={"file": ("charges.txt", "not,a,supported,file", "text/plain")},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "unsupported table file extension"


@pytest.mark.parametrize(
    "headers",
    (
        {},
        {"X-User-Id": "   ", "X-Role": "member"},
        {"X-User-Id": "anonymous", "X-Role": "member"},
    ),
)
def test_analytics_table_upload_rejects_missing_identity_before_retention(
    tmp_path: Path,
    headers: dict[str, str],
) -> None:
    upload_root = tmp_path / "retained-uploads"
    state = _api_state(tmp_path)
    state.analytics_upload_store = InMemoryAnalyticsUploadStore(upload_root=upload_root)
    client = TestClient(create_app(state))

    response = client.post(
        "/analytics/table-upload",
        headers=headers,
        files={"file": ("charges.csv", "item,amount\nA,10", "text/csv")},
    )

    assert response.status_code == 401
    assert response.json()["detail"] == "X-User-Id header is required"
    assert state.analytics_upload_store.list_uploads() == []
    assert list(upload_root.glob("**/*")) == []


def test_analytics_upload_history_is_owner_scoped_and_redacts_storage_path(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.analytics_upload_store = InMemoryAnalyticsUploadStore(
        upload_root=tmp_path / "retained-uploads"
    )
    client = TestClient(create_app(state))

    for user_identifier, file_name in (
        ("analytics-owner-a", "owner-a.csv"),
        ("analytics-owner-b", "owner-b.csv"),
    ):
        response = client.post(
            "/analytics/table-upload",
            headers={"X-User-Id": user_identifier, "X-Role": "member"},
            files={"file": (file_name, "item,amount\nA,10", "text/csv")},
        )
        assert response.status_code == 200

    missing_identity = client.get("/analytics/table-uploads")
    owner_history = client.get(
        "/analytics/table-uploads",
        headers={"X-User-Id": "analytics-owner-a", "X-Role": "member"},
    )
    admin_history = client.get(
        "/analytics/table-uploads",
        headers={"X-User-Id": "analytics-admin", "X-Role": "admin"},
    )

    assert missing_identity.status_code == 401
    assert owner_history.status_code == 200
    assert [item["name"] for item in owner_history.json()["items"]] == ["owner-a.csv"]
    assert admin_history.status_code == 200
    assert {item["name"] for item in admin_history.json()["items"]} == {
        "owner-a.csv",
        "owner-b.csv",
    }
    assert "storage_path" not in json.dumps(owner_history.json())
    assert "storage_path" not in json.dumps(admin_history.json())


def test_sql_analytics_upload_removes_retained_file_when_database_write_fails(
    tmp_path: Path,
) -> None:
    upload_root = tmp_path / "retained-uploads"
    store = SqlAlchemyAnalyticsUploadStore(
        database_url=f"sqlite:///{tmp_path / 'analytics-write-failure.db'}",
        upload_root=upload_root,
        create_schema=True,
    )
    AnalyticsUploadRecord.__table__.drop(store._engine)

    with pytest.raises(SQLAlchemyError):
        store.add_upload(
            file_name="must-not-remain.csv",
            extension="csv",
            content=b"item,amount\nA,10",
            analysis_summary={"status": "parsed"},
            created_by="analytics-owner-a",
        )

    assert list(upload_root.rglob("*.csv")) == []


def test_documents_governance_status_is_redacted_get_only(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={
            "document_storage": DocumentStorageSettings(
                provider="tencent-cos",
                cos_bucket="medical-audit-prod",
                cos_region="ap-guangzhou",
                cos_prefix="personal-materials/prod",
                cos_secret_id_env="COS_SECRET_ID",
                cos_secret_key_env="COS_SECRET_KEY",
                cos_sdk_bootstrap_enabled=True,
                signed_url_ttl_seconds=180,
                object_retention_days=365,
                local_quarantine_retention_days=14,
                record_storage_objects=True,
            ),
            "document_upload_governance": DocumentUploadGovernanceSettings(
                virus_scan_provider="tencent-ci-virus",
                virus_scan_job_endpoint_env="VIRUS_SCAN_ENDPOINT",
                virus_scan_job_secret_env="VIRUS_SCAN_SECRET",
                dlp_review_provider="external-dlp",
                dlp_review_job_endpoint_env="DLP_REVIEW_ENDPOINT",
                dlp_review_job_secret_env="DLP_REVIEW_SECRET",
                redaction_rewrite_enabled=True,
                redaction_policy_version="redaction-v2026",
                redaction_manual_review_required=True,
                governance_audit_event_required=True,
            ),
        }
    )

    def fake_schema_ready(database_url: str) -> bool:
        assert database_url == state.settings.database_url
        return True

    monkeypatch.setenv("COS_SECRET_ID", "cos-id-secret-sentinel")
    monkeypatch.setenv("COS_SECRET_KEY", "cos-key-secret-sentinel")
    monkeypatch.setenv("VIRUS_SCAN_SECRET", "virus-secret-sentinel")
    monkeypatch.setenv("DLP_REVIEW_SECRET", "dlp-secret-sentinel")
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_documents.document_storage_objects_schema_ready",
        fake_schema_ready,
    )
    client = TestClient(create_app(state))

    response = client.get("/documents/governance/status")
    prefixed_response = client.get("/api/v1/documents/governance/status")

    assert response.status_code == 200
    assert prefixed_response.status_code == 200
    body = response.json()
    serialized = json.dumps(body, ensure_ascii=False)
    fields = body["required_report_fields"]

    assert body["status"] == "readonly_status_available"
    assert body["evidence_grade"] == "L1-public-or-runtime"
    assert body["storage"]["provider"] == "tencent-cos"
    assert body["storage"]["cos_bucket_status"] == "set"
    assert body["storage"]["cos_region_status"] == "set"
    assert body["storage"]["cos_prefix_status"] == "set"
    assert body["storage"]["cos_secret_id"] == {
        "env_name_status": "set",
        "referenced_secret_status": "set",
    }
    assert body["governance"]["virus_scan"]["provider"] == "tencent-ci-virus"
    assert body["governance"]["dlp_review"]["provider"] == "external-dlp"
    assert fields["document_storage_provider"] == "tencent-cos"
    assert fields["document_storage_objects_schema_ready"] is True
    assert fields["document_upload_list_readonly_status"] == ("blocked_by_audit_log_side_effect")
    assert fields["download_metadata_readonly_status"] == ("blocked_by_audit_log_side_effect")
    assert fields["audit_log_readonly_status"] == "available_no_event_written"
    assert body["boundaries"] == {
        "production_write": False,
        "document_upload_write": False,
        "document_upload_list_api_called": False,
        "download_metadata_api_called": False,
        "audit_log_write_expected": False,
        "provider_call": False,
        "object_storage_write": False,
        "secret_values_reported": False,
        "allowed_http_methods": ["GET"],
        "non_get_http_methods_allowed": False,
    }
    assert state.operation_logs == []
    for hidden in (
        "medical-audit-prod",
        "ap-guangzhou",
        "personal-materials/prod",
        "COS_SECRET_ID",
        "COS_SECRET_KEY",
        "VIRUS_SCAN_ENDPOINT",
        "VIRUS_SCAN_SECRET",
        "DLP_REVIEW_ENDPOINT",
        "DLP_REVIEW_SECRET",
        "redaction-v2026",
        "cos-id-secret-sentinel",
        "cos-key-secret-sentinel",
        "virus-secret-sentinel",
        "dlp-secret-sentinel",
    ):
        assert hidden not in serialized


def test_knowledge_base_catalog_reports_index_layers_without_writes(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.search_backend = "postgres"
    state.search_backend_details = {
        "embedding_provider": "openai",
        "embedding_model": "kimi-for-coding",
        "matching_embedding_count": 49051,
        "collection_metrics": {
            "medical-insurance-laws": {
                "document_count": 503,
                "chunk_count": 49051,
                "embedding_count": 49051,
                "active_embedding_count": 49051,
                "candidate_chunk_count": 727214,
                "latest_index_version_key": "incremental-20260615",
                "latest_index_status": "active",
            }
        },
        "postgres_totals": {
            "source_documents": 20054,
            "document_chunks": 923288,
            "chunk_embeddings": 923288,
        },
    }
    client = TestClient(create_app(state))

    response = client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"})

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["contract_version"] == "knowledge-base-catalog-v1"
    assert body["summary"]["source_collection_count"] == 25
    assert body["summary"]["total_document_count"] == 20054
    assert body["summary"]["total_chunk_count"] == 923288
    assert body["summary"]["current_search_embedding_count"] == 49051
    assert body["summary"]["candidate_chunk_count"] == 727214
    assert body["boundaries"]["production_write"] is False
    assert body["boundaries"]["database_write"] is False
    assert body["boundaries"]["source"] == "runtime_state_and_postgres_catalog"
    assert body["store"] == {
        "ready": True,
        "catalog_ready": True,
        "metrics_ready": True,
        "backend": "runtime_state_and_postgres_catalog",
    }
    law_item = next(
        item for item in body["items"] if item["source_collection"] == "medical-insurance-laws"
    )
    assert law_item["metrics"]["document_count"] == 503
    assert law_item["metrics"]["active_embedding_count"] == 49051
    assert law_item["metrics"]["candidate_chunk_count"] == 727214
    assert law_item["index"]["latest_status"] == "active"
    assert state.operation_logs == []


def test_knowledge_base_catalog_controlled_success_writes_no_audit_event(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    persisted_events: list[tuple[str, dict[str, object]]] = []

    class AuditSpy:
        def add_event(self, action: str, payload: dict[str, object]) -> dict[str, object]:
            persisted_events.append((action, payload))
            return {"action": action, "payload": payload}

    class SearchEngineThatMustNotRun:
        def search(self, *_args: object, **_kwargs: object) -> object:
            raise AssertionError("catalog must not invoke the search engine or provider")

    state.audit_log_store = AuditSpy()  # type: ignore[assignment]
    state.search_engine = SearchEngineThatMustNotRun()  # type: ignore[assignment]
    client = TestClient(create_app(state, enforce_controlled_api_auth=True))

    response = client.get(
        "/api/v1/knowledge-base/catalog",
        headers={
            "X-User-Id": "deployment-state-auditor-contract-test",
            "X-Role": "it-admin",
            "X-Tenant-Id": "hospital-demo",
        },
    )

    assert response.status_code == 200, response.text
    assert response.json()["boundaries"]["database_write"] is False
    assert response.json()["boundaries"]["provider_call"] is False
    assert state.operation_logs == []
    assert persisted_events == []


def test_search_backend_status_records_audit_event_even_when_reading(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    persisted_events: list[tuple[str, dict[str, object]]] = []
    acceptance_user_id = "frontend-acceptance-fa-20260720t003332z-c030f1ad"

    class AuditSpy:
        def add_event(self, action: str, payload: dict[str, object]) -> dict[str, object]:
            persisted_events.append((action, payload))
            return {"action": action, "payload": payload}

    state.audit_log_store = AuditSpy()  # type: ignore[assignment]
    client = TestClient(create_app(state, enforce_controlled_api_auth=True))

    response = client.get(
        "/api/v1/index/search-backend",
        headers={
            "X-User-Id": acceptance_user_id,
            "X-Role": "it-admin",
            "X-Tenant-Id": "hospital-demo",
        },
    )

    assert response.status_code == 200, response.text
    assert state.operation_logs[-1]["action"] == "search-backend-status-view"
    assert persisted_events[-1][0] == "search-backend-status-view"
    assert persisted_events[-1][1]["user_identifier"] == acceptance_user_id
    assert persisted_events[-1][1]["role"] == "admin"


def test_knowledge_base_catalog_marks_registry_only_metrics_unready(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    client = TestClient(create_app(state))

    response = client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"})

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["boundaries"]["source"] == "runtime_state_and_registry_only"
    assert body["store"] == {
        "ready": False,
        "catalog_ready": True,
        "metrics_ready": False,
        "backend": "runtime_state_and_registry_only",
    }


def test_document_source_and_knowledge_base_catalog_scrub_sensitive_diagnostics(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    state.search_backend_details = {
        "embedding_provider": "deterministic-fake",
        "token": "dummy-top-token-sentinel",
        "password": "dummy-top-password-sentinel",
        "api_key": "dummy-top-api-key-sentinel",
        "private_key": "dummy-top-private-key-sentinel",
        "credential": "dummy-top-credential-sentinel",
        "secret": "dummy-top-secret-sentinel",
        "nested_mapping": {
            "safe_value": "dummy-nested-safe-sentinel",
            "token": "dummy-nested-token-sentinel",
            "password": "dummy-nested-password-sentinel",
            "api_key": "dummy-nested-api-key-sentinel",
            "private_key": "dummy-nested-private-key-sentinel",
            "credential": "dummy-nested-credential-sentinel",
            "secret": "dummy-nested-secret-sentinel",
        },
        "nested_list": [
            {
                "safe_value": 7,
                "token": "dummy-list-token-sentinel",
                "password": "dummy-list-password-sentinel",
                "api_key": "dummy-list-api-key-sentinel",
                "private_key": "dummy-list-private-key-sentinel",
                "credential": "dummy-list-credential-sentinel",
                "secret": "dummy-list-secret-sentinel",
            },
            ("dummy-tuple-safe-sentinel", {"secret": "dummy-tuple-secret-sentinel"}),
        ],
        "endpoint_url": (
            "https://dummy-user-sentinel:dummy-password-sentinel@example.invalid/search"
            "?safe=one&token=dummy-query-token-sentinel&safe=two"
            "&password=dummy-query-password-sentinel&blank="
            "#token=dummy-fragment-token-sentinel"
        ),
        "ipv6_url": (
            "https://dummy-ipv6-user-sentinel:dummy-ipv6-password-sentinel@"
            "[2001:db8::1]:8443/path"
            "?api_key=dummy-ipv6-api-key-sentinel&safe=value"
        ),
        "invalid_port_url": (
            "https://dummy-invalid-user-sentinel:dummy-invalid-password-sentinel@"
            "example.invalid:notaport/path?credential=dummy-invalid-query-sentinel"
        ),
        "malformed_bracket_url": (
            "https://dummy-bracket-user-sentinel:dummy-bracket-password-sentinel@"
            "[2001:db8::1/path?secret=dummy-bracket-query-sentinel"
        ),
    }
    client = TestClient(create_app(state))

    responses = (
        client.get("/documents/source-collections", headers={"X-Role": "it-admin"}),
        client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"}),
    )

    for response in responses:
        assert response.status_code == 200, response.text
        details = response.json()["search_backend"]["details"]
        assert details == {
            "embedding_provider": "deterministic-fake",
            "nested_mapping": {"safe_value": "dummy-nested-safe-sentinel"},
            "nested_list": [
                {"safe_value": 7},
                ["dummy-tuple-safe-sentinel", {}],
            ],
            "endpoint_url": "https://example.invalid/search?safe=one&safe=two&blank=",
            "ipv6_url": "https://[2001:db8::1]:8443/path?safe=value",
            "invalid_port_url": "<redacted-invalid-url>",
            "malformed_bracket_url": "<redacted-invalid-url>",
        }
        serialized = json.dumps(response.json(), ensure_ascii=False)
        for sensitive_key in (
            "token",
            "password",
            "api_key",
            "private_key",
            "credential",
            "secret",
        ):
            assert sensitive_key not in serialized.lower()
        for sensitive_value in (
            "dummy-top-token-sentinel",
            "dummy-top-password-sentinel",
            "dummy-top-api-key-sentinel",
            "dummy-top-private-key-sentinel",
            "dummy-top-credential-sentinel",
            "dummy-top-secret-sentinel",
            "dummy-nested-token-sentinel",
            "dummy-nested-password-sentinel",
            "dummy-nested-api-key-sentinel",
            "dummy-nested-private-key-sentinel",
            "dummy-nested-credential-sentinel",
            "dummy-nested-secret-sentinel",
            "dummy-list-token-sentinel",
            "dummy-list-password-sentinel",
            "dummy-list-api-key-sentinel",
            "dummy-list-private-key-sentinel",
            "dummy-list-credential-sentinel",
            "dummy-list-secret-sentinel",
            "dummy-tuple-secret-sentinel",
            "dummy-user-sentinel",
            "dummy-password-sentinel",
            "dummy-query-token-sentinel",
            "dummy-query-password-sentinel",
            "dummy-fragment-token-sentinel",
            "dummy-ipv6-user-sentinel",
            "dummy-ipv6-password-sentinel",
            "dummy-ipv6-api-key-sentinel",
            "dummy-invalid-user-sentinel",
            "dummy-invalid-password-sentinel",
            "dummy-invalid-query-sentinel",
            "dummy-bracket-user-sentinel",
            "dummy-bracket-password-sentinel",
            "dummy-bracket-query-sentinel",
        ):
            assert sensitive_value not in serialized


def test_document_source_and_knowledge_base_catalog_replace_unsupported_diagnostics(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    state.search_backend_details = {
        "nested": {"unsupported": object()},
    }
    client = TestClient(create_app(state))

    responses = (
        client.get("/documents/source-collections", headers={"X-Role": "it-admin"}),
        client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"}),
    )

    for response in responses:
        assert response.status_code == 200, response.text
        assert response.json()["search_backend"]["details"] == {
            "nested": {"unsupported": "<unsupported-diagnostic-value>"}
        }


def test_document_source_and_knowledge_base_catalog_scrub_protocol_relative_urls(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    state.search_backend_details = {
        "protocol_relative_url": (
            "//dummy-user-sentinel:dummy-pass-sentinel@example.invalid/path"
            "?safe=ok&token=dummy-query-token-sentinel"
            "#password=dummy-fragment-password-sentinel"
        ),
        "empty_host_url": (
            "https://?token=dummy-empty-token-sentinel#password=dummy-empty-password-sentinel"
        ),
        "missing_protocol_relative_host_url": "//",
        "plain_diagnostic": "plain diagnostic dummy-safe-sentinel",
    }
    client = TestClient(create_app(state))

    responses = (
        client.get("/documents/source-collections", headers={"X-Role": "it-admin"}),
        client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"}),
    )

    for response in responses:
        assert response.status_code == 200, response.text
        assert response.json()["search_backend"]["details"] == {
            "protocol_relative_url": "//example.invalid/path?safe=ok",
            "empty_host_url": "<redacted-invalid-url>",
            "missing_protocol_relative_host_url": "<redacted-invalid-url>",
            "plain_diagnostic": "plain diagnostic dummy-safe-sentinel",
        }
        serialized = json.dumps(response.json(), ensure_ascii=False)
        for sensitive_value in (
            "dummy-user-sentinel",
            "dummy-pass-sentinel",
            "dummy-query-token-sentinel",
            "dummy-fragment-password-sentinel",
            "dummy-empty-token-sentinel",
            "dummy-empty-password-sentinel",
        ):
            assert sensitive_value not in serialized


def test_document_source_and_knowledge_base_catalog_scrub_double_encoded_url_parts(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    state.search_backend_details = {
        "encoded_url": (
            "https://example.invalid/path?safe=ok"
            "&to%256ben=dummy-double-token-sentinel"
            "&Pa%2573sword=dummy-double-password-sentinel"
            "#Pa%2573sword=dummy-double-fragment-sentinel"
        ),
    }
    client = TestClient(create_app(state))

    responses = (
        client.get("/documents/source-collections", headers={"X-Role": "it-admin"}),
        client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"}),
    )

    for response in responses:
        assert response.status_code == 200, response.text
        assert response.json()["search_backend"]["details"] == {
            "encoded_url": "https://example.invalid/path?safe=ok"
        }
        serialized = json.dumps(response.json(), ensure_ascii=False)
        for sensitive_value in (
            "dummy-double-token-sentinel",
            "dummy-double-password-sentinel",
            "dummy-double-fragment-sentinel",
        ):
            assert sensitive_value not in serialized


def test_document_source_and_knowledge_base_catalog_replace_nonfinite_diagnostics(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    state.search_backend_details = {
        "nested": {
            "nan": float("nan"),
            "positive_infinity": float("inf"),
            "negative_infinity": float("-inf"),
        }
    }
    client = TestClient(create_app(state))

    responses = (
        client.get("/documents/source-collections", headers={"X-Role": "it-admin"}),
        client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"}),
    )

    for response in responses:
        assert response.status_code == 200, response.text
        assert response.json()["search_backend"]["details"] == {
            "nested": {
                "nan": "<unsupported-diagnostic-value>",
                "positive_infinity": "<unsupported-diagnostic-value>",
                "negative_infinity": "<unsupported-diagnostic-value>",
            }
        }


def test_document_source_and_knowledge_base_catalog_normalize_prefixed_url_candidates(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    state.search_backend_details = {
        "prefixed_empty_host_url": (
            " \x00https://?token=dummy-prefixed-empty-token-sentinel"
            "#password=dummy-prefixed-empty-password-sentinel"
        ),
        "prefixed_protocol_relative_url": (
            " \t//dummy-prefixed-user-sentinel:dummy-prefixed-pass-sentinel@"
            "example.invalid/path?safe=ok&token=dummy-prefixed-query-token-sentinel"
            "#password=dummy-prefixed-fragment-password-sentinel"
        ),
    }
    client = TestClient(create_app(state))

    responses = (
        client.get("/documents/source-collections", headers={"X-Role": "it-admin"}),
        client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"}),
    )

    for response in responses:
        assert response.status_code == 200, response.text
        assert response.json()["search_backend"]["details"] == {
            "prefixed_empty_host_url": "<redacted-invalid-url>",
            "prefixed_protocol_relative_url": "//example.invalid/path?safe=ok",
        }
        serialized = json.dumps(response.json(), ensure_ascii=False)
        for sensitive_value in (
            "dummy-prefixed-empty-token-sentinel",
            "dummy-prefixed-empty-password-sentinel",
            "dummy-prefixed-user-sentinel",
            "dummy-prefixed-pass-sentinel",
            "dummy-prefixed-query-token-sentinel",
            "dummy-prefixed-fragment-password-sentinel",
        ):
            assert sensitive_value not in serialized


def test_document_source_and_knowledge_base_catalog_fail_closed_beyond_decode_cap(
    tmp_path: Path,
) -> None:
    overencoded_token_key = "To%4beN"
    overencoded_password_fragment = "Pa%73sWoRd"
    for _ in range(5):
        overencoded_token_key = overencoded_token_key.replace("%", "%25")
        overencoded_password_fragment = overencoded_password_fragment.replace("%", "%25")

    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    state.search_backend_details = {
        "overencoded_url": (
            "https://example.invalid/path?safe=ok"
            f"&{overencoded_token_key}=dummy-overencoded-query-sentinel"
            f"#{overencoded_password_fragment}=dummy-overencoded-fragment-sentinel"
        ),
        "double_encoded_url": (
            "https://example.invalid/path?safe=ok"
            "&to%256ben=dummy-double-encoded-query-sentinel"
            "#Pa%2573sword=dummy-double-encoded-fragment-sentinel"
        ),
        "plain_diagnostic": "plain diagnostic dummy-safe-sentinel",
    }
    client = TestClient(create_app(state))

    responses = (
        client.get("/documents/source-collections", headers={"X-Role": "it-admin"}),
        client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"}),
    )

    for response in responses:
        assert response.status_code == 200, response.text
        assert response.json()["search_backend"]["details"] == {
            "overencoded_url": "https://example.invalid/path?safe=ok",
            "double_encoded_url": "https://example.invalid/path?safe=ok",
            "plain_diagnostic": "plain diagnostic dummy-safe-sentinel",
        }
        serialized = json.dumps(response.json(), ensure_ascii=False)
        for sensitive_value in (
            "dummy-overencoded-query-sentinel",
            "dummy-overencoded-fragment-sentinel",
            "dummy-double-encoded-query-sentinel",
            "dummy-double-encoded-fragment-sentinel",
        ):
            assert sensitive_value not in serialized


def test_document_source_and_knowledge_base_catalog_normalize_embedded_url_controls(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.settings = state.settings.model_copy(
        update={"database_url": "registry-only://dummy-diagnostic-sentinel"}
    )
    state.search_backend_details = {
        "embedded_tab_url": (
            "h\tttps://?token=dummy-embedded-tab-token-sentinel"
            "#password=dummy-embedded-tab-password-sentinel"
        ),
        "embedded_carriage_return_url": (
            "ht\rtps://?token=dummy-embedded-cr-token-sentinel"
            "#password=dummy-embedded-cr-password-sentinel"
        ),
        "embedded_line_feed_url": (
            "htt\nps://?token=dummy-embedded-lf-token-sentinel"
            "#password=dummy-embedded-lf-password-sentinel"
        ),
    }
    client = TestClient(create_app(state))

    responses = (
        client.get("/documents/source-collections", headers={"X-Role": "it-admin"}),
        client.get("/knowledge-base/catalog", headers={"X-Role": "it-admin"}),
    )

    for response in responses:
        assert response.status_code == 200, response.text
        assert response.json()["search_backend"]["details"] == {
            "embedded_tab_url": "<redacted-invalid-url>",
            "embedded_carriage_return_url": "<redacted-invalid-url>",
            "embedded_line_feed_url": "<redacted-invalid-url>",
        }
        serialized = json.dumps(response.json(), ensure_ascii=False)
        for sensitive_value in (
            "dummy-embedded-tab-token-sentinel",
            "dummy-embedded-tab-password-sentinel",
            "dummy-embedded-cr-token-sentinel",
            "dummy-embedded-cr-password-sentinel",
            "dummy-embedded-lf-token-sentinel",
            "dummy-embedded-lf-password-sentinel",
        ):
            assert sensitive_value not in serialized


def test_documents_search_is_readonly_and_scoped_to_source_collection(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.get(
        "/documents/search",
        headers={"X-Role": "it-admin", "X-User-Id": "auditor-1"},
        params={
            "q": "医保基金审核依据",
            "source_collection": SourceCollection.MEDICAL_INSURANCE_LAWS.value,
            "limit": 3,
        },
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["contract_version"] == "document-search-v2"
    assert body["query"] == "医保基金审核依据"
    assert body["effective_source_collections"] == [SourceCollection.MEDICAL_INSURANCE_LAWS.value]
    assert body["boundaries"]["query_history_write"] is False
    assert body["items"][0]["source_collection"] == SourceCollection.MEDICAL_INSURANCE_LAWS.value
    assert body["items"][0]["preview_url"].startswith("/api/v1/preview/")
    assert body["items"][0]["download_url"].startswith("/api/v1/documents/source/")
    assert body["items"][0]["match_count"] == 1
    assert body["items"][0]["matched_snippets"]
    download_response = client.get(
        body["items"][0]["download_url"],
        headers={"X-Role": "it-admin", "X-User-Id": "auditor-1"},
    )
    assert download_response.status_code == 200
    assert download_response.headers["content-disposition"].startswith("attachment;")
    assert "医疗机构应当保留医保基金审核依据" in download_response.text
    assert state.query_history_store.list_queries() == []


def test_personal_document_source_download_is_owner_scoped(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    personal_source_path = "personal-materials/document-upload-private/note.txt"
    _write_text(
        state.source_root / personal_source_path,
        "院内个人材料提示：医保基金审核依据需要结合内部台账。",  # noqa: RUF001
    )
    state.search_engine = _search_engine_with_personal_materials(
        system_chunk_id=uuid4(),
        personal_chunk_id=uuid4(),
        source_path=personal_source_path,
    )
    client = TestClient(create_app(state))

    search_response = client.get(
        "/documents/search",
        headers={"X-Role": "auditor", "X-User-Id": "auditor-1"},
        params={
            "q": "院内个人材料提示",
            "source_collection": SourceCollection.PERSONAL_MATERIALS.value,
            "limit": 3,
        },
    )

    assert search_response.status_code == 200, search_response.text
    chunk_id = search_response.json()["items"][0]["chunk_id"]
    download_url = f"/api/v1/documents/source/{chunk_id}/download"
    assert (
        client.get(
            download_url,
            headers={"X-Role": "auditor", "X-User-Id": "auditor-1"},
        ).status_code
        == 200
    )
    assert (
        client.get(
            download_url,
            headers={"X-Role": "auditor", "X-User-Id": "auditor-2"},
        ).status_code
        == 404
    )
    assert (
        client.get(
            download_url,
            headers={"X-Role": "it-admin", "X-User-Id": "admin-1"},
        ).status_code
        == 200
    )


def test_documents_permissions_and_uploads_are_role_scoped(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'document-uploads.db'}"
    upload_root = tmp_path / "document-uploads"
    state = _api_state(tmp_path)
    state.document_upload_store = SqlAlchemyDocumentUploadStore(
        database_url=database_url,
        upload_root=upload_root,
        create_schema=True,
    )
    client = TestClient(create_app(state))

    permissions_response = client.get("/documents/permissions", headers={"X-Role": "auditor"})

    assert permissions_response.status_code == 200
    permissions_body = permissions_response.json()
    assert permissions_body["role"] == "auditor"
    assert [item["source_collection"] for item in permissions_body["source_collections"]] == [
        *[definition.collection.value for definition in SYSTEM_SOURCE_COLLECTION_DEFINITIONS],
        SourceCollection.PERSONAL_MATERIALS.value,
    ]
    assert permissions_body["source_collections"][-1]["access"] == "explicit-owner-read"
    assert permissions_body["upload_permissions"] == {
        "can_upload_personal": True,
        "can_read_all_personal_uploads": False,
        "can_govern_personal_uploads": False,
    }

    upload_response = client.post(
        "/documents/uploads",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        files={"file": ("policy.pdf", b"%PDF-1.4 policy", "application/pdf")},
    )

    assert upload_response.status_code == 200
    upload_body = upload_response.json()
    uploaded = upload_body["item"]
    assert uploaded["id"].startswith(DOCUMENT_UPLOAD_ID_PREFIX)
    assert uploaded["created_by"] == "auditor-1"
    assert uploaded["extension"] == "pdf"
    assert uploaded["retention_status"] == "retained"
    assert uploaded["index_status"] == "not-indexed"
    assert uploaded["governance_status"] == "pending-review"
    assert uploaded["governance_note"] == ""
    assert uploaded["governed_by"] is None
    assert uploaded["governed_at"] is None
    assert uploaded["security_scan_status"] == "local-policy-passed"
    assert uploaded["security_scan_provider"] == "local-policy"
    assert uploaded["dlp_status"] == "clear"
    assert uploaded["security_findings"] == []
    assert uploaded["download_url"] == f"/api/v1/documents/uploads/{uploaded['id']}/download"
    assert uploaded["sha256"]
    assert upload_body["store"]["backend"] == "SqlAlchemyDocumentUploadStore"
    assert state.operation_logs[-1]["action"] == "document-upload"
    assert state.operation_logs[-1]["payload"]["index_status"] == "not-indexed"
    staged_payload = {**uploaded, "index_status": "staged-for-index"}
    assert DocumentUploadItem.model_validate(staged_payload).index_status == "staged-for-index"

    retained_path = upload_root / uploaded["storage_path"]
    assert retained_path.exists()
    assert retained_path.read_bytes() == b"%PDF-1.4 policy"

    owner_response = client.get(
        "/documents/uploads",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    )
    assert owner_response.status_code == 200
    assert owner_response.json()["items"][0]["id"] == uploaded["id"]

    owner_download_response = client.get(
        f"/documents/uploads/{uploaded['id']}/download",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    )
    assert owner_download_response.status_code == 200
    assert owner_download_response.content == b"%PDF-1.4 policy"
    assert owner_download_response.headers["x-document-upload-id"] == uploaded["id"]
    assert owner_download_response.headers["x-document-security-scan"] == "local-policy-passed"
    assert state.operation_logs[-1]["action"] == "document-upload-download"

    other_auditor_response = client.get(
        "/documents/uploads",
        headers={"X-User-Id": "auditor-2", "X-Role": "auditor"},
    )
    assert other_auditor_response.status_code == 200
    assert other_auditor_response.json()["items"] == []

    other_download_response = client.get(
        f"/documents/uploads/{uploaded['id']}/download",
        headers={"X-User-Id": "auditor-2", "X-Role": "auditor"},
    )
    assert other_download_response.status_code == 404
    assert state.operation_logs[-1]["action"] == "authorization-denied"
    assert state.operation_logs[-1]["payload"]["attempted_action"] == "document-upload-download"

    admin_response = client.get(
        "/documents/uploads",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
    )
    assert admin_response.status_code == 200
    admin_body = admin_response.json()
    assert admin_body["permissions"]["can_read_all_personal_uploads"] is True
    assert admin_body["permissions"]["can_govern_personal_uploads"] is True
    assert admin_body["items"][0]["id"] == uploaded["id"]
    admin_permissions_response = client.get(
        "/documents/permissions",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
    )
    assert admin_permissions_response.status_code == 200
    assert admin_permissions_response.json()["source_collections"][-1]["access"] == (
        "explicit-read-all"
    )

    admin_download_response = client.get(
        f"/documents/uploads/{uploaded['id']}/download",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
    )
    assert admin_download_response.status_code == 200
    assert admin_download_response.content == b"%PDF-1.4 policy"

    member_governance_response = client.post(
        f"/documents/uploads/{uploaded['id']}/governance",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"governance_status": "approved-for-index", "note": "member should not approve"},
    )
    assert member_governance_response.status_code == 403

    governance_response = client.post(
        f"/documents/uploads/{uploaded['id']}/governance",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
        json={"governance_status": "approved-for-index", "note": "已完成材料治理。"},
    )
    assert governance_response.status_code == 200
    governed = governance_response.json()["item"]
    assert governed["governance_status"] == "approved-for-index"
    assert governed["governance_note"] == "已完成材料治理。"
    assert governed["governed_by"] == "admin-1"
    assert governed["governed_at"]
    assert governed["index_status"] == "index-ready"
    assert state.operation_logs[-1]["action"] == "document-upload-governance-update"

    governed_admin_items = client.get(
        "/documents/uploads",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
    ).json()["items"]
    assert governed_admin_items[0]["governance_status"] == "approved-for-index"
    assert governed_admin_items[0]["index_status"] == "index-ready"

    second_state = _api_state(tmp_path / "second")
    second_state.document_upload_store = SqlAlchemyDocumentUploadStore(
        database_url=database_url,
        upload_root=upload_root,
    )
    second_client = TestClient(create_app(second_state))
    persisted_items = second_client.get(
        "/documents/uploads",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    ).json()["items"]
    assert persisted_items[0]["id"] == uploaded["id"]
    assert persisted_items[0]["governance_status"] == "approved-for-index"
    assert persisted_items[0]["index_status"] == "index-ready"


def test_documents_permissions_use_persistent_role_and_status_gate(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    assert state.auth_user_store is not None
    state.auth_user_store.add_user(
        {
            "user_key": "persistent-document-director",
            "display_name": "文档主任",
            "status": "active",
        }
    )
    state.auth_user_store.assign_role(
        "persistent-document-director",
        {"role": "director", "scope_type": "global"},
    )
    client = TestClient(create_app(state))

    permissions_response = client.get(
        "/documents/permissions",
        headers={"X-User-Id": "persistent-document-director", "X-Role": "auditor"},
    )
    disabled_user = state.auth_user_store.update_user(
        "persistent-document-director",
        {"status": "disabled"},
    )
    denied_response = client.get(
        "/documents/permissions",
        headers={"X-User-Id": "persistent-document-director", "X-Role": "director"},
    )

    assert permissions_response.status_code == 200
    permissions_body = permissions_response.json()
    assert permissions_body["role"] == "department-head"
    assert permissions_body["upload_permissions"]["can_read_all_personal_uploads"] is True
    assert permissions_body["upload_permissions"]["can_govern_personal_uploads"] is True
    assert disabled_user["status"] == "disabled"
    assert denied_response.status_code == 403
    assert denied_response.json()["detail"] == "auth user status is disabled"


def test_documents_upload_rejects_unsupported_extension(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.document_upload_store = InMemoryDocumentUploadStore(
        upload_root=tmp_path / "document-uploads",
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/documents/uploads",
        files={"file": ("policy.exe", b"binary", "application/octet-stream")},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "unsupported document file extension"


def test_documents_upload_local_policy_blocks_index_approval_until_review(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.document_upload_store = InMemoryDocumentUploadStore(
        upload_root=tmp_path / "document-uploads",
    )
    client = TestClient(create_app(state))

    upload_response = client.post(
        "/documents/uploads",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        files={"file": ("notes.txt", "password=example", "text/plain")},
    )

    assert upload_response.status_code == 200
    uploaded = upload_response.json()["item"]
    assert uploaded["security_scan_status"] == "local-policy-review"
    assert uploaded["dlp_status"] == "needs-review"
    assert uploaded["security_findings"] == ["sensitive-keyword:credential"]

    governance_response = client.post(
        f"/documents/uploads/{uploaded['id']}/governance",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
        json={"governance_status": "approved-for-index", "note": "review later"},
    )

    assert governance_response.status_code == 409
    assert governance_response.json()["detail"] == (
        "document upload security review is required before index approval"
    )
    assert state.operation_logs[-1]["action"] == "document-upload-governance-blocked"
    assert state.operation_logs[-1]["payload"]["security_finding_count"] == 1


def test_documents_index_readiness_governance_result_and_manual_approval(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.document_upload_store = InMemoryDocumentUploadStore(
        upload_root=tmp_path / "document-uploads",
    )
    client = TestClient(create_app(state))

    upload_response = client.post(
        "/documents/uploads",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        files={"file": ("policy.txt", "controlled evidence", "text/plain")},
    )
    assert upload_response.status_code == 200
    uploaded = upload_response.json()["item"]
    upload_id = uploaded["id"]
    assert uploaded["index_readiness"]["blockers"] == [
        "virus-scan-required",
        "dlp-review-required",
        "manual-index-approval-required",
    ]

    auditor_update = client.post(
        f"/documents/uploads/{upload_id}/index-readiness/governance-result",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={
            "check_type": "dlp-review",
            "provider": "external-dlp",
            "status": "passed",
            "detail": "auditor should not write governance result",
        },
    )
    assert auditor_update.status_code == 403
    assert auditor_update.json()["detail"] == (
        "document upload governance result update requires department-head or system-admin role"
    )
    assert state.operation_logs[-1]["action"] == ("document-upload-governance-result-access-denied")

    virus_response = client.post(
        f"/documents/uploads/{upload_id}/index-readiness/governance-result",
        headers={"X-User-Id": "head-1", "X-Role": "department-head"},
        json={
            "check_type": "virus-scan",
            "provider": "tencent-ci-virus",
            "status": "passed",
            "detail": "controlled virus result writeback",
            "result_code": "normal",
            "external_job_id": "job-virus-1",
            "finished_at": "2026-06-27T15:00:00Z",
        },
    )
    assert virus_response.status_code == 200
    after_virus = virus_response.json()["item"]
    assert after_virus["index_status"] == "not-indexed"
    assert after_virus["index_readiness"]["status"] == "blocked"
    assert "virus-scan-required" not in after_virus["index_readiness"]["blockers"]
    assert state.operation_logs[-1]["action"] == "document-upload-governance-result-update"

    dlp_response = client.post(
        f"/documents/uploads/{upload_id}/index-readiness/governance-result",
        headers={"X-User-Id": "head-1", "X-Role": "department-head"},
        json={
            "check_type": "dlp-review",
            "provider": "external-dlp",
            "status": "passed",
            "detail": "controlled dlp result writeback",
            "result_code": "no-sensitive-marker",
            "external_job_id": "job-dlp-1",
        },
    )
    assert dlp_response.status_code == 200
    after_dlp = dlp_response.json()["item"]
    assert after_dlp["index_readiness"]["status"] == "blocked"
    assert after_dlp["index_readiness"]["blockers"] == ["manual-index-approval-required"]

    manual_response = client.post(
        f"/documents/uploads/{upload_id}/index-readiness/manual-approval",
        headers={"X-User-Id": "head-1", "X-Role": "department-head"},
        json={
            "decision": "approved",
            "note": "controlled manual approval",
        },
    )
    assert manual_response.status_code == 200
    ready = manual_response.json()["item"]
    assert ready["index_status"] == "not-indexed"
    assert ready["personal_index_status"] == "not-indexed"
    assert ready["index_readiness"]["status"] == "ready"
    assert ready["index_readiness"]["blockers"] == []
    assert ready["index_readiness"]["next_action"] == "ingest-personal-upload"
    assert all(check["status"] == "passed" for check in ready["index_readiness"]["checks"])
    assert state.operation_logs[-1]["action"] == "document-upload-index-readiness-update"
    assert state.operation_logs[-1]["payload"]["index_status"] == "not-indexed"


def test_personal_document_index_ingestion_requires_enabled_indexer(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.document_upload_store = InMemoryDocumentUploadStore(
        upload_root=tmp_path / "document-uploads",
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/documents/uploads/document-upload-missing/index-ingestion",
        headers={"X-User-Id": "head-1", "X-Role": "department-head"},
    )

    assert response.status_code == 409
    assert response.json()["detail"] == "document upload indexing is not enabled"
    assert state.operation_logs[-1]["action"] == "document-upload-index-ingestion-blocked"
    assert state.operation_logs[-1]["payload"]["reason"] == "document-upload-indexing-disabled"


def test_personal_document_index_ingestion_stages_candidate(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'document-index-ingestion.db'}"
    upload_root = tmp_path / "document-uploads"
    state = _api_state(tmp_path)
    state.document_upload_store = SqlAlchemyDocumentUploadStore(
        database_url=database_url,
        upload_root=upload_root,
        create_schema=True,
    )
    state.document_upload_indexer = SqlAlchemyDocumentUploadIndexer(
        database_url=database_url,
        upload_root=upload_root,
        settings=DocumentUploadIndexingSettings(
            enabled=True,
            embedding_dimension=1024,
            source_package_version_key="personal-materials-test-package",
            index_version_key="personal-materials-test-candidate",
        ),
    )
    client = TestClient(create_app(state))
    upload_response = client.post(
        "/documents/uploads",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        files={
            "file": (
                "personal-note.txt",
                "医保基金审核个人补充材料，需核对院内台账。",
                "text/plain",
            )
        },
    )
    assert upload_response.status_code == 200
    upload_id = upload_response.json()["item"]["id"]
    for check_type, provider in (
        ("virus-scan", "clamav-sidecar"),
        ("dlp-review", "ruleset-v1"),
    ):
        result_response = client.post(
            f"/documents/uploads/{upload_id}/index-readiness/governance-result",
            headers={"X-User-Id": "head-1", "X-Role": "department-head"},
            json={
                "check_type": check_type,
                "provider": provider,
                "status": "passed",
                "detail": f"{check_type} passed in controlled test",
                "result_code": "clean" if check_type == "virus-scan" else "no-sensitive-marker",
            },
        )
        assert result_response.status_code == 200
    manual_response = client.post(
        f"/documents/uploads/{upload_id}/index-readiness/manual-approval",
        headers={"X-User-Id": "head-1", "X-Role": "department-head"},
        json={"decision": "approved", "note": "approved for candidate staging"},
    )
    assert manual_response.status_code == 200
    assert manual_response.json()["item"]["index_readiness"]["status"] == "ready"

    denied_response = client.post(
        f"/documents/uploads/{upload_id}/index-ingestion",
        headers={"X-User-Id": "auditor-2", "X-Role": "auditor"},
    )
    assert denied_response.status_code == 403
    assert denied_response.json()["detail"] == (
        "document upload index ingestion requires governance role"
    )

    index_response = client.post(
        f"/documents/uploads/{upload_id}/index-ingestion",
        headers={"X-User-Id": "head-1", "X-Role": "department-head"},
    )

    assert index_response.status_code == 200
    body = index_response.json()
    item = body["item"]
    ingestion = body["ingestion"]
    assert item["index_status"] == "staged-for-index"
    assert ingestion["status"] == "staged-for-index"
    assert ingestion["upload_key"] == upload_id
    assert ingestion["source_collection"] == "personal-materials"
    assert ingestion["source_package_version_key"] == "personal-materials-test-package"
    assert ingestion["index_version_key"] == "personal-materials-test-candidate"
    assert ingestion["index_version_status"] == "candidate"
    assert ingestion["chunk_count"] == 1
    assert ingestion["embedding_count"] == 1
    assert ingestion["embedding_dimension"] == 1024
    assert ingestion["external_provider_call_performed"] is False
    assert ingestion["live_retrieval_activated"] is False
    assert state.operation_logs[-1]["action"] == "document-upload-index-ingestion"

    already_response = client.post(
        f"/documents/uploads/{upload_id}/index-ingestion",
        headers={"X-User-Id": "head-1", "X-Role": "department-head"},
    )
    assert already_response.status_code == 200
    assert already_response.json()["ingestion"]["status"] == "already-staged"


def test_personal_document_index_is_governed_and_query_scoped(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.document_upload_store = InMemoryDocumentUploadStore(
        upload_root=tmp_path / "document-uploads",
    )
    client = TestClient(create_app(state))

    upload_response = client.post(
        "/documents/uploads",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        files={
            "file": (
                "local-note.txt",
                "院内个人材料提示：医保基金审核依据需核对院内报销清单。",
                "text/plain",
            )
        },
    )

    assert upload_response.status_code == 200
    uploaded = upload_response.json()["item"]
    assert uploaded["personal_index_status"] == "not-indexed"
    assert uploaded["personal_index_chunk_count"] == 0

    unapproved_index_response = client.post(
        f"/documents/uploads/{uploaded['id']}/index",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    )
    assert unapproved_index_response.status_code == 409
    assert unapproved_index_response.json()["detail"] == (
        "document upload must be approved before personal index"
    )

    governance_response = client.post(
        f"/documents/uploads/{uploaded['id']}/governance",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
        json={"governance_status": "approved-for-index", "note": "可进入个人材料索引。"},
    )
    assert governance_response.status_code == 200

    index_response = client.post(
        f"/documents/uploads/{uploaded['id']}/index",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    )
    assert index_response.status_code == 200
    indexed = index_response.json()["item"]
    assert indexed["personal_index_status"] == "indexed"
    assert indexed["personal_index_chunk_count"] == 1
    assert indexed["personal_indexed_by"] == "auditor-1"
    assert state.operation_logs[-1]["action"] == "document-upload-index"

    owner_query_response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2},
    )
    assert owner_query_response.status_code == 200
    owner_body = owner_query_response.json()
    assert owner_body["personal_upload_matches"][0]["upload_id"] == uploaded["id"]
    assert "院内个人材料提示" in owner_body["personal_upload_matches"][0]["snippet"]
    assert state.query_logs[-1]["personal_upload_match_count"] == 1

    other_query_response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-2", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2},
    )
    assert other_query_response.status_code == 200
    assert other_query_response.json()["personal_upload_matches"] == []

    admin_query_response = client.post(
        "/query",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
        json={"question": "医保基金审核依据", "top_k": 2},
    )
    assert admin_query_response.status_code == 200
    assert admin_query_response.json()["personal_upload_matches"][0]["upload_id"] == uploaded["id"]


def test_query_endpoint_returns_citation_answer_and_records_query_log(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["contract_version"] == KNOWLEDGE_QUERY_CONTRACT_VERSION
    assert "[C1]" in body["answer"]
    assert "medical-insurance-laws" in body["effective_source_collections"]
    assert "personal-materials" not in body["effective_source_collections"]
    assert body["citations"][0]["source_collection"] == "medical-insurance-laws"
    assert body["citations"][0]["index_version_key"] == "index-v1"
    assert body["citations"][0]["source_package_version_key"] == "package-v1"
    assert body["basis_groups"][0]["title"] == "法规依据"
    assert body["basis_groups"][0]["items"][0]["source_collection"] == "medical-insurance-laws"
    assert body["generation_status"] == "not_requested"
    assert body["generation_failure_code"] is None
    assert body["query_log_index"] == 0

    logs_response = client.get(
        "/query/logs",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    )
    assert logs_response.status_code == 200
    assert logs_response.json()["items"][0]["user_identifier"] == "auditor-1"
    assert logs_response.json()["items"][0]["filters"]["top_k"] == 2


def test_query_endpoint_excludes_personal_materials_from_default_retrieval(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.search_engine = _search_engine_with_personal_materials(
        system_chunk_id=uuid4(),
        personal_chunk_id=uuid4(),
        source_path="全量法律/law.md",
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据 院内个人材料提示", "top_k": 5},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["contract_version"] == KNOWLEDGE_QUERY_CONTRACT_VERSION
    assert body["citations"]
    assert {item["source_collection"] for item in body["citations"]} == {"medical-insurance-laws"}
    assert state.query_logs[-1]["filters"]["source_collections"] == []
    assert (
        "personal-materials" not in state.query_logs[-1]["filters"]["effective_source_collections"]
    )
    assert state.query_logs[-1]["filters"]["personal_material_scope"] == "none"

    explicit_personal_response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={
            "question": "医保基金审核依据",
            "top_k": 5,
            "source_collections": ["personal-materials"],
        },
    )
    assert explicit_personal_response.status_code == 200
    explicit_body = explicit_personal_response.json()
    assert explicit_body["effective_source_collections"] == ["personal-materials"]
    assert {item["source_collection"] for item in explicit_body["citations"]} == {
        "personal-materials"
    }
    assert state.query_logs[-1]["filters"]["source_collections"] == ["personal-materials"]
    assert state.query_logs[-1]["filters"]["effective_source_collections"] == ["personal-materials"]
    assert state.query_logs[-1]["filters"]["personal_material_scope"] == "self"

    other_auditor_response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-2", "X-Role": "auditor"},
        json={
            "question": "医保基金审核依据",
            "top_k": 5,
            "source_collections": ["personal-materials"],
        },
    )
    assert other_auditor_response.status_code == 404

    admin_personal_response = client.post(
        "/query",
        headers={"X-User-Id": "admin-1", "X-Role": "it-admin"},
        json={
            "question": "医保基金审核依据",
            "top_k": 5,
            "source_collections": ["personal-materials"],
        },
    )
    assert admin_personal_response.status_code == 200
    assert {item["source_collection"] for item in admin_personal_response.json()["citations"]} == {
        "personal-materials"
    }
    assert state.query_logs[-1]["filters"]["personal_material_scope"] == "all"


def test_query_endpoint_records_selected_agent_invocation(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(
        f"sqlite:///{tmp_path / 'query-agent-invocations.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={
            "X-User-Id": "auditor-1",
            "X-Role": "auditor",
            "X-Project-Name": PROJECT_NAME_HEADER,
        },
        json={
            "question": "医保基金审核依据",
            "top_k": 2,
            "agent": "agent-citation-check",
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["agent_invocation_id"]
    assert state.agent_store is not None
    invocations = state.agent_store.list_invocations("agent-citation-check")
    assert invocations[0]["id"] == body["agent_invocation_id"]
    assert invocations[0]["invocation_source"] == "/query"
    assert invocations[0]["question"] == "医保基金审核依据"
    assert invocations[0]["metadata"]["filters"]["agent"] == "agent-citation-check"
    assert state.operation_logs[-2]["action"] == "agent-invocation-create"
    assert state.operation_logs[-1]["action"] == "query"
    assert state.operation_logs[-1]["payload"]["agent_invocation_id"] == body["agent_invocation_id"]


def test_query_endpoint_records_installed_catalog_agent_invocation(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.agent_store = SqlAlchemyAgentStore(
        f"sqlite:///{tmp_path / 'query-installed-catalog-agent.db'}",
        create_schema=True,
    )
    client = TestClient(create_app(state))
    project_headers = {
        "X-User-Id": "director-1",
        "X-Role": "director",
        "X-Project-Name": PROJECT_NAME_HEADER,
    }
    catalog_metadata = {
        "catalog_source": "audit-agent-prompts-0613",
        "catalog_row_id": "audit-agent-prompts-0613-118",
        "source_key": "audit-agent-prompts-0613-118",
        "source_row_index": 118,
        "legacy_source_key": "工具智能体|医保基金问答助手",
        "source_title": "医保基金问答助手",
        "display_name": "基金问答助手",
        "avatar_seed": "工具智能体-医保基金问答助手-118",
    }
    create_response = client.post(
        "/agents",
        headers=project_headers,
        json={
            "name": "基金问答助手",
            "category": "效率类",
            "topic": "医保基金问答",
            "prompt": "围绕医保基金审核依据回答，并保留引用边界。",
            "knowledge_base": "系统医保审计知识库",
            "project_name": "医保基金使用合规专项自查",
            "visibility_scope": "project",
            "metadata": catalog_metadata,
        },
    )
    assert create_response.status_code == 200
    agent_id = str(create_response.json()["item"]["id"])

    missing_scope_response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={
            "question": "医保基金审核依据",
            "top_k": 2,
            "agent": agent_id,
        },
    )

    response = client.post(
        "/query",
        headers={
            "X-User-Id": "auditor-1",
            "X-Role": "auditor",
            "X-Project-Name": PROJECT_NAME_HEADER,
        },
        json={
            "question": "医保基金审核依据",
            "top_k": 2,
            "topic": "medical-insurance-fund",
            "source_collections": ["medical-insurance-laws"],
            "agent": agent_id,
        },
    )

    assert missing_scope_response.status_code == 403
    assert missing_scope_response.json()["detail"] == "agent project scope requires current project"
    assert response.status_code == 200
    body = response.json()
    assert body["agent_invocation_id"]
    assert state.agent_store is not None
    invocations = state.agent_store.list_invocations(agent_id)
    assert invocations[0]["id"] == body["agent_invocation_id"]
    assert invocations[0]["agent_key"] == agent_id
    assert invocations[0]["invocation_source"] == "/query"
    assert invocations[0]["question"] == "医保基金审核依据"
    invocation_metadata = invocations[0]["metadata"]
    assert invocation_metadata["filters"]["agent"] == agent_id
    assert invocation_metadata["filters"]["source_collections"] == ["medical-insurance-laws"]
    assert invocation_metadata["query_log_index"] == body["query_log_index"]
    assert invocation_metadata["project_name"] == "医保基金使用合规专项自查"
    stored_agent = state.agent_store.get_agent(agent_id)
    assert stored_agent is not None
    assert stored_agent["metadata"]["catalog_row_id"] == "audit-agent-prompts-0613-118"
    assert stored_agent["metadata"]["source_row_index"] == 118


def test_query_endpoint_supports_title_only_filter(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "title_only": True},
    )
    unmatched_response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "不存在的标题", "top_k": 2, "title_only": True},
    )

    assert response.status_code == 200
    assert response.json()["citations"][0]["source_collection"] == "medical-insurance-laws"
    assert state.query_logs[-1]["filters"]["title_only"] is True
    assert state.operation_logs[-1]["payload"]["filters"]["title_only"] is True
    assert unmatched_response.status_code == 404


def test_query_endpoint_scopes_to_topic_and_rejects_unknown(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    scoped = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "topic": "medical-insurance-fund"},
    )
    # 存量医保 chunk 无 domain 标签，靠 source_collection 兜底仍进专题。
    assert scoped.status_code == 200
    assert scoped.json()["citations"]
    assert state.query_logs[-1]["filters"]["topic"] == "medical-insurance-fund"

    unknown = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "topic": "no-such-topic"},
    )
    assert unknown.status_code == 400


def test_query_endpoint_uses_persistent_user_status_gate(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.auth_user_store = SqlAlchemyAuthUserStore(
        f"sqlite:///{tmp_path / 'auth-users.db'}",
        create_schema=True,
    )
    assert state.auth_user_store is not None
    state.auth_user_store.add_user(
        {
            "user_key": "disabled-query-user",
            "display_name": "停用查询用户",
            "status": "disabled",
        }
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "disabled-query-user", "X-Role": "it-admin"},
        json={"question": "医保基金审核依据", "top_k": 2},
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "auth user status is disabled"


def test_query_endpoint_persists_query_history(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'query-history.db'}"
    state = _api_state(tmp_path)
    state.query_history_store = SqlAlchemyQueryHistoryStore(database_url, create_schema=True)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={
            "question": "医保基金审核依据",
            "top_k": 2,
            "source_collections": ["medical-insurance-laws"],
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["query_log_id"]

    history_response = client.get(
        "/query/logs?limit=5",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    )
    assert history_response.status_code == 200
    history_body = history_response.json()
    assert history_body["store"]["backend"] == "SqlAlchemyQueryHistoryStore"
    assert history_body["items"][0]["id"] == body["query_log_id"]
    assert history_body["items"][0]["question"] == "医保基金审核依据"
    assert history_body["items"][0]["filters"]["source_collections"] == ["medical-insurance-laws"]
    assert history_body["items"][0]["citation_count"] == 1
    assert history_body["items"][0]["answer_summary"]
    assert history_body["items"][0]["generation_status"] == "not_requested"
    assert history_body["items"][0]["generation_failure_code"] is None

    second_state = _api_state(tmp_path / "second")
    second_state.query_history_store = SqlAlchemyQueryHistoryStore(database_url)
    second_client = TestClient(create_app(second_state))
    persisted_items = second_client.get(
        "/query/logs",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    ).json()["items"]
    assert persisted_items[0]["id"] == body["query_log_id"]
    assert persisted_items[0]["generation_status"] == "not_requested"
    assert persisted_items[0]["generation_failure_code"] is None


def test_query_endpoint_persists_generation_fallback_after_restart(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_query.answer_generation_provider_for_alias",
        lambda _alias: FailingApiAnswerProvider(),
        raising=False,
    )
    database_url = f"sqlite:///{tmp_path / 'query-generation-history.db'}"
    state = _api_state(tmp_path)
    state.query_history_store = SqlAlchemyQueryHistoryStore(database_url, create_schema=True)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "model": "kimi-2.7"},
    )

    assert response.status_code == 200
    second_state = _api_state(tmp_path / "second-generation-history")
    second_state.query_history_store = SqlAlchemyQueryHistoryStore(database_url)
    persisted = (
        TestClient(create_app(second_state))
        .get(
            "/query/logs",
            headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        )
        .json()["items"][0]
    )
    assert persisted["generation_status"] == "retrieval_fallback"
    assert persisted["generation_failure_code"] == "provider_exception"


def test_query_history_store_failure_does_not_block_query(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.query_history_store = FailingQueryHistoryStore()
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["query_log_id"] is None
    assert state.operation_logs[-1]["payload"]["query_history_error"] == {
        "error_type": "RuntimeError",
        "message": "query history store operation failed",
    }

    logs_response = client.get(
        "/query/logs",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
    )
    assert logs_response.status_code == 200
    logs_body = logs_response.json()
    assert logs_body["store"] == {
        "ready": False,
        "backend": "FailingQueryHistoryStore",
        "error": {
            "error_type": "RuntimeError",
            "message": "query history store operation failed",
        },
    }
    assert logs_body["items"][0]["question"] == "医保基金审核依据"


def test_query_endpoint_uses_configured_answer_generation_provider(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.app.answer_generation_provider_from_settings",
        lambda _settings: StaticApiAnswerProvider(),
        raising=False,
    )
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["fallback_used"] is False
    assert body["generation_status"] == "generated"
    assert body["generation_failure_code"] is None
    assert body["answer"] == "生成模型回答：应核验医保基金审核依据 [C1]。"


def test_query_models_reports_alias_availability_without_secret_values(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _clear_chat_model_env(monkeypatch)
    monkeypatch.setenv("MEDICAL_AUDIT_KB_CHAT_MODEL_KIMI_2_7_API_KEY_ENV", "TEST_KIMI_CHAT_KEY")
    monkeypatch.setenv("MEDICAL_AUDIT_KB_CHAT_MODEL_KIMI_2_7_MODEL", "moonshot-test-model")
    monkeypatch.setenv("TEST_KIMI_CHAT_KEY", "secret-value-not-in-response")
    client = TestClient(create_app(_api_state(tmp_path)))

    response = client.get("/query/models")

    assert response.status_code == 200
    body = response.json()
    serialized = json.dumps(body, ensure_ascii=False)
    assert body["contract_version"] == "chat-model-catalog-v1"
    assert body["default_model"] == "kimi-2.7"
    kimi = next(item for item in body["items"] if item["alias"] == "kimi-2.7")
    deepseek = next(item for item in body["items"] if item["alias"] == "deepseek-v4-pro")
    assert kimi["available"] is True
    assert kimi["provider"] == "kimi"
    assert deepseek["available"] is False
    assert "secret-value-not-in-response" not in serialized


@pytest.mark.parametrize(
    (
        "alias",
        "api_key_env",
        "expected_model",
        "expected_base_url",
        "expected_temperature",
        "expected_max_output_tokens",
        "expected_thinking_mode",
    ),
    (
        (
            ChatModelAlias.KIMI_2_7,
            "TEST_KIMI_CHAT_KEY",
            "kimi-k2.6",
            "https://api.moonshot.cn/v1",
            1.0,
            4096,
            "enabled",
        ),
        (
            ChatModelAlias.DEEPSEEK_V4_PRO,
            "TEST_DEEPSEEK_CHAT_KEY",
            "deepseek-v4-pro",
            "https://api.deepseek.com",
            0.0,
            900,
            "disabled",
        ),
    ),
)
def test_chat_model_config_uses_verified_provider_defaults(
    monkeypatch: pytest.MonkeyPatch,
    alias: ChatModelAlias,
    api_key_env: str,
    expected_model: str,
    expected_base_url: str,
    expected_temperature: float,
    expected_max_output_tokens: int,
    expected_thinking_mode: str,
) -> None:
    _clear_chat_model_env(monkeypatch)
    env_slug = alias.value.upper().replace("-", "_").replace(".", "_")
    monkeypatch.setenv(
        f"MEDICAL_AUDIT_KB_CHAT_MODEL_{env_slug}_API_KEY_ENV",
        api_key_env,
    )
    monkeypatch.setenv(api_key_env, "secret-value-not-returned")

    config, reason = chat_model_config_from_env(alias)

    assert reason is None
    assert config is not None
    assert config.model_name == expected_model
    assert config.base_url == expected_base_url
    assert config.temperature == expected_temperature
    assert config.max_output_tokens == expected_max_output_tokens
    assert config.thinking_mode == expected_thinking_mode


def test_kimi_chat_model_rejects_insufficient_output_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _clear_chat_model_env(monkeypatch)
    monkeypatch.setenv("MEDICAL_AUDIT_KB_CHAT_MODEL_KIMI_2_7_API_KEY_ENV", "TEST_KIMI_CHAT_KEY")
    monkeypatch.setenv("MEDICAL_AUDIT_KB_CHAT_MODEL_KIMI_2_7_MAX_OUTPUT_TOKENS", "900")
    monkeypatch.setenv("TEST_KIMI_CHAT_KEY", "secret-value-not-returned")

    config, reason = chat_model_config_from_env(ChatModelAlias.KIMI_2_7)

    assert config is None
    assert reason == "insufficient_output_budget"


def test_deepseek_chat_model_rejects_enabled_thinking(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _clear_chat_model_env(monkeypatch)
    monkeypatch.setenv(
        "MEDICAL_AUDIT_KB_CHAT_MODEL_DEEPSEEK_V4_PRO_API_KEY_ENV",
        "TEST_DEEPSEEK_CHAT_KEY",
    )
    monkeypatch.setenv(
        "MEDICAL_AUDIT_KB_CHAT_MODEL_DEEPSEEK_V4_PRO_THINKING_MODE",
        "enabled",
    )
    monkeypatch.setenv("TEST_DEEPSEEK_CHAT_KEY", "secret-value-not-returned")

    config, reason = chat_model_config_from_env(ChatModelAlias.DEEPSEEK_V4_PRO)

    assert config is None
    assert reason == "unsupported_thinking_mode"


def test_query_models_fake_provider_requires_explicit_local_gate(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _clear_chat_model_env(monkeypatch)
    monkeypatch.setenv("MEDICAL_AUDIT_KB_CHAT_MODEL_KIMI_2_7_API_KEY_ENV", "TEST_KIMI_CHAT_KEY")
    monkeypatch.setenv("MEDICAL_AUDIT_KB_CHAT_MODEL_KIMI_2_7_PROVIDER", "fake")
    monkeypatch.setenv("TEST_KIMI_CHAT_KEY", "local-fake-key")
    client = TestClient(create_app(_api_state(tmp_path)))

    blocked_response = client.get("/query/models")

    blocked_kimi = next(
        item for item in blocked_response.json()["items"] if item["alias"] == "kimi-2.7"
    )
    assert blocked_kimi["available"] is False
    assert blocked_kimi["unavailable_reason"] == "fake_provider_not_allowed"

    monkeypatch.setenv("MEDICAL_AUDIT_KB_ALLOW_FAKE_CHAT_MODELS", "1")
    allowed_response = client.post(
        "/query",
        headers={"X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "model": "kimi-2.7"},
    )

    assert allowed_response.status_code == 200
    body = allowed_response.json()
    assert body["model_alias"] == "kimi-2.7"
    assert body["model_status"] == "selected_provider"
    assert body["answer"].startswith("本地验收模型回答")


def test_query_endpoint_rejects_unconfigured_selected_model_before_logging(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _clear_chat_model_env(monkeypatch)
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "model": "kimi-2.7"},
    )

    assert response.status_code == 409
    assert response.json()["detail"] == {
        "code": "chat_model_unavailable",
        "model": "kimi-2.7",
        "reason": "missing_api_key_env",
    }
    assert state.query_logs == []


def test_query_endpoint_uses_selected_chat_model_alias(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_query.answer_generation_provider_for_alias",
        lambda _alias: StaticApiAnswerProvider(),
        raising=False,
    )
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-Role": "auditor"},
        json={
            "question": "医保基金审核依据",
            "top_k": 2,
            "model": "kimi-2.7",
            "source_collections": ["medical-insurance-laws"],
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["model_alias"] == "kimi-2.7"
    assert body["model_status"] == "selected_provider"
    assert body["fallback_used"] is False
    assert body["generation_status"] == "generated"
    assert body["generation_failure_code"] is None
    assert state.query_logs[-1]["filters"]["model"] == "kimi-2.7"
    assert state.query_logs[-1]["generation_status"] == "generated"
    assert state.query_logs[-1]["generation_failure_code"] is None
    assert state.operation_logs[-1]["payload"]["model"] == "kimi-2.7"


def test_query_endpoint_reports_sanitized_generation_fallback(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_query.answer_generation_provider_for_alias",
        lambda _alias: FailingApiAnswerProvider(),
        raising=False,
    )
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "model": "kimi-2.7"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["fallback_used"] is True
    assert body["generation_status"] == "retrieval_fallback"
    assert body["generation_failure_code"] == "provider_exception"
    assert "private provider detail" not in json.dumps(body, ensure_ascii=False)
    assert state.query_logs[-1]["generation_status"] == "retrieval_fallback"
    assert state.query_logs[-1]["generation_failure_code"] == "provider_exception"
    operation_payload = state.operation_logs[-1]["payload"]
    assert isinstance(operation_payload, dict)
    assert operation_payload["generation_status"] == "retrieval_fallback"
    assert operation_payload["generation_failure_code"] == "provider_exception"


def test_query_endpoint_reports_and_persists_safe_provider_http_status(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_query.answer_generation_provider_for_alias",
        lambda _alias: HttpStatusFailingApiAnswerProvider(),
        raising=False,
    )
    database_url = f"sqlite:///{tmp_path / 'query-http-status-history.db'}"
    state = _api_state(tmp_path)
    state.query_history_store = SqlAlchemyQueryHistoryStore(database_url, create_schema=True)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "model": "kimi-2.7"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["generation_failure_code"] == "provider_http_status"
    assert body["generation_http_status"] == 429
    assert state.query_logs[-1]["generation_http_status"] == 429
    assert state.operation_logs[-1]["payload"]["generation_http_status"] == 429

    second_state = _api_state(tmp_path / "second-http-status-history")
    second_state.query_history_store = SqlAlchemyQueryHistoryStore(database_url)
    persisted = (
        TestClient(create_app(second_state))
        .get(
            "/query/logs",
            headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        )
        .json()["items"][0]
    )
    assert persisted["generation_http_status"] == 429


def test_query_endpoint_reports_and_persists_safe_provider_failure_reason(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_query.answer_generation_provider_for_alias",
        lambda _alias: InvalidResponseApiAnswerProvider(),
        raising=False,
    )
    database_url = f"sqlite:///{tmp_path / 'query-failure-reason-history.db'}"
    state = _api_state(tmp_path)
    state.query_history_store = SqlAlchemyQueryHistoryStore(database_url, create_schema=True)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "model": "deepseek-v4-pro"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["generation_failure_code"] == "provider_response_invalid"
    assert body["generation_failure_reason"] == "deepseek_content_invalid_json"
    assert "private provider sentinel" not in json.dumps(body, ensure_ascii=False)
    assert state.query_logs[-1]["generation_failure_reason"] == "deepseek_content_invalid_json"
    assert (
        state.operation_logs[-1]["payload"]["generation_failure_reason"]
        == "deepseek_content_invalid_json"
    )

    second_state = _api_state(tmp_path / "second-failure-reason-history")
    second_state.query_history_store = SqlAlchemyQueryHistoryStore(database_url)
    persisted = (
        TestClient(create_app(second_state))
        .get(
            "/query/logs",
            headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        )
        .json()["items"][0]
    )
    assert persisted["generation_failure_reason"] == "deepseek_content_invalid_json"


def test_query_endpoint_propagates_and_persists_provider_abstention(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_query.answer_generation_provider_for_alias",
        lambda _alias: AbstainingApiAnswerProvider(),
        raising=False,
    )
    database_url = f"sqlite:///{tmp_path / 'query-provider-abstention-history.db'}"
    state = _api_state(tmp_path)
    state.query_history_store = SqlAlchemyQueryHistoryStore(database_url, create_schema=True)
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={"question": "医保基金审核依据", "top_k": 2, "model": "deepseek-v4-pro"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["fallback_used"] is True
    assert body["generation_status"] == "retrieval_fallback"
    assert body["generation_failure_code"] == "provider_abstention"
    assert body["generation_failure_reason"] == "deepseek_strict_insufficient_evidence"
    assert "private provider abstention sentinel" not in json.dumps(body, ensure_ascii=False)

    query_log = state.query_logs[-1]
    assert query_log["generation_failure_code"] == "provider_abstention"
    assert query_log["generation_failure_reason"] == "deepseek_strict_insufficient_evidence"
    assert "private provider abstention sentinel" not in json.dumps(
        query_log,
        ensure_ascii=False,
    )
    operation_payload = state.operation_logs[-1]["payload"]
    assert operation_payload["generation_failure_code"] == "provider_abstention"
    assert operation_payload["generation_failure_reason"] == (
        "deepseek_strict_insufficient_evidence"
    )
    assert "private provider abstention sentinel" not in json.dumps(
        operation_payload,
        ensure_ascii=False,
    )

    second_state = _api_state(tmp_path / "second-provider-abstention-history")
    second_state.query_history_store = SqlAlchemyQueryHistoryStore(database_url)
    persisted = (
        TestClient(create_app(second_state))
        .get(
            "/query/logs",
            headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        )
        .json()["items"][0]
    )
    assert persisted["generation_failure_code"] == "provider_abstention"
    assert persisted["generation_failure_reason"] == ("deepseek_strict_insufficient_evidence")
    assert "private provider abstention sentinel" not in json.dumps(
        persisted,
        ensure_ascii=False,
    )


def test_chat_attachment_analyzes_table_with_selected_model(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_chat.answer_generation_provider_for_alias",
        lambda _alias: StaticApiAnswerProvider(),
        raising=False,
    )
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.post(
        "/chat/attachments/analyze",
        data={"model": "kimi-2.7", "mode": "auto"},
        files={
            "file": (
                "charge-sample.csv",
                "patient_id,charge_amount\nP001,120\nP002,80\n",
                "text/csv",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["contract_version"] == "chat-attachment-analysis-v1"
    assert body["mode"] == "table-analysis"
    assert body["model_alias"] == "kimi-2.7"
    assert "行数：2" in body["summary_items"]
    assert "patient_id" in body["extracted_preview"]
    assert body["answer"] == "生成模型回答：应核验医保基金审核依据 [C1]。"
    assert body["boundaries"]["object_storage_write"] is False
    assert state.operation_logs[-1]["action"] == "chat-attachment-analyze"


def test_chat_attachment_falls_back_without_selected_model(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    response = client.post(
        "/chat/attachments/analyze",
        data={"mode": "auto"},
        files={
            "file": (
                "charge-sample.csv",
                "patient_id,charge_amount\nP001,120\nP002,80\n",
                "text/csv",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["contract_version"] == "chat-attachment-analysis-v1"
    assert body["mode"] == "table-analysis"
    assert body["model_alias"] is None
    assert body["model_status"] == "default_fallback"
    assert body["boundaries"]["provider_call"] is False
    assert "未调用外部模型" in body["answer"]
    assert "行数：2" in body["summary_items"]
    assert state.operation_logs[-1]["payload"]["provider_call"] is False


def test_chat_attachment_analyzes_text_document_with_selected_model(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_chat.answer_generation_provider_for_alias",
        lambda _alias: StaticApiAnswerProvider(),
        raising=False,
    )
    client = TestClient(create_app(_api_state(tmp_path)))

    response = client.post(
        "/chat/attachments/analyze",
        data={"model": "deepseek-v4-pro", "mode": "document-summary"},
        files={
            "file": (
                "meeting-summary.txt",
                "医保基金审计会议纪要：要求复核高频收费项目和异常结算。",
                "text/plain",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["mode"] == "document-summary"
    assert body["model_alias"] == "deepseek-v4-pro"
    assert "字符数" in body["summary_items"][2]
    assert "医保基金审计会议纪要" in body["extracted_preview"]


def test_chat_attachment_rejects_webp_before_ocr(tmp_path: Path) -> None:
    class TrackingUnlimitedOcrClient:
        calls = 0

        async def extract_text(
            self,
            *,
            file_name: str,
            extension: str,
            content: bytes,
        ) -> UnlimitedOcrResult:
            self.calls += 1
            return UnlimitedOcrResult(
                text="不应调用 OCR",
                page_count=1,
                model="baidu/Unlimited-OCR",
                source_commit="d49ff64afffc1f47ab563dc1c589bc2f78808fa4",
            )

    ocr_client = TrackingUnlimitedOcrClient()
    state = _api_state(tmp_path)
    state.ocr_client = ocr_client
    client = TestClient(create_app(state))

    response = client.post(
        "/chat/attachments/analyze",
        data={"mode": "auto"},
        files={"file": ("unsupported.webp", b"not-a-webp", "image/webp")},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "unsupported chat attachment extension"
    assert ocr_client.calls == 0


def test_chat_attachment_rejects_image_only_pdf_with_actionable_ocr_message(
    tmp_path: Path,
) -> None:
    pdf_bytes = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(pdf_bytes)
    client = TestClient(create_app(_api_state(tmp_path)))

    response = client.post(
        "/chat/attachments/analyze",
        data={"model": "deepseek-v4-pro", "mode": "auto"},
        files={
            "file": (
                "scanned-audit-material.pdf",
                pdf_bytes.getvalue(),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 422
    assert response.json()["detail"] == (
        "PDF 未检测到可读取文字，可能是扫描件或图片型 PDF。"
        "请先进行 OCR 识别，或上传可搜索文字版 PDF。"
    )


def test_chat_attachment_uses_unlimited_ocr_for_image_only_pdf(
    tmp_path: Path,
) -> None:
    class StaticUnlimitedOcrClient:
        async def extract_text(
            self,
            *,
            file_name: str,
            extension: str,
            content: bytes,
        ) -> UnlimitedOcrResult:
            assert file_name == "scanned-audit-material.pdf"
            assert extension == "pdf"
            assert content.startswith(b"%PDF")
            return UnlimitedOcrResult(
                text="扫描件识别结果：医保基金审计应复核高频收费项目。",
                page_count=1,
                model="baidu/Unlimited-OCR",
                source_commit="d49ff64afffc1f47ab563dc1c589bc2f78808fa4",
            )

    pdf_bytes = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(pdf_bytes)
    state = _api_state(tmp_path)
    state.ocr_client = StaticUnlimitedOcrClient()
    client = TestClient(create_app(state))

    response = client.post(
        "/chat/attachments/analyze",
        data={"mode": "auto"},
        files={
            "file": (
                "scanned-audit-material.pdf",
                pdf_bytes.getvalue(),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["mode"] == "document-summary"
    assert body["boundaries"]["ocr_call"] is True
    assert body["boundaries"]["ocr_engine"] == "baidu/Unlimited-OCR"
    assert body["boundaries"]["answer_provider_call"] is False
    assert body["boundaries"]["provider_call"] is True
    assert body["summary_items"][0] == "OCR 页数：1"
    assert "扫描件识别结果" in body["extracted_preview"]
    assert state.operation_logs[-1]["payload"]["ocr_source_commit"] == (
        "d49ff64afffc1f47ab563dc1c589bc2f78808fa4"
    )


def test_query_endpoint_blocks_unknown_role(tmp_path: Path) -> None:
    client = TestClient(create_app(_api_state(tmp_path)))

    response = client.post(
        "/query",
        headers={"X-Role": "guest"},
        json={"question": "医保基金审核依据"},
    )

    assert response.status_code == 403


def test_preview_endpoint_resolves_chunk_reference_after_query(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))
    query_response = client.post(
        "/query",
        headers={"X-Role": "auditor"},
        json={"question": "医保基金审核依据"},
    )
    chunk_id = query_response.json()["citations"][0]["chunk_id"]

    preview_response = client.get(f"/preview/{chunk_id}")

    assert preview_response.status_code == 200
    body = preview_response.json()
    assert body["chunk_id"] == chunk_id
    assert "医疗机构应当保留医保基金审核依据" in body["preview_text"]
    assert body["highlights"]
    assert body["line_start"] == 1


def test_index_rebuild_incremental_lists_and_permissions(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    source_root = state.source_root
    _write_text(source_root / "医保目录" / "catalog.md", "# 医保目录\n医保目录内容")
    _write_text(source_root / "风险负面清单" / "risk.png", "pending")
    client = TestClient(create_app(state))

    forbidden_response = client.post(
        "/index/rebuild",
        headers={"X-Role": "auditor"},
        json={"package_version_key": "api-package"},
    )
    assert forbidden_response.status_code == 403

    rebuild_response = client.post(
        "/index/rebuild",
        headers={"X-Role": "it-admin"},
        json={"package_version_key": "api-package"},
    )
    assert rebuild_response.status_code == 200
    rebuild_summary = rebuild_response.json()["summary"]
    assert rebuild_summary["job_type"] == "full-rebuild"
    # 全量入库后 全量法律/law.md（医保内容、非医保文件名）也成候选 → catalog.md + law.md = 2。
    assert rebuild_summary["index_candidate_file_count"] == 2
    assert rebuild_summary["pending_file_count"] == 1

    incremental_response = client.post(
        "/index/incremental",
        headers={"X-Role": "it-admin"},
        json={"package_version_key": "api-package-next"},
    )
    assert incremental_response.status_code == 200
    assert incremental_response.json()["summary"]["job_type"] == "incremental"

    versions = client.get("/index/versions").json()["items"]
    jobs = client.get("/index/jobs").json()["items"]
    pending = client.get("/index/pending").json()["items"]
    failures = client.get("/index/failures").json()["items"]

    assert len(versions) == 2
    assert len(jobs) == 2
    assert jobs[-1]["status"] == "succeeded"
    assert pending[0]["relative_path"] == "风险负面清单/risk.png"
    assert failures == []


def test_index_rebuild_allows_technician_role(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    _write_text(state.source_root / "医保目录" / "catalog.md", "# 医保目录\n医保目录内容")
    client = TestClient(create_app(state))

    response = client.post(
        "/index/rebuild",
        headers={"X-User-Id": "technician-1", "X-Role": "technician"},
        json={"package_version_key": "technician-package"},
    )

    assert response.status_code == 200
    assert response.json()["summary"]["job_type"] == "full-rebuild"


def test_index_retry_file_reports_missing_target(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    _write_text(state.source_root / "医保目录" / "catalog.md", "# 医保目录\n医保目录内容")
    client = TestClient(create_app(state))

    response = client.post(
        "/index/retry-file",
        headers={"X-Role": "it-admin"},
        json={"package_version_key": "retry-package", "relative_path": "医保目录/missing.md"},
    )

    assert response.status_code == 200
    assert response.json()["summary"]["failed_file_count"] == 1
    assert (
        client.get("/index/failures").json()["items"][0]["relative_path"] == "医保目录/missing.md"
    )


def test_index_postgres_status_reports_database_summary(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)

    def fake_load_postgres_index_status(database_url: str) -> dict[str, object]:
        return {
            "available": True,
            "row_counts": {
                "source_package_versions": 1,
                "source_documents": 486,
                "document_chunks": 48985,
                "chunk_embeddings": 48985,
                "index_versions": 1,
                "index_jobs": 1,
                "failed_files": 0,
                "pending_files": 13,
            },
            "embedding_sets": [
                {
                    "provider": "openai",
                    "model_name": "kimi-for-coding",
                    "provider_version": "v1",
                    "dimension": 1024,
                    "embedding_count": 48985,
                }
            ],
            "index_versions": [
                {
                    "version_key": "source-package-real-data-kimi-20260531",
                    "status": "active",
                    "vector_provider": "openai",
                    "vector_model": "kimi-for-coding",
                    "chunk_count": 48985,
                    "document_count": 486,
                }
            ],
            "source_packages": [
                {
                    "version_key": "source-package-real-data-kimi-20260531",
                    "source_root_path": "data/医保审核前期资料",
                }
            ],
        }

    monkeypatch.setattr(
        "medical_audit_kb.api.routes_index.load_postgres_index_status",
        fake_load_postgres_index_status,
    )
    client = TestClient(create_app(state))

    response = client.get("/index/postgres-status")

    assert response.status_code == 200
    body = response.json()
    assert body["row_counts"]["document_chunks"] == 48985
    assert body["row_counts"]["pending_files"] == 13
    assert body["embedding_sets"][0]["model_name"] == "kimi-for-coding"
    assert state.operation_logs[-1]["action"] == "postgres-index-status-view"


def test_index_postgres_search_backend_loads_with_admin_role(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)
    state.search_engine = None
    loaded_engine = _search_engine(_chunk_id(), "全量法律/law.md")
    captured: dict[str, object] = {}

    def fake_load_postgres_hybrid_search_engine(
        *,
        database_url: str,
        embedding_provider: DeterministicFakeEmbeddingProvider,
    ) -> HybridSearchEngine:
        captured["database_url"] = database_url
        captured["embedding_provider"] = embedding_provider
        return loaded_engine

    monkeypatch.setattr(
        "medical_audit_kb.api.routes_index.load_postgres_hybrid_search_engine",
        fake_load_postgres_hybrid_search_engine,
    )
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_index.count_postgres_embeddings",
        lambda database_url, embedding_provider: 48985,
    )
    client = TestClient(create_app(state))

    status_response = client.get("/index/search-backend")
    assert status_response.status_code == 200
    assert status_response.json() == {"backend": "none", "ready": False, "details": {}}

    forbidden_response = client.post(
        "/index/search-backend/postgres",
        headers={"X-Role": "auditor"},
        json={
            "embedding_provider": "fake",
            "embedding_model": "fake",
            "embedding_dimension": 32,
        },
    )
    assert forbidden_response.status_code == 403

    response = client.post(
        "/index/search-backend/postgres",
        headers={"X-Role": "it-admin"},
        json={
            "embedding_provider": "fake",
            "embedding_model": "fake",
            "embedding_dimension": 32,
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["backend"] == "postgres"
    assert body["ready"] is True
    assert body["details"]["embedding_provider"] == "fake"
    assert body["details"]["embedding_dimension"] == 32
    assert body["details"]["matching_embedding_count"] == 48985
    assert id(state.search_engine) == id(loaded_engine)
    assert captured["database_url"] == state.settings.database_url


def test_index_postgres_search_backend_rejects_unmatched_metadata(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)
    state.search_engine = None

    def fake_count_postgres_embeddings(
        database_url: str,
        embedding_provider: DeterministicFakeEmbeddingProvider,
    ) -> int:
        raise ValueError("no postgres embeddings match requested provider metadata")

    monkeypatch.setattr(
        "medical_audit_kb.api.routes_index.count_postgres_embeddings",
        fake_count_postgres_embeddings,
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/index/search-backend/postgres",
        headers={"X-Role": "it-admin"},
        json={
            "embedding_provider": "fake",
            "embedding_model": "fake",
            "embedding_dimension": 1536,
        },
    )

    assert response.status_code == 409
    assert "no postgres embeddings match" in response.json()["detail"]
    assert state.search_engine is None
    assert state.search_backend == "none"


def test_index_version_activate_and_rollback_actions_require_admin(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)
    client = TestClient(create_app(state))

    monkeypatch.setattr(
        "medical_audit_kb.api.routes_index.activate_index_version",
        lambda **kwargs: IndexActivationResult(
            index_version_key=str(kwargs["index_version_key"]),
            vector_provider="openai",
            vector_model="kimi-for-coding",
            previous_status="candidate",
            deactivated_index_version_keys=("active-old",),
        ),
    )
    monkeypatch.setattr(
        "medical_audit_kb.api.routes_index.rollback_index_version",
        lambda **kwargs: IndexRollbackResult(
            index_version_key=str(kwargs["index_version_key"]),
            vector_provider="openai",
            vector_model="kimi-for-coding",
            previous_status="inactive",
            deactivated_index_version_keys=("active-current",),
        ),
    )

    forbidden_response = client.post(
        "/index/versions/activate",
        headers={"X-Role": "auditor"},
        json={"index_version_key": "candidate-next"},
    )
    assert forbidden_response.status_code == 403

    activate_response = client.post(
        "/index/versions/activate",
        headers={"X-Role": "it-admin"},
        json={"index_version_key": "candidate-next"},
    )
    assert activate_response.status_code == 200
    assert activate_response.json()["result"]["index_version_key"] == "candidate-next"
    assert str(state.operation_logs[-1]["action"]) == "index-version-activate"

    rollback_response = client.post(
        "/index/versions/rollback",
        headers={"X-Role": "it-admin"},
        json={"index_version_key": "active-old"},
    )
    assert rollback_response.status_code == 200
    assert rollback_response.json()["result"]["previous_status"] == "inactive"
    assert str(state.operation_logs[-1]["action"]) == "index-version-rollback"


def test_index_version_switch_returns_conflict_for_invalid_state(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client = TestClient(create_app(_api_state(tmp_path)))

    def fake_activate_index_version(**kwargs: object) -> IndexActivationResult:
        raise IndexActivationError("index version must be candidate or active")

    monkeypatch.setattr(
        "medical_audit_kb.api.routes_index.activate_index_version",
        fake_activate_index_version,
    )

    response = client.post(
        "/index/versions/activate",
        headers={"X-Role": "it-admin"},
        json={"index_version_key": "inactive-old"},
    )

    assert response.status_code == 409
    assert "candidate or active" in response.json()["detail"]


def test_index_evaluation_run_scores_runtime_backend_and_records_status(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    retrieval_cases = _write_text(
        tmp_path / "retrieval-cases.yaml",
        """
cases:
  - case_id: retrieval-case-001
    question: 医疗机构需要保留什么审核依据？
    expected_evidence:
      - source_collection: medical-insurance-laws
        source_path: 全量法律/law.md
        article_or_rule: 第一条
""".strip(),
    )
    answer_cases = _write_text(
        tmp_path / "answer-cases.yaml",
        """
cases:
  - case_id: answer-case-001
    question: 医疗机构需要保留什么审核依据？
    expected_behavior: answer
    required_evidence_terms: [第一条]
    required_answer_terms: [审核依据]
    required_citation_terms: [医疗机构]
""".strip(),
    )
    client = TestClient(create_app(state))

    forbidden_response = client.post(
        "/index/evaluation/run",
        headers={"X-Role": "auditor"},
        json={},
    )
    assert forbidden_response.status_code == 403

    response = client.post(
        "/index/evaluation/run",
        headers={"X-Role": "it-admin"},
        json={
            "retrieval_cases_file": str(retrieval_cases),
            "answer_cases_file": str(answer_cases),
            "max_retrieval_cases": 1,
            "max_answer_cases": 1,
            "top_k": 2,
            "smoke_question": "医保基金审核依据",
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "pass"
    assert body["retrieval"]["case_count"] == 1
    assert body["retrieval"]["recall_at_k"] == 1.0
    assert body["answer"]["case_count"] == 1
    assert body["answer"]["pass_rate"] == 1.0
    assert body["ui_smoke"]["success"] is True
    assert state.evaluation_runs[-1]["status"] == "pass"
    assert state.operation_logs[-1]["action"] == "index-evaluation-run"


def test_index_evaluation_run_persists_json_report_and_exports_latest(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    retrieval_cases = _write_text(
        tmp_path / "retrieval-cases.yaml",
        """
cases:
  - case_id: retrieval-case-001
    question: 医疗机构需要保留什么审核依据？
    expected_evidence:
      - source_collection: medical-insurance-laws
        source_path: 全量法律/law.md
        article_or_rule: 第一条
""".strip(),
    )
    answer_cases = _write_text(
        tmp_path / "answer-cases.yaml",
        """
cases:
  - case_id: answer-case-001
    question: 医疗机构需要保留什么审核依据？
    expected_behavior: answer
    required_evidence_terms: [第一条]
    required_answer_terms: [审核依据]
    required_citation_terms: [医疗机构]
""".strip(),
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/index/evaluation/run",
        headers={"X-Role": "it-admin"},
        json={
            "retrieval_cases_file": str(retrieval_cases),
            "answer_cases_file": str(answer_cases),
            "max_retrieval_cases": 1,
            "max_answer_cases": 1,
            "top_k": 2,
            "smoke_question": "医保基金审核依据",
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["report"]["download_path"] == "/index/evaluation/latest/export"
    report_path = Path(body["report"]["path"])
    assert report_path.is_file()
    assert report_path.parent == state.settings.index_root / "evaluation-runs"
    persisted = json.loads(report_path.read_text(encoding="utf-8"))
    assert persisted["status"] == "pass"
    assert persisted["request"]["max_retrieval_cases"] == 1
    assert persisted["search_backend"]["backend"] == "none"

    export_response = client.get("/index/evaluation/latest/export")

    assert export_response.status_code == 200
    exported = export_response.json()
    assert exported["run_id"] == body["report"]["run_id"]
    assert exported["status"] == "pass"
    assert state.operation_logs[-1]["action"] == "index-evaluation-report-export"


def test_index_evaluation_run_records_postgres_history_when_backend_is_postgres(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)
    state.search_backend = "postgres"
    state.search_backend_details = {
        "embedding_model": "kimi-for-coding",
        "embedding_dimension": 1024,
        "matching_embedding_count": 48985,
    }
    retrieval_cases = _write_text(
        tmp_path / "retrieval-cases.yaml",
        """
cases:
  - case_id: retrieval-case-001
    question: 医疗机构需要保留什么审核依据？
    expected_evidence:
      - source_collection: medical-insurance-laws
        source_path: 全量法律/law.md
        article_or_rule: 第一条
""".strip(),
    )
    answer_cases = _write_text(
        tmp_path / "answer-cases.yaml",
        """
cases:
  - case_id: answer-case-001
    question: 医疗机构需要保留什么审核依据？
    expected_behavior: answer
    required_evidence_terms: [第一条]
    required_answer_terms: [审核依据]
    required_citation_terms: [医疗机构]
""".strip(),
    )
    captured: dict[str, object] = {}

    def fake_persist_evaluation_history(
        history_state: ApiState,
        report: dict[str, object],
    ) -> dict[str, object]:
        captured["state"] = history_state
        captured["report"] = report
        return {
            "backend": "postgres",
            "persisted": True,
            "table": "index_evaluation_runs",
            "run_id": report["run_id"],
        }

    monkeypatch.setattr(
        "medical_audit_kb.api.evaluation_reports.persist_evaluation_history",
        fake_persist_evaluation_history,
    )
    client = TestClient(create_app(state))

    response = client.post(
        "/index/evaluation/run",
        headers={"X-Role": "it-admin"},
        json={
            "retrieval_cases_file": str(retrieval_cases),
            "answer_cases_file": str(answer_cases),
            "max_retrieval_cases": 1,
            "max_answer_cases": 1,
            "top_k": 2,
            "smoke_question": "医保基金审核依据",
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["history"]["persisted"] is True
    assert body["history"]["table"] == "index_evaluation_runs"
    assert body["report"]["history"]["persisted"] is True
    assert captured["state"] is state
    report = captured["report"]
    assert isinstance(report, dict)
    assert report["request"]["top_k"] == 2
    assert report["search_backend"]["backend"] == "postgres"
    assert report["report"]["path"].endswith(".json")


def test_index_evaluation_history_lists_runs_and_records_log(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    state = _api_state(tmp_path)

    def fake_list_evaluation_history(
        history_state: ApiState,
        *,
        limit: int = 20,
    ) -> list[dict[str, object]]:
        assert history_state is state
        assert limit == 20
        return [
            {
                "run_id": "11111111-1111-4111-8111-111111111111",
                "status": "pass",
                "generated_at": "2026-06-02T00:00:00+00:00",
                "retrieval_case_count": 52,
                "answer_case_count": 8,
                "ui_smoke_success": True,
                "report_path": "tmp/knowledge-query-indexes/evaluation-runs/report.json",
                "download_path": "/index/evaluation/latest/export",
                "source": "postgres",
            }
        ]

    monkeypatch.setattr(
        "medical_audit_kb.api.routes_index.list_evaluation_history",
        fake_list_evaluation_history,
    )
    client = TestClient(create_app(state))

    response = client.get("/index/evaluation/history")

    assert response.status_code == 200
    body = response.json()
    assert body["items"][0]["run_id"] == "11111111-1111-4111-8111-111111111111"
    assert body["items"][0]["source"] == "postgres"
    assert state.operation_logs[-1]["action"] == "index-evaluation-history-view"


def test_index_evaluation_export_returns_not_found_before_first_report(tmp_path: Path) -> None:
    client = TestClient(create_app(_api_state(tmp_path)))

    response = client.get("/index/evaluation/latest/export")

    assert response.status_code == 404
    assert "evaluation report not found" in response.json()["detail"]


def test_index_evaluation_run_requires_ready_search_backend(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.search_engine = None
    client = TestClient(create_app(state))

    response = client.post(
        "/index/evaluation/run",
        headers={"X-Role": "it-admin"},
        json={},
    )

    assert response.status_code == 409
    assert "search backend is not ready" in response.json()["detail"]


def _legacy_project_member(
    member_key: str,
    user_identifier: str | None,
    status: str,
) -> dict[str, object]:
    metadata: dict[str, object] = {"fixture": "legacy-duplicate"}
    if user_identifier is not None:
        metadata["user_identifier"] = user_identifier
    member: dict[str, object] = {
        "id": member_key,
        "project_key": "CATALOG-LIMIT-202606",
        "name": member_key,
        "role": "审计员",
        "department": "内审部",
        "status": status,
        "created_by": "legacy-import",
        "source": "custom",
        "metadata": metadata,
    }
    if user_identifier is not None:
        member["user_identifier"] = user_identifier
    return member


def _seed_legacy_project_members(
    database_url: str,
    project_key: str,
    raw_members: list[dict[str, object]],
) -> None:
    engine = create_engine(database_url, connect_args={"check_same_thread": False})
    newest_timestamp = datetime(2026, 7, 12, 12, 0, tzinfo=UTC)
    with Session(engine) as session:
        for index, member in enumerate(raw_members):
            timestamp = newest_timestamp - timedelta(minutes=index)
            metadata = member["metadata"]
            assert isinstance(metadata, dict)
            session.add(
                AuditProjectMember(
                    member_key=str(member["id"]),
                    project_key=project_key,
                    name=str(member["name"]),
                    role=str(member["role"]),
                    department=str(member["department"]),
                    status=str(member["status"]),
                    created_by=str(member["created_by"]),
                    extra_metadata=dict(metadata),
                    created_at=timestamp,
                    updated_at=timestamp,
                )
            )
        session.commit()


class _RecordingProjectMemberStore:
    def __init__(self) -> None:
        self.list_reads = 0

    def list_members(self, project_key: str) -> list[dict[str, object]]:
        self.list_reads += 1
        return []

    def add_member(
        self,
        project_key: str,
        values: dict[str, object],
    ) -> dict[str, object]:
        raise AssertionError("graph tests do not add project members")

    def member_counts(self) -> dict[str, int]:
        raise AssertionError("graph tests do not read member counts")


class _RecordingGraphFindingStore:
    def __init__(
        self,
        items_by_project: dict[str, list[dict[str, object]]] | None = None,
    ) -> None:
        self.items_by_project = items_by_project or {}
        self.requested_project_keys: list[str | None] = []

    def list_findings(
        self,
        *,
        review_status: str | None = None,
        project_key: str | None = None,
        limit: int = 100,
    ) -> list[dict[str, object]]:
        self.requested_project_keys.append(project_key)
        assert review_status is None
        assert project_key is not None
        return [dict(item) for item in self.items_by_project.get(project_key, [])[:limit]]


class _RecordingGraphReviewTaskStore:
    def __init__(self, tasks: list[dict[str, object]] | None = None) -> None:
        self.tasks = tasks or []
        self.list_calls = 0

    def list_tasks(self) -> list[dict[str, object]]:
        self.list_calls += 1
        return [dict(task) for task in self.tasks]


class _CountingSqlAlchemyReviewTaskStore:
    def __init__(self, database_url: str) -> None:
        self._store = SqlAlchemyReviewTaskStore(database_url)
        self.list_calls = 0

    def list_tasks(self) -> list[dict[str, object]]:
        self.list_calls += 1
        return self._store.list_tasks()


def _seed_project_findings(database_url: str) -> None:
    engine = create_engine(database_url, connect_args={"check_same_thread": False})
    with Session(engine) as session:
        rule = AuditRule(
            rule_key="PROJECT-SCOPE-RULE",
            scenario_key="project-scope-test",
            name="项目隔离测试规则",
            status="active",
            owner="unit-test",
        )
        session.add(rule)
        session.flush()
        rule_version = RuleVersion(
            audit_rule_id=rule.id,
            version_key="PROJECT-SCOPE-RULE@v1",
            rule_key=rule.rule_key,
            status="active",
            logic={"fixture": True},
            evidence_links={},
            created_by="unit-test",
        )
        session.add(rule_version)
        session.flush()

        for suffix, project_key in (
            ("self", "SELF-CHECK-FUND-20260607"),
            ("catalog", "CATALOG-LIMIT-202606"),
        ):
            linked_review_task = ReviewTask(
                external_task_id=f"review-{suffix}-sql",
                question=f"{suffix} SENSITIVE-GRAPH-BODY question",
                status="confirmed-violation",
                status_label="确认违规",
                citation_count=1,
                review_gate="负责人确认",
                confidence_label="high",
                fallback_label="manual-review",
                reviewer_note="SENSITIVE-GRAPH-BODY reviewer note",
                conclusion="SENSITIVE-GRAPH-BODY conclusion",
                created_by="unit-test",
                assigned_to="unit-test",
                source="chat-dossier",
                dossier={
                    "report_draft": {
                        "title": f"{suffix} SENSITIVE-GRAPH-BODY report",
                        "summary": "SENSITIVE-GRAPH-BODY summary",
                        "rectification_request": "SENSITIVE-GRAPH-BODY request",
                    },
                    "signed_report": {
                        "status": "signed",
                        "report_id": f"signed-report-{suffix}-sql",
                        "content_sha256": f"sha256-{suffix}",
                        "content": "SENSITIVE-GRAPH-BODY signed content",
                    },
                    "rectification": {
                        "rectification_id": f"rectification-{suffix}-sql",
                        "status": "in-progress",
                        "progress_note": "SENSITIVE-GRAPH-BODY progress",
                    },
                },
            )
            session.add(linked_review_task)
            session.flush()
            session.add(
                ReviewTask(
                    external_task_id=f"report-draft-{suffix}-sql",
                    question="SENSITIVE-GRAPH-BODY template question",
                    status="pending-review",
                    status_label="待复核",
                    citation_count=0,
                    review_gate="manual-review",
                    confidence_label="pending",
                    fallback_label="manual-review",
                    reviewer_note="",
                    conclusion="",
                    created_by="unit-test",
                    assigned_to=None,
                    source="report-template-draft",
                    dossier={
                        "report_template_draft": {
                            "status": "draft",
                            "template_id": f"workpaper-{suffix}",
                            "project_key": project_key,
                            "field_values": {"sensitive": "SENSITIVE-GRAPH-BODY field value"},
                        }
                    },
                )
            )
            project = AuditProject(
                project_key=project_key,
                name=f"{suffix} project",
                scenario_key="project-scope-test",
                status="active",
                created_by="unit-test",
            )
            session.add(project)
            session.flush()
            snapshot = AuditDataSnapshot(
                snapshot_key=f"snapshot-{suffix}",
                project_id=project.id,
                source_batch_key=f"batch-{suffix}",
                time_range={},
                row_counts={"fixture": 1},
                checksum=f"sha256:{suffix}",
                status="validated",
            )
            session.add(snapshot)
            session.flush()
            task = AuditTask(
                task_key=f"task-{suffix}",
                project_id=project.id,
                snapshot_id=snapshot.id,
                topic=f"{suffix} topic",
                department_scope={},
                date_range={},
                status="ready",
                created_by="unit-test",
            )
            session.add(task)
            session.flush()
            run = AuditRun(
                run_key=f"run-{suffix}",
                audit_task_id=task.id,
                snapshot_id=snapshot.id,
                rule_version_key=rule_version.version_key,
                status="succeeded",
                summary={"fixture": True},
            )
            session.add(run)
            session.flush()
            session.add(
                AuditFinding(
                    finding_key=f"finding-{suffix}-sql",
                    audit_run_id=run.id,
                    audit_task_id=task.id,
                    rule_version_id=rule_version.id,
                    snapshot_id=snapshot.id,
                    status="open",
                    finding_type="project-scope-test",
                    severity="medium",
                    source_record_locator={"project_key": project_key},
                    calculation_trace={},
                    review_status="pending-review",
                    review_task_id=linked_review_task.id,
                    extra_metadata={"owner": f"{suffix}-owner"},
                )
            )
        session.add(
            ReviewTask(
                external_task_id="legacy-project-string-sql",
                question=("SELF-CHECK-FUND-20260607 and CATALOG-LIMIT-202606 are text only"),
                status="confirmed-violation",
                status_label="确认违规",
                citation_count=0,
                review_gate="manual-review",
                confidence_label="pending",
                fallback_label="manual-review",
                reviewer_note="",
                conclusion="",
                created_by="unit-test",
                assigned_to=None,
                source="legacy",
                dossier={
                    "metadata": {
                        "project_key": "SELF-CHECK-FUND-20260607",
                        "project_hint": "CATALOG-LIMIT-202606",
                    },
                    "report_draft": {"title": "SENSITIVE-GRAPH-BODY legacy report"},
                },
            )
        )
        session.commit()


def test_agent_market_catalog_exposes_132_plus_featured_contract_agent(
    tmp_path: Path,
) -> None:
    client = TestClient(create_app(_api_state(tmp_path)))

    response = client.get("/api/v1/agent-market/catalog")

    assert response.status_code == 200
    body = response.json()
    assert body["contract_version"] == "agent-market-catalog-v2"
    assert body["count"] == 133
    assert body["featured_count"] == 1
    assert body["items"][0]["id"] == "agent-contract-audit-v2"
    assert body["items"][0]["featured"] is True
    assert all("prompt" not in item for item in body["items"])


def test_agent_market_install_materializes_server_prompt_only(tmp_path: Path) -> None:
    state = _api_state(tmp_path)
    state.agent_store = InMemoryAgentStore()
    client = TestClient(create_app(state))
    headers = {
        "X-User-Id": "admin-1",
        "X-Role": "admin",
        "X-Project-Name": urllib.parse.quote("医保基金使用合规专项自查"),
    }

    response = client.post(
        "/api/v1/agent-market/templates/agent-contract-audit-v2/install",
        headers=headers,
        json={"project_name": "医保基金使用合规专项自查"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["created"] is True
    assert body["item"]["metadata"]["template_id"] == "agent-contract-audit-v2"
    assert "页面证据" in body["item"]["prompt"]


def test_selected_agent_prompt_is_bound_to_generation_call(tmp_path: Path) -> None:
    class AgentAwareProvider:
        provider = "fake"
        model_name = "fake-agent-aware"
        provider_version = "v1"

        def __init__(self) -> None:
            self.prompt = ""
            self.prompt_version_key = ""

        def generate_answer(self, question: str, citations: Sequence[Citation]) -> str:
            raise AssertionError("agent-aware path must be used")

        def generate_answer_with_agent(
            self,
            question: str,
            citations: Sequence[Citation],
            *,
            agent_prompt: str,
            prompt_version_key: str,
        ) -> str:
            self.prompt = agent_prompt
            self.prompt_version_key = prompt_version_key
            return f"已按合同审计智能体执行 {citations[0].marker}"

    provider = AgentAwareProvider()
    state = _api_state(tmp_path)
    state.agent_store = InMemoryAgentStore()
    state.answer_generation_provider = provider
    client = TestClient(create_app(state))

    response = client.post(
        "/query",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        json={
            "question": "请复核合同依据",
            "top_k": 2,
            "agent": "agent-contract-audit-v2",
        },
    )

    assert response.status_code == 200
    assert "合同审计智能体" in provider.prompt
    assert provider.prompt_version_key == "contract-audit-v2@2.0.0"


def test_contract_audit_job_persists_and_downloads_without_provider(
    tmp_path: Path,
) -> None:
    state = _api_state(tmp_path)
    state.answer_generation_provider = None
    client = TestClient(create_app(state))
    headers = {"X-User-Id": "auditor-1", "X-Role": "auditor"}

    response = client.post(
        "/api/v1/contract-audits",
        headers=headers,
        data={"project_name": "采购合同专项", "audit_stage": "签约前"},
        files={
            "file": ("采购合同.txt", "甲方向乙方采购设备，合同价款100万元。".encode(), "text/plain")
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "insufficient_evidence"
    assert body["pages"][0]["mapping_status"] == "resolved"
    job_id = body["job_id"]
    assert (tmp_path / "index" / "contract-audit-jobs" / f"{job_id}.json").is_file()

    persisted = client.get(f"/api/v1/contract-audits/{job_id}", headers=headers)
    markdown = client.get(
        f"/api/v1/contract-audits/{job_id}/report?format=markdown",
        headers=headers,
    )
    docx = client.get(
        f"/api/v1/contract-audits/{job_id}/report?format=docx",
        headers=headers,
    )
    assert persisted.status_code == 200
    assert markdown.status_code == 200
    assert "合同审计报告" in markdown.text
    assert docx.status_code == 200
    assert docx.content.startswith(b"PK")


def test_contract_audit_job_calls_versioned_agent_and_completes(tmp_path: Path) -> None:
    class ContractProvider:
        provider = "fake"
        model_name = "fake-contract-model"
        provider_version = "v1"

        def generate_answer(self, question: str, citations: Sequence[Citation]) -> str:
            raise AssertionError("contract audit must bind its agent prompt")

        def generate_answer_with_agent(
            self,
            question: str,
            citations: Sequence[Citation],
            *,
            agent_prompt: str,
            prompt_version_key: str,
        ) -> str:
            assert "合同审计智能体" in agent_prompt
            assert prompt_version_key == "contract-audit-v2@2.0.0"
            return f"# 审计发现\n付款安排需人工复核 {citations[0].marker}"

    state = _api_state(tmp_path)
    state.answer_generation_provider = ContractProvider()
    client = TestClient(create_app(state))

    response = client.post(
        "/api/v1/contract-audits",
        headers={"X-User-Id": "auditor-1", "X-Role": "auditor"},
        files={"file": ("采购合同.txt", "合同价款100万元，验收后付款。".encode(), "text/plain")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "completed"
    assert body["agent"]["prompt_version_key"] == "contract-audit-v2@2.0.0"
    assert body["result"]["conclusion"]["analysis_markdown"].endswith("[C1]")


def _api_state(tmp_path: Path) -> ApiState:
    source_root = tmp_path / "data"
    source_file = source_root / "全量法律" / "law.md"
    _write_text(
        source_file,
        "\n".join(
            [
                "第一条 医疗机构应当保留医保基金审核依据。",
                "第二条 其他条款。",
            ]
        ),
    )
    settings = KnowledgeQuerySettings(
        data_root=source_root,
        index_root=tmp_path / "index",
        database_url="postgresql+psycopg://user:pass@localhost:5433/db",
        model_provider=ModelProviderSettings(
            provider="fake",
            api_key_env="OPENAI_API_KEY",
            embedding_model="fake",
            chat_model="fake",
        ),
        source_collection_weights={
            "medical-insurance-catalog": 1.25,
            "supervision-rules-knowledge": 1.35,
            "risk-negative-list": 1.1,
            "medical-insurance-laws": 1.0,
        },
    )
    state = ApiState.from_settings(settings)
    state.audit_log_store = None
    state.analytics_upload_store = InMemoryAnalyticsUploadStore(
        upload_root=settings.index_root / "analytics-uploads"
    )
    state.document_upload_store = InMemoryDocumentUploadStore(
        upload_root=settings.index_root / "document-uploads"
    )
    state.query_history_store = InMemoryQueryHistoryStore()
    state.search_engine = _search_engine(
        _chunk_id(),
        source_file.relative_to(source_root).as_posix(),
    )
    return state


def _search_engine(chunk_id: UUID, source_path: str) -> HybridSearchEngine:
    provider = DeterministicFakeEmbeddingProvider(dimension=32)
    chunk = ChunkEmbeddingInput(
        chunk_id=chunk_id,
        text="第一条 医疗机构应当保留医保基金审核依据。",
        metadata={
            "source_collection": SourceCollection.MEDICAL_INSURANCE_LAWS.value,
            "locator": {
                "type": "law-article",
                "source_path": source_path,
                "line_start": 1,
                "line_end": 1,
                "article_number": "第一条",
            },
            "index_version_key": "index-v1",
            "source_package_version_key": "package-v1",
            "title": "医保基金审核依据",
            "source_path": source_path,
            "title_path": ["医保基金审核依据"],
            "year": 2024,
            "region": "国家",
            "document_type": "law",
            "business_topic": "fund-supervision",
        },
    )
    vector_index = InMemoryVectorIndex(dimension=provider.dimension)
    vector_index.upsert(build_chunk_embedding_records([chunk], provider=provider))
    bm25_index = InMemoryBM25Index()
    bm25_index.upsert(
        [
            BM25Document(
                chunk_id=chunk.chunk_id,
                text=chunk.text,
                metadata=chunk.metadata,
            )
        ]
    )
    return HybridSearchEngine(
        embedding_provider=provider,
        vector_index=vector_index,
        bm25_index=bm25_index,
        rerank_provider=FakeRerankProvider(),
    )


def _search_engine_with_personal_materials(
    *,
    system_chunk_id: UUID,
    personal_chunk_id: UUID,
    source_path: str,
) -> HybridSearchEngine:
    provider = DeterministicFakeEmbeddingProvider(dimension=32)
    chunks = [
        ChunkEmbeddingInput(
            chunk_id=system_chunk_id,
            text="第一条 医疗机构应当保留医保基金审核依据。",
            metadata={
                "source_collection": SourceCollection.MEDICAL_INSURANCE_LAWS.value,
                "locator": {
                    "type": "law-article",
                    "source_path": source_path,
                    "line_start": 1,
                    "line_end": 1,
                    "article_number": "第一条",
                },
                "index_version_key": "index-v1",
                "source_package_version_key": "package-v1",
                "title": "医保基金审核依据",
                "source_path": source_path,
                "title_path": ["医保基金审核依据"],
            },
        ),
        ChunkEmbeddingInput(
            chunk_id=personal_chunk_id,
            text="院内个人材料提示：医保基金审核依据需要结合内部台账。",
            metadata={
                "source_collection": SourceCollection.PERSONAL_MATERIALS.value,
                "locator": {
                    "type": "personal-material",
                    "upload_key": "document-upload-private",
                    "source_path": "personal-materials/document-upload-private/note.txt",
                },
                "index_version_key": "personal-materials-test-active",
                "source_package_version_key": "personal-materials-test-active",
                "title": "院内个人材料提示",
                "source_path": "personal-materials/document-upload-private/note.txt",
                "created_by": "auditor-1",
                "visibility": "private",
            },
        ),
    ]
    vector_index = InMemoryVectorIndex(dimension=provider.dimension)
    vector_index.upsert(build_chunk_embedding_records(chunks, provider=provider))
    bm25_index = InMemoryBM25Index()
    bm25_index.upsert(
        [
            BM25Document(
                chunk_id=chunk.chunk_id,
                text=chunk.text,
                metadata=chunk.metadata,
            )
            for chunk in chunks
        ]
    )
    return HybridSearchEngine(
        embedding_provider=provider,
        vector_index=vector_index,
        bm25_index=bm25_index,
        rerank_provider=FakeRerankProvider(),
    )


def _chunk_id() -> UUID:
    return uuid4()


def _clear_chat_model_env(monkeypatch: pytest.MonkeyPatch) -> None:
    aliases = ("KIMI_2_7", "DEEPSEEK_V4_PRO")
    suffixes = (
        "PROVIDER",
        "API_KEY_ENV",
        "MODEL",
        "BASE_URL",
        "MAX_OUTPUT_TOKENS",
        "TEMPERATURE",
        "THINKING_MODE",
    )
    for alias in aliases:
        for suffix in suffixes:
            monkeypatch.delenv(f"MEDICAL_AUDIT_KB_CHAT_MODEL_{alias}_{suffix}", raising=False)
    monkeypatch.delenv("TEST_KIMI_CHAT_KEY", raising=False)
    monkeypatch.delenv("TEST_DEEPSEEK_CHAT_KEY", raising=False)
    monkeypatch.delenv("MEDICAL_AUDIT_KB_ALLOW_FAKE_CHAT_MODELS", raising=False)


class StaticApiAnswerProvider:
    provider = "fake"
    model_name = "fake-chat"
    provider_version = "v1"

    def generate_answer(self, question: str, citations: Sequence[Citation]) -> str:
        return f"生成模型回答：应核验医保基金审核依据 {citations[0].marker}。"


class FailingApiAnswerProvider:
    provider = "fake"
    model_name = "fake-chat"
    provider_version = "v1"

    def generate_answer(self, question: str, citations: Sequence[Citation]) -> str:
        _ = question, citations
        raise RuntimeError("private provider detail")


class HttpStatusFailingApiAnswerProvider:
    provider = "fake"
    model_name = "fake-chat"
    provider_version = "v1"

    def generate_answer(self, question: str, citations: Sequence[Citation]) -> str:
        _ = question, citations
        raise AnswerProviderError(
            "answer generation request failed: HTTP 429",
            code="provider_http_status",
            http_status=429,
        )


class InvalidResponseApiAnswerProvider:
    provider = "fake"
    model_name = "invalid-response-chat"
    provider_version = "v1"

    def generate_answer(self, question: str, citations: Sequence[Citation]) -> str:
        _ = question, citations
        raise AnswerProviderError(
            "private provider sentinel",
            code="provider_response_invalid",
            reason="deepseek_content_invalid_json",
        )


class AbstainingApiAnswerProvider:
    provider = "deepseek"
    model_name = "deepseek-v4-pro"
    provider_version = "v1"

    def generate_answer(self, question: str, citations: Sequence[Citation]) -> str:
        _ = question, citations
        raise AnswerProviderError(
            "private provider abstention sentinel",
            code="provider_abstention",
            reason="deepseek_strict_insufficient_evidence",
        )


class FailingQueryHistoryStore:
    def add_query(self, _values: dict[str, object]) -> dict[str, object]:
        raise RuntimeError("history database unavailable")

    def list_queries(
        self,
        *,
        limit: int = 20,
        user_identifier: str | None = None,
    ) -> list[dict[str, object]]:
        _ = limit, user_identifier
        raise RuntimeError("history database unavailable")


def _write_text(path: Path, content: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    return path
